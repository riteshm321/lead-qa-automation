import datetime
import os
from unittest.mock import patch

import openpyxl
import pandas as pd
from streamlit.testing.v1 import AppTest

from core.app_settings import get_clients_dir, save_enhancio_client_id, save_app_settings
from core.enhancio_sync import save_pending_leads, load_pending_leads, save_uploaded_emails, load_uploaded_emails
from core.models import ClientProfile, FieldMapping, EnhancioConfig, EnhancioAllocationMapping
from core.profile_store import save_profile

_PAGE_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "pages", "8_Enhancio.py")


def _make_accumulated(path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accumulated"
    ws.append(["Date", "CID", "Email", "First Name", "Last Name", "Company"])
    wb.create_sheet("Refund").append(
        ["Date", "CID", "Email", "First Name", "Last Name", "Company", "Refund Reason"])
    wb.save(path)


def _save_profile(acc_path: str, jira_ticket_key: str = "", accumulated_report_link: str = "") -> ClientProfile:
    fm = FieldMapping(email="Email", first_name="First Name", last_name="Last Name", company="Company", cid="CID")
    profile = ClientProfile(
        name="Amazon Business EMEA", accumulated_report_path=acc_path, field_mapping=fm,
        jira_ticket_key=jira_ticket_key, accumulated_report_link=accumulated_report_link,
        enhancio=EnhancioConfig(
            enabled=True,
            allocations=[
                EnhancioAllocationMapping(cid="120022", allocation_uid="L-22256"),
                EnhancioAllocationMapping(cid="120028", allocation_uid="L-22257"),
            ],
            field_mapping={"Email": "Email Address", "First Name": "First Name"},
        ),
    )
    save_profile(profile, get_clients_dir())
    return profile


def test_upload_fills_micro_audience_for_a_cid_covered_by_the_box_tracker_rule(tmp_path, monkeypatch):
    # CID 118741 ("Bob") is one of IBM APAC's fixed micro_audience CIDs
    # (core.box_tracker._MICRO_AUDIENCE_BY_CID -> "Platform_SWE") -- the
    # same rule the Box Tracker Lead Template already applies must also
    # fill it in here, purely from the CID, with no LOB/micro_audience
    # column in the leadfile at all.
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    fm = FieldMapping(email="Email", first_name="First Name", last_name="Last Name", company="Company", cid="CID")
    profile = ClientProfile(
        name="IBM APAC Interactive Avenues Pvt Ltd", accumulated_report_path=acc_path, field_mapping=fm,
        enhancio=EnhancioConfig(
            enabled=True,
            allocations=[EnhancioAllocationMapping(cid="118741", allocation_uid="L-22SD7")],
            field_mapping={"Email": "Email Address", "micro_audience": "Micro Audience"},
        ),
    )
    save_profile(profile, get_clients_dir())
    save_enhancio_client_id("CID123")

    leads_csv = tmp_path / "leads.csv"
    pd.DataFrame([{"CID": "118741", "Email": "a@x.com", "First Name": "A"}]).to_csv(leads_csv, index=False)

    captured_leads = {}

    def _fake_import_leads(token, allocation_uid, leads):
        captured_leads["leads"] = leads
        return {"submitted": [
            {"leadId": "lead-1", "status": "Submitted", "email": leads[0]["Email Address"]},
        ], "errors": []}

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.import_leads", side_effect=_fake_import_leads):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value(
            "IBM APAC Interactive Avenues Pvt Ltd").run()
        with open(leads_csv, "rb") as f:
            at.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()
        next(b for b in at.button if b.label == "Upload to Enhancio").click().run()
        assert not at.exception

    assert captured_leads["leads"][0]["Micro Audience"] == "Platform_SWE"


def _make_accumulated_with_status(path: str, rows: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accumulated"
    ws.append(["Date", "CID", "Email", "First Name", "Last Name", "Company", "Status"])
    for row in rows:
        ws.append([
            row.get("Date"), row["CID"], row["Email"], row.get("First Name", ""),
            row.get("Last Name", ""), row.get("Company", ""), row.get("Status", ""),
        ])
    wb.create_sheet("Refund").append(
        ["Date", "CID", "Email", "First Name", "Last Name", "Company", "Refund Reason"])
    wb.save(path)


def test_pull_from_accumulated_report_filters_by_date_range_and_stamps_status(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    today = datetime.date.today()
    yesterday = today - datetime.timedelta(days=1)
    too_old = today - datetime.timedelta(days=10)
    _make_accumulated_with_status(acc_path, [
        {"Date": yesterday, "CID": "120022", "Email": "in_range@x.com",
         "First Name": "A", "Last Name": "One", "Company": "Acme"},
        {"Date": too_old, "CID": "120022", "Email": "too_old@x.com",
         "First Name": "B", "Last Name": "Two", "Company": "Acme"},
    ])
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")

    def _fake_import_leads(token, allocation_uid, leads):
        return {"submitted": [
            {"leadId": f"lead-{i}", "status": "Submitted", "email": lead["Email Address"]}
            for i, lead in enumerate(leads)
        ], "errors": []}

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.import_leads", side_effect=_fake_import_leads):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        at.radio(key="enhancio_lead_source").set_value("Pull from Accumulated Report by date range").run()
        at.date_input(key="enhancio_range_start").set_value(yesterday).run()
        at.date_input(key="enhancio_range_end").set_value(today).run()
        next(b for b in at.button if b.label == "Upload to Enhancio").click().run()
        assert not at.exception

    results_df = at.session_state["enhancio_upload_results"]
    assert results_df["Result"].str.startswith("✅").sum() == 1
    assert "in_range@x.com" in results_df["Email"].values
    assert "too_old@x.com" not in results_df["Email"].values

    accumulated_df = pd.read_excel(acc_path, sheet_name="Accumulated")
    status = accumulated_df.loc[accumulated_df["Email"] == "in_range@x.com", "Status"].iloc[0]
    assert status.startswith("Uploaded to Enhancio")
    old_status = accumulated_df.loc[accumulated_df["Email"] == "too_old@x.com", "Status"].iloc[0]
    assert old_status == "" or pd.isna(old_status)


def test_warns_when_no_client_has_enhancio_enabled(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    assert not at.exception
    assert any("No client has Enhancio enabled" in w.value for w in at.warning)


def test_upload_batches_leads_by_allocation_and_reports_results(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")

    leads_csv = tmp_path / "leads.csv"
    pd.DataFrame([
        {"CID": "120022", "Email": "a@x.com", "First Name": "A", "Last Name": "One", "Company": "Acme"},
        {"CID": "120028", "Email": "b@x.com", "First Name": "B", "Last Name": "Two", "Company": "Acme"},
    ]).to_csv(leads_csv, index=False)

    def _fake_import_leads(token, allocation_uid, leads):
        return {"submitted": [
            {"leadId": f"lead-{allocation_uid}-{i}", "status": "Submitted", "email": lead["Email Address"]}
            for i, lead in enumerate(leads)
        ], "errors": []}

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.import_leads", side_effect=_fake_import_leads):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(leads_csv, "rb") as f:
            at.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()
        next(b for b in at.button if b.label == "Upload to Enhancio").click().run()
        assert not at.exception

    results_df = at.session_state["enhancio_upload_results"]
    assert results_df["Result"].str.startswith("✅").sum() == 2
    assert any("lead-L-22256" in r for r in results_df["Result"])
    assert any("lead-L-22257" in r for r in results_df["Result"])


def test_preview_shows_leads_to_send_without_calling_the_api(tmp_path, monkeypatch):
    # The user must be able to see exactly what would be sent, and
    # download it, before ever clicking "Upload to Enhancio" -- this
    # must never call import_leads.
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")

    leads_csv = tmp_path / "leads.csv"
    pd.DataFrame([
        {"CID": "120022", "Email": "a@x.com", "First Name": "A", "Last Name": "One", "Company": "Acme"},
        {"CID": "120028", "Email": "b@x.com", "First Name": "B", "Last Name": "Two", "Company": "Acme"},
    ]).to_csv(leads_csv, index=False)

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.import_leads") as mock_import_leads:
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(leads_csv, "rb") as f:
            at.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()

        assert not at.exception
        mock_import_leads.assert_not_called()

    preview_expander = next(e for e in at.expander if e.label.startswith("📋 Preview leads to send"))
    assert "2 lead(s)" in preview_expander.label
    assert "2 allocation(s)" in preview_expander.label
    assert any(dl.key == "enhancio_preview_download" for dl in at.download_button)


def test_upload_keeps_and_saves_successes_when_batch_also_has_rejected_leads(tmp_path, monkeypatch):
    # Regression test for a real production incident: Enhancio accepted 29
    # of 160 leads in one batch and rejected 131 as duplicates, all in the
    # SAME response. The tool must still report the 29 as uploaded and save
    # their lead IDs for Reconcile, instead of discarding everything and
    # reporting 0 uploaded just because the response also carried errors.
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    profile = _save_profile(acc_path)
    save_enhancio_client_id("CID123")

    leads_csv = tmp_path / "leads.csv"
    pd.DataFrame([
        {"CID": "120022", "Email": "ok1@x.com", "First Name": "A", "Last Name": "One", "Company": "Acme"},
        {"CID": "120022", "Email": "ok2@x.com", "First Name": "B", "Last Name": "Two", "Company": "Acme"},
        {"CID": "120022", "Email": "dup@x.com", "First Name": "C", "Last Name": "Three", "Company": "Acme"},
    ]).to_csv(leads_csv, index=False)

    def _fake_import_leads(token, allocation_uid, leads):
        return {
            "submitted": [
                {"leadId": "lead-ok1", "status": "Submitted", "email": "ok1@x.com"},
                {"leadId": "lead-ok2", "status": "Submitted", "email": "ok2@x.com"},
            ],
            "errors": [
                {"message": "Duplicate lead within the campaign allocation", "email": "dup@x.com"},
            ],
        }

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.import_leads", side_effect=_fake_import_leads):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(leads_csv, "rb") as f:
            at.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()
        next(b for b in at.button if b.label == "Upload to Enhancio").click().run()
        assert not at.exception

    results_df = at.session_state["enhancio_upload_results"]
    assert results_df["Result"].str.startswith("✅").sum() == 2
    assert results_df["Result"].str.startswith("❌").sum() == 1
    assert any("Duplicate lead within the campaign allocation" in w.value for w in at.warning)

    pending = load_pending_leads(profile.name)
    assert set(pending.keys()) == {"lead-ok1", "lead-ok2"}


def test_reset_button_clears_an_allocations_already_uploaded_memory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")
    save_uploaded_emails("Amazon Business EMEA", "L-22256", {"a@x.com"})

    leads_csv = tmp_path / "leads.csv"
    pd.DataFrame([{"CID": "120022", "Email": "a@x.com", "First Name": "A"}]).to_csv(leads_csv, index=False)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
    with open(leads_csv, "rb") as f:
        at.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()

    assert load_uploaded_emails("Amazon Business EMEA", "L-22256") == {"a@x.com"}
    at.button(key="enhancio_reset_L-22256").click().run()
    assert not at.exception

    assert load_uploaded_emails("Amazon Business EMEA", "L-22256") == set()
    assert any("Cleared already-uploaded memory" in t.value for t in at.toast)


def test_reuploading_the_same_file_skips_leads_already_uploaded_to_that_allocation(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")

    leads_csv = tmp_path / "leads.csv"
    pd.DataFrame([
        {"CID": "120022", "Email": "a@x.com", "First Name": "A", "Last Name": "One", "Company": "Acme"},
        {"CID": "120028", "Email": "b@x.com", "First Name": "B", "Last Name": "Two", "Company": "Acme"},
    ]).to_csv(leads_csv, index=False)

    import_calls = []

    def _fake_import_leads(token, allocation_uid, leads):
        import_calls.append((allocation_uid, leads))
        return {"submitted": [
            {"leadId": f"lead-{allocation_uid}-{i}", "status": "Submitted", "email": lead["Email Address"]}
            for i, lead in enumerate(leads)
        ], "errors": []}

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.import_leads", side_effect=_fake_import_leads):
        # First upload: both leads go through.
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(leads_csv, "rb") as f:
            at.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()
        next(b for b in at.button if b.label == "Upload to Enhancio").click().run()
        assert not at.exception
        assert len(import_calls) == 2  # one call per allocation, L-22256 and L-22257

        import_calls.clear()

        # Second upload of the SAME file: nothing should reach import_leads
        # at all -- both leads are already-uploaded for their allocation.
        at2 = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at2.run()
        next(s for s in at2.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(leads_csv, "rb") as f:
            at2.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()
        next(b for b in at2.button if b.label == "Upload to Enhancio").click().run()
        assert not at2.exception

    assert import_calls == []
    results_df = at2.session_state["enhancio_upload_results"]
    assert results_df["Result"].str.startswith("⏭️").sum() == 2
    assert all("already uploaded to this allocation previously" in r for r in results_df["Result"])


def test_reupload_checkbox_lets_you_resend_an_already_uploaded_lead(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")

    leads_csv = tmp_path / "leads.csv"
    pd.DataFrame([{"CID": "120022", "Email": "a@x.com", "First Name": "A"}]).to_csv(leads_csv, index=False)

    import_calls = []

    def _fake_import_leads(token, allocation_uid, leads):
        import_calls.append((allocation_uid, leads))
        return {"submitted": [
            {"leadId": f"lead-{allocation_uid}-{len(import_calls)}-{i}", "status": "Submitted",
             "email": lead["Email Address"]} for i, lead in enumerate(leads)
        ], "errors": []}

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.import_leads", side_effect=_fake_import_leads):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(leads_csv, "rb") as f:
            at.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()
        next(b for b in at.button if b.label == "Upload to Enhancio").click().run()
        assert not at.exception
        assert len(import_calls) == 1

        # Re-uploading the same file without the checkbox: skipped.
        at2 = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at2.run()
        next(s for s in at2.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(leads_csv, "rb") as f:
            at2.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()
        assert any("already uploaded" in w.value for w in at2.warning)
        next(b for b in at2.button if b.label == "Upload to Enhancio").click().run()
        assert not at2.exception
        assert len(import_calls) == 1

        # Checking "upload anyway" resends it to the same allocation.
        at3 = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at3.run()
        next(s for s in at3.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(leads_csv, "rb") as f:
            at3.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()
        at3.checkbox(key="enhancio_reupload_duplicates").set_value(True).run()
        next(b for b in at3.button if b.label == "Upload to Enhancio").click().run()
        assert not at3.exception

    assert len(import_calls) == 2


def test_same_email_can_still_upload_to_a_different_allocation(tmp_path, monkeypatch):
    # The same email routed through two different CIDs mapped to two
    # different allocations must succeed for BOTH -- dedupe is per
    # allocation, not per client.
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")

    def _fake_import_leads(token, allocation_uid, leads):
        return {"submitted": [
            {"leadId": f"lead-{allocation_uid}-{i}", "status": "Submitted", "email": lead["Email Address"]}
            for i, lead in enumerate(leads)
        ], "errors": []}

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.import_leads", side_effect=_fake_import_leads):
        first_csv = tmp_path / "first.csv"
        pd.DataFrame([{"CID": "120022", "Email": "shared@x.com"}]).to_csv(first_csv, index=False)
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(first_csv, "rb") as f:
            at.get("file_uploader")[0].set_value(("first.csv", f.read(), "text/csv")).run()
        next(b for b in at.button if b.label == "Upload to Enhancio").click().run()
        assert not at.exception

        second_csv = tmp_path / "second.csv"
        pd.DataFrame([{"CID": "120028", "Email": "shared@x.com"}]).to_csv(second_csv, index=False)
        at2 = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at2.run()
        next(s for s in at2.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(second_csv, "rb") as f:
            at2.get("file_uploader")[0].set_value(("second.csv", f.read(), "text/csv")).run()
        next(b for b in at2.button if b.label == "Upload to Enhancio").click().run()
        assert not at2.exception

    results_df = at2.session_state["enhancio_upload_results"]
    assert results_df["Result"].str.startswith("✅").sum() == 1


def test_upload_applies_fixed_field_values_only_to_the_matching_allocation(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    fm = FieldMapping(email="Email", first_name="First Name", last_name="Last Name", company="Company", cid="CID")
    profile = ClientProfile(
        name="Amazon Business EMEA", accumulated_report_path=acc_path, field_mapping=fm,
        enhancio=EnhancioConfig(
            enabled=True,
            allocations=[
                EnhancioAllocationMapping(cid="120022", allocation_uid="L-22256"),
                EnhancioAllocationMapping(cid="120028", allocation_uid="L-22257"),
            ],
            field_mapping={"Email": "Email Address"},
            fixed_field_values={"L-22256": {"Company Size": "1M - 5M"}},
        ),
    )
    save_profile(profile, get_clients_dir())
    save_enhancio_client_id("CID123")

    leads_csv = tmp_path / "leads.csv"
    pd.DataFrame([
        {"CID": "120022", "Email": "a@x.com"},
        {"CID": "120028", "Email": "b@x.com"},
    ]).to_csv(leads_csv, index=False)

    captured_calls = {}

    def _fake_import_leads(token, allocation_uid, leads):
        captured_calls[allocation_uid] = leads
        return {"submitted": [
            {"leadId": f"lead-{allocation_uid}-{i}", "status": "Submitted", "email": lead["Email Address"]}
            for i, lead in enumerate(leads)
        ], "errors": []}

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.import_leads", side_effect=_fake_import_leads):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(leads_csv, "rb") as f:
            at.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()
        next(b for b in at.button if b.label == "Upload to Enhancio").click().run()
        assert not at.exception

    assert captured_calls["L-22256"][0]["Company Size"] == "1M - 5M"
    assert "Company Size" not in captured_calls["L-22257"][0]


def test_upload_reformats_a_date_field_to_enhancios_required_format(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    fm = FieldMapping(email="Email", first_name="First Name", last_name="Last Name", company="Company", cid="CID")
    profile = ClientProfile(
        name="Amazon Business EMEA", accumulated_report_path=acc_path, field_mapping=fm,
        enhancio=EnhancioConfig(
            enabled=True,
            allocations=[EnhancioAllocationMapping(cid="120022", allocation_uid="L-22256")],
            field_mapping={"Email": "Email Address", "Created Timestamp": "Created Timestamp"},
        ),
    )
    save_profile(profile, get_clients_dir())
    save_enhancio_client_id("CID123")

    leads_csv = tmp_path / "leads.csv"
    # A CSV has no native date type -- this is the common real case where a
    # leadfile's date column round-trips as plain text, not a real datetime.
    pd.DataFrame([{"CID": "120022", "Email": "a@x.com", "Created Timestamp": "2026-03-05 14:30:00"}]).to_csv(
        leads_csv, index=False)

    captured = {}

    def _fake_import_leads(token, allocation_uid, leads):
        captured["leads"] = leads
        return {"submitted": [{"leadId": "lead-1", "status": "Submitted", "email": leads[0]["Email Address"]}],
                "errors": []}

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.import_leads", side_effect=_fake_import_leads):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        with open(leads_csv, "rb") as f:
            at.get("file_uploader")[0].set_value(("leads.csv", f.read(), "text/csv")).run()
        next(b for b in at.button if b.label == "Upload to Enhancio").click().run()
        assert not at.exception

    assert captured["leads"][0]["Created Timestamp"] == "03-05-2026 14:30:00"


def test_reconcile_writes_accepted_to_accumulated_and_rejected_to_refund_with_cid_and_reason(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")
    save_pending_leads("Amazon Business EMEA", {
        "101": {"Email": "accepted@x.com", "CID": "120022", "First Name": "A", "Last Name": "One"},
        "102": {"Email": "rejected@x.com", "CID": "120028", "First Name": "B", "Last Name": "Two"},
    })

    def _fake_get_lead_status(token, lead_ids):
        return [
            {"leadId": "101", "status": "Accepted", "email": "accepted@x.com"},
            {"leadId": "102", "status": "Rejected", "email": "rejected@x.com", "rejectionReason": "Lead Duplicate"},
        ]

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.get_lead_status", side_effect=_fake_get_lead_status):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        fetch_button = next(b for b in at.button if b.label == "Fetch decisions from Enhancio")
        fetch_button.click().run()
        assert not at.exception

        write_button = next(b for b in at.button if b.label == "Write to Accumulated & Refund")
        write_button.click().run()
        assert not at.exception

    accumulated_df = pd.read_excel(acc_path, sheet_name="Accumulated")
    assert len(accumulated_df) == 1
    assert accumulated_df.loc[0, "Email"] == "accepted@x.com"
    assert str(accumulated_df.loc[0, "CID"]) == "120022"
    assert pd.notna(accumulated_df.loc[0, "Date"])

    refund_df = pd.read_excel(acc_path, sheet_name="Refund")
    assert len(refund_df) == 1
    assert refund_df.loc[0, "Email"] == "rejected@x.com"
    assert str(refund_df.loc[0, "CID"]) == "120028"
    assert refund_df.loc[0, "Refund Reason"] == "Lead Duplicate"


def test_reconcile_prefers_comments_over_rejection_reason_for_refund_reason(tmp_path, monkeypatch):
    # comments carries the actual detail; rejectionReason is often just a
    # terse code -- comments must win when both are present.
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")
    save_pending_leads("Amazon Business EMEA", {"102": {"Email": "rejected@x.com", "CID": "120028"}})

    def _fake_get_lead_status(token, lead_ids):
        return [{"leadId": "102", "status": "Rejected", "email": "rejected@x.com",
                  "rejectionReason": "Lead Duplicate",
                  "comments": "Lead validation failed: Duplicate lead within the campaign allocation"}]

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.get_lead_status", side_effect=_fake_get_lead_status):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        next(b for b in at.button if b.label == "Fetch decisions from Enhancio").click().run()
        next(b for b in at.button if b.label == "Write to Accumulated & Refund").click().run()
        assert not at.exception

    refund_df = pd.read_excel(acc_path, sheet_name="Refund")
    assert refund_df.loc[0, "Refund Reason"] == "Lead validation failed: Duplicate lead within the campaign allocation"


def test_write_to_accumulated_shows_a_toast_that_survives_the_rerun(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")
    save_pending_leads("Amazon Business EMEA", {"101": {"Email": "accepted@x.com", "CID": "120022"}})

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.get_lead_status", return_value=[
             {"leadId": "101", "status": "Accepted", "email": "accepted@x.com"}]):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        next(b for b in at.button if b.label == "Fetch decisions from Enhancio").click().run()
        next(b for b in at.button if b.label == "Write to Accumulated & Refund").click().run()
        assert not at.exception

    assert any("1 accepted" in t.value and "0 rejected" in t.value for t in at.toast)


def test_reconcile_leaves_unresolved_leads_pending(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path)
    save_enhancio_client_id("CID123")
    save_pending_leads("Amazon Business EMEA", {"201": {"Email": "still@x.com", "CID": "120022"}})

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.get_lead_status", return_value=[]):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        next(b for b in at.button if b.label == "Fetch decisions from Enhancio").click().run()
        assert not at.exception
        assert any("No new decided leads" in c.value for c in at.caption)

    from core.enhancio_sync import load_pending_leads
    assert load_pending_leads("Amazon Business EMEA") == {"201": {"Email": "still@x.com", "CID": "120022"}}


def test_jira_section_hidden_without_a_ticket_configured(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    _make_accumulated(acc_path)
    _save_profile(acc_path)  # no jira_ticket_key

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    assert not at.exception
    assert any("No Jira ticket configured" in c.value for c in at.caption)
    assert not any(t.key == "enhancio_jira_message" for t in at.text_area)


def test_jira_section_posts_a_summary_after_reconcile(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    save_app_settings({"shared_root_dir": str(tmp_path / "Shared")})
    _make_accumulated(acc_path)
    _save_profile(acc_path, jira_ticket_key="PROJ-1234",
                  accumulated_report_link="https://madlog.sharepoint.com/:x:/s/Team/AccLink")
    save_enhancio_client_id("CID123")
    save_pending_leads("Amazon Business EMEA", {"301": {"Email": "accepted@x.com", "CID": "120022"}})

    from core.app_settings import save_jira_settings
    save_jira_settings("https://example.atlassian.net", "me@example.com", "token123")

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.get_lead_status",
               return_value=[{"leadId": "301", "status": "Accepted", "email": "accepted@x.com"}]):
        at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
        at.run()
        next(s for s in at.selectbox if s.label == "Client").set_value("Amazon Business EMEA").run()
        next(b for b in at.button if b.label == "Fetch decisions from Enhancio").click().run()
        next(b for b in at.button if b.label == "Write to Accumulated & Refund").click().run()

        message_box = next(t for t in at.text_area if t.key == "enhancio_jira_message")
        assert "1 accepted, 0 rejected" in message_box.value

        with patch("core.jira_client.post_comment_body") as mock_post:
            post_button = next(b for b in at.button if b.key == "enhancio_jira_post")
            post_button.click().run()
            assert not at.exception
            mock_post.assert_called_once()
            args, _ = mock_post.call_args
            assert args[0] == "https://example.atlassian.net"
            assert args[3] == "PROJ-1234"
            adf_body = args[4]
            assert "https://madlog.sharepoint.com/:x:/s/Team/AccLink" in str(adf_body)
            assert "Accumulated File" in str(adf_body)
