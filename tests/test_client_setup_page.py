import os
from unittest.mock import patch

import openpyxl
from streamlit.testing.v1 import AppTest

_PAGE_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "pages", "1_Client_Setup.py")


def test_saving_a_jira_ticket_link_normalizes_to_the_bare_key(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    next(t for t in at.text_input if t.label == "Client name").set_value("Jira Link Test Client").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    next(t for t in at.text_input if t.label == "Jira ticket key or link (optional)").set_value(
        "https://yourteam.atlassian.net/browse/PROJ-9876"
    ).run()

    save_button = next(b for b in at.button if "Save Client Profile" in b.label)
    save_button.click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Jira Link Test Client", get_clients_dir())
    assert loaded.jira_ticket_key == "PROJ-9876"


def test_client_name_with_path_separator_is_rejected(tmp_path, monkeypatch):
    # Regression test: the client name becomes a bare "<name>.json" filename
    # under clients_dir with no sanitizing — a "/" silently created a
    # nested, orphaned profile file that the client dropdown's flat
    # directory scan could never show again, and ".." could escape
    # clients_dir onto an arbitrary path on disk.
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    next(t for t in at.text_input if t.label == "Client name").set_value("Acme/Corp").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()

    save_button = next(b for b in at.button if "Save Client Profile" in b.label)
    save_button.click().run()
    assert not at.exception

    assert any("can't contain" in e.value for e in at.error)
    from core.app_settings import get_clients_dir
    from core.profile_store import list_profile_names
    assert list_profile_names(get_clients_dir()) == []


def test_lead_template_mapping_reads_from_first_tabs_own_file_when_shared_path_is_blank(tmp_path, monkeypatch):
    # Regression test: with per-CID Lead Template routing, a client can have
    # every tab point at its own separate file and leave the shared "Lead
    # Template path" completely blank (no default needed). The column
    # mapping expander previously always read headers from that shared path
    # only, so it silently showed "enter a valid path" even though the
    # tab's own file was perfectly valid.
    monkeypatch.chdir(tmp_path)
    tab_file = str(tmp_path / "emea_only.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EMEA"
    ws.append(["Email_Address", "First_Name", "Last_Name", "Company_Name", "CID"])
    wb.save(tab_file)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    multi_tab_checkbox = next(c for c in at.checkbox if "Route different CIDs" in c.label)
    multi_tab_checkbox.set_value(True).run()
    assert not at.exception

    add_tab_button = next(b for b in at.button if b.key == "lead_template_tabs_add")
    add_tab_button.click().run()
    assert not at.exception

    tab_file_input = next(t for t in at.text_input if t.label.startswith("File for this tab"))
    tab_file_input.set_value(tab_file).run()
    assert not at.exception

    sheet_select = next(s for s in at.selectbox if s.label == "Tab (sheet) name")
    assert "EMEA" in sheet_select.options
    sheet_select.set_value("EMEA").run()
    assert not at.exception

    email_select = at.selectbox(key="tmpl_map_email")
    assert "Email_Address" in email_select.options


def test_saving_lead_template_mapping_persists_mandatory_and_override(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    # Build a minimal real Lead Template workbook so _safe_read_template_headers
    # has something to read.
    template_path = str(tmp_path / "template.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Email", "Company Size"])
    wb.save(template_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    next(t for t in at.text_input if t.label == "Client name").set_value("LTM Test Client").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_input(key="lead_template_path_input").set_value(template_path).run()
    at.selectbox(key="lead_template_sheet_select").set_value("Sheet1").run()

    at.checkbox(key="ltm_mandatory_Company Size").set_value(True).run()

    next(b for b in at.button if "Save Client Profile" in b.label).click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    saved = load_profile("LTM Test Client", get_clients_dir())
    rule = next(r for r in saved.lead_template_mapping.rules if r.template_column == "Company Size")
    assert rule.mandatory is True


def test_a_lead_template_column_left_at_every_default_is_not_saved_as_a_rule(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    template_path = str(tmp_path / "template.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Email", "Company Size"])
    wb.save(template_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    next(t for t in at.text_input if t.label == "Client name").set_value("LTM Default Client").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_input(key="lead_template_path_input").set_value(template_path).run()
    at.selectbox(key="lead_template_sheet_select").set_value("Sheet1").run()

    # Deliberately do NOT touch the "Company Size" mandatory checkbox,
    # source override, or date format -- every Lead Template column left at
    # its default must not turn into a saved rule.

    next(b for b in at.button if "Save Client Profile" in b.label).click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    saved = load_profile("LTM Default Client", get_clients_dir())
    assert saved.lead_template_mapping.rules == []


def test_accumulated_field_mapping_saves_as_none_when_every_dropdown_left_unset(tmp_path, monkeypatch):
    # Regression test for a real bug: this section is documented as
    # "(optional)" -- leaving every dropdown at "No mapping" (the default)
    # used to still save a FieldMapping with every field blank instead of
    # None. Since `saved_mapping or fallback`-style code elsewhere treats
    # any FieldMapping object as valid regardless of its field values,
    # that blank-but-present mapping silently won over the fallback it
    # was supposed to defer to -- confirmed in a real client's profile
    # (Company/CID came out blank everywhere it was consulted).
    monkeypatch.chdir(tmp_path)
    acc_path = str(tmp_path / "accumulated.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accumulated"
    ws.append(["Email_Address", "First_Name", "Last_Name", "Company_Name", "CID"])
    wb.save(acc_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    next(t for t in at.text_input if t.label == "Client name").set_value("Blank Mapping Test Client").run()
    at.text_input(key="accumulated_path_input").set_value(acc_path).run()

    # Explicitly set every dropdown to "no mapping" -- some header names
    # here are auto-guessable (a seeded UI default, not a user choice), so
    # this is the only way to reliably simulate "the user left this whole
    # optional section unmapped" regardless of what got guessed.
    _no_mapping = "(none — this file has no such column)"
    for key in ("acc_map_email", "acc_map_first", "acc_map_last", "acc_map_company", "acc_map_cid"):
        at.selectbox(key=key).set_value(_no_mapping).run()

    next(b for b in at.button if "Save Client Profile" in b.label).click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Blank Mapping Test Client", get_clients_dir())
    assert loaded.accumulated_field_mapping is None


def test_exclusion_source_accepts_a_csv_file_and_reads_its_columns_and_saves(tmp_path, monkeypatch):
    # Reference sources (Exclusion, TAL, Suppression, Dedupe) were Excel-only
    # -- picking a .csv here used to either error or leave the sheet/column
    # pickers empty. A CSV has no real sheets, so the picker must fall back
    # to one fixed pseudo-sheet rather than breaking.
    monkeypatch.chdir(tmp_path)
    csv_path = str(tmp_path / "exclusion.csv")
    with open(csv_path, "w", encoding="utf-8") as f:
        f.write("Account Name,Domain\nExcluded Co,excluded.com\n")

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    next(t for t in at.text_input if t.label == "Client name").set_value("CSV Test Client").run()
    next(c for c in at.checkbox if c.label == "Enable Exclusion check").set_value(True).run()
    assert not at.exception

    add_button = next(b for b in at.button if b.label == "➕ Add Exclusion Source")
    add_button.click().run()
    assert not at.exception

    next(t for t in at.text_input if t.label == "Name").set_value("Global").run()
    file_path_input = next(t for t in at.text_input if t.label == "File path")
    file_path_input.set_value(csv_path).run()
    assert not at.exception

    sheet_select = next(s for s in at.selectbox if s.label == "Sheet")
    assert sheet_select.options == ["(CSV file)"]

    domain_select = next(s for s in at.selectbox if s.label == "Domain column")
    assert "Domain" in domain_select.options
    assert "Account Name" in domain_select.options

    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    save_button = next(b for b in at.button if "Save Client Profile" in b.label)
    save_button.click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("CSV Test Client", get_clients_dir())
    assert loaded.exclusion.enabled is True
    assert loaded.exclusion.sources[0].file_path == csv_path
    assert loaded.exclusion.sources[0].sheet_name == "(CSV file)"


def test_collation_checkbox_saves(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    collation_checkbox = next(c for c in at.checkbox if c.label == "Enable file collation for this client")
    assert collation_checkbox.value is False
    collation_checkbox.set_value(True).run()
    assert not at.exception

    next(t for t in at.text_input if t.label == "Client name").set_value("Amazon Business EMEA").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()

    save_button = next(b for b in at.button if "Save Client Profile" in b.label)
    save_button.click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Amazon Business EMEA", get_clients_dir())
    assert loaded.collation_enabled is True


def test_complex_account_checkbox_reveals_file_path_fields_and_saves(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    assert not any(t.label == "TAL file path" for t in at.text_input)

    complex_checkbox = next(c for c in at.checkbox if c.label == "This is a complex account")
    complex_checkbox.set_value(True).run()
    assert not at.exception

    next(t for t in at.text_input if t.label == "Client name").set_value("Dell APAC").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_input(key="complex_account_tal_path_input").set_value(str(tmp_path / "TAL.csv")).run()
    at.text_input(key="complex_account_specs_path_input").set_value(str(tmp_path / "specs.xlsx")).run()

    save_button = next(b for b in at.button if "Save Client Profile" in b.label)
    save_button.click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Dell APAC", get_clients_dir())
    assert loaded.complex_account.enabled is True
    assert loaded.complex_account.tal_path == str(tmp_path / "TAL.csv")
    assert loaded.complex_account.specifications_path == str(tmp_path / "specs.xlsx")


def test_box_tracker_checkbox_reveals_fields_and_saves(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    assert not any(t.label == "Local mirror workbook path" for t in at.text_input)

    box_tracker_checkbox = next(c for c in at.checkbox if c.label == "This client uses a Box Tracker")
    box_tracker_checkbox.set_value(True).run()
    assert not at.exception

    next(t for t in at.text_input if t.label == "Client name").set_value("IBM APAC").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_input(key="box_tracker_mirror_path_input").set_value(str(tmp_path / "mirror.xlsx")).run()
    at.text_area(key="box_tracker_cid_map_input").set_value("118741,Bob\n118742,CXO").run()

    save_button = next(b for b in at.button if "Save Client Profile" in b.label)
    save_button.click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("IBM APAC", get_clients_dir())
    assert loaded.box_tracker.enabled is True
    assert loaded.box_tracker.mirror_workbook_path == str(tmp_path / "mirror.xlsx")
    assert loaded.box_tracker.cid_campaign_map == {"118741": "Bob", "118742": "CXO"}


def test_convertr_checkbox_reveals_fields_and_saves(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    assert not any("Convertr enterprise subdomain" in t.label for t in at.text_input)

    convertr_checkbox = next(c for c in at.checkbox if c.label == "This client uploads to Convertr")
    convertr_checkbox.set_value(True).run()
    assert not at.exception

    next(t for t in at.text_input if t.label == "Client name").set_value("Amazon Business EMEA").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    next(t for t in at.text_input if "Convertr enterprise subdomain" in t.label).set_value("amazonbusiness").run()
    next(t for t in at.text_input if "Convertr Publisher ID" in t.label).set_value("11003").run()
    at.text_area(key="convertr_campaigns_input").set_value(
        "120022,44709\n120021,44709\n120028,44706,80").run()
    at.text_area(key="convertr_field_map_cols_input").set_value("Email\nFirst Name").run()
    at.text_area(key="convertr_field_map_targets_input").set_value("email\nfirstName").run()

    # No per-campaign API key inputs anymore -- the Publisher API uses the
    # one account login for every campaign, so only a single "Test
    # connection" button should appear per unique campaign_id.
    assert not any("campaign 44709" in t.label for t in at.text_input)
    assert sum(1 for b in at.button if "Test connection" in b.label) == 2
    at.text_input(key="convertr_account_username").set_value("me@x.com").run()
    at.text_input(key="convertr_account_password").set_value("hunter2").run()

    save_button = next(b for b in at.button if "Save Client Profile" in b.label)
    save_button.click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir, get_convertr_account_credentials
    from core.profile_store import load_profile

    loaded = load_profile("Amazon Business EMEA", get_clients_dir())
    assert loaded.convertr.enabled is True
    assert loaded.convertr.enterprise == "amazonbusiness"
    assert loaded.convertr.publisher_id == "11003"
    assert loaded.convertr.field_mapping == {"Email": "email", "First Name": "firstName"}
    campaigns_by_cid = {c.cid: c for c in loaded.convertr.campaigns}
    assert campaigns_by_cid["120022"].campaign_id == "44709"
    assert campaigns_by_cid["120028"].campaign_id == "44706"
    assert campaigns_by_cid["120028"].global_form_id == "80"

    # The account credentials must NOT be in the shared client profile
    # JSON, only in the local app_settings.json.
    with open(get_clients_dir() + "/Amazon Business EMEA.json", encoding="utf-8") as f:
        assert "hunter2" not in f.read()
    assert get_convertr_account_credentials("Amazon Business EMEA") == {
        "username": "me@x.com", "password": "hunter2"}


def test_convertr_leadfile_column_mapping_saves_independently_of_qa_field_mapping(tmp_path, monkeypatch):
    # This is what lets a client with no QA at all (e.g. Amazon) use
    # Convertr without ever visiting Run Check first -- the leadfile column
    # mapping lives right here on Convertr's own section.
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    next(c for c in at.checkbox if c.label == "This client uploads to Convertr").set_value(True).run()

    next(t for t in at.text_input if t.label == "Client name").set_value("Amazon Business EMEA").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    next(t for t in at.text_input if "Convertr enterprise subdomain" in t.label).set_value("amazonbusiness").run()
    next(t for t in at.text_input if "Convertr Publisher ID" in t.label).set_value("11003").run()
    at.text_input(key="convertr_lf_email").set_value("Email").run()
    at.text_input(key="convertr_lf_first").set_value("First Name").run()
    at.text_input(key="convertr_lf_last").set_value("Last Name").run()
    at.text_input(key="convertr_lf_company").set_value("Company").run()
    at.text_input(key="convertr_lf_cid").set_value("CID").run()

    next(b for b in at.button if "Save Client Profile" in b.label).click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Amazon Business EMEA", get_clients_dir())
    assert loaded.field_mapping is None  # no QA configured
    assert loaded.convertr.leadfile_field_mapping.email == "Email"
    assert loaded.convertr.leadfile_field_mapping.cid == "CID"


def test_enhancio_checkbox_reveals_fields_and_saves(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    assert not any(t.key == "enhancio_allocations_input" for t in at.text_area)

    enhancio_checkbox = next(c for c in at.checkbox if c.label == "This client uploads to Enhancio")
    enhancio_checkbox.set_value(True).run()
    assert not at.exception

    next(t for t in at.text_input if t.label == "Client name").set_value("Amazon Business EMEA").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_area(key="enhancio_allocations_input").set_value("120022,L-22256\n120028,L-22257").run()
    at.text_area(key="enhancio_field_map_cols_input").set_value("Email\nFirst Name").run()
    at.text_area(key="enhancio_field_map_targets_input").set_value("Email Address\nFirst Name").run()
    at.text_area(key="enhancio_fixed_values_input").set_value(
        "L-22256,Company Size,1M - 5M\nL-22256,Lead Source,Website").run()
    at.text_input(key="enhancio_lf_email").set_value("Email").run()
    at.text_input(key="enhancio_lf_cid").set_value("CID").run()

    save_button = next(b for b in at.button if "Save Client Profile" in b.label)
    save_button.click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Amazon Business EMEA", get_clients_dir())
    assert loaded.enhancio.enabled is True
    assert loaded.enhancio.field_mapping == {"Email": "Email Address", "First Name": "First Name"}
    allocations_by_cid = {a.cid: a for a in loaded.enhancio.allocations}
    assert allocations_by_cid["120022"].allocation_uid == "L-22256"
    assert allocations_by_cid["120028"].allocation_uid == "L-22257"
    assert loaded.enhancio.fixed_field_values == {
        "L-22256": {"Company Size": "1M - 5M", "Lead Source": "Website"}}
    assert loaded.enhancio.leadfile_field_mapping.email == "Email"
    assert loaded.enhancio.leadfile_field_mapping.cid == "CID"


def test_enhancio_field_mapping_handles_a_column_containing_commas_or_any_punctuation(tmp_path, monkeypatch):
    # A real leadfile column is often a verbatim survey/consent question
    # ("I'd like to receive news..., and I agree to...") that can contain
    # commas, arrows, or any other punctuation a delimiter-based format
    # might pick as "the separator." The paired-list design (two plain
    # lists, matched by line position) has no delimiter to confuse at
    # all -- each line is taken verbatim, in either box.
    monkeypatch.chdir(tmp_path)
    consent_text = "I'd like to receive news, and I agree to the -> policy, for more details."

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    next(c for c in at.checkbox if c.label == "This client uploads to Enhancio").set_value(True).run()
    next(t for t in at.text_input if t.label == "Client name").set_value("Comma Client").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_area(key="enhancio_field_map_cols_input").set_value(f"Email\n{consent_text}").run()
    at.text_area(key="enhancio_field_map_targets_input").set_value(f"Email Address\n{consent_text}").run()

    next(b for b in at.button if "Save Client Profile" in b.label).click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Comma Client", get_clients_dir())
    assert loaded.enhancio.field_mapping == {"Email": "Email Address", consent_text: consent_text}


def test_enhancio_field_mapping_round_trips(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    consent_text = "I'd like to receive news, and I agree, for more details, read the policy."

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    next(c for c in at.checkbox if c.label == "This client uploads to Enhancio").set_value(True).run()
    next(t for t in at.text_input if t.label == "Client name").set_value("Comma Client").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_area(key="enhancio_field_map_cols_input").set_value(consent_text).run()
    at.text_area(key="enhancio_field_map_targets_input").set_value(consent_text).run()
    next(b for b in at.button if "Save Client Profile" in b.label).click().run()

    at2 = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at2.run()
    next(r for r in at2.radio if r.label == "Mode").set_value("Edit existing client").run()
    next(s for s in at2.selectbox if s.label == "Client").set_value("Comma Client").run()

    cols_area = next(t for t in at2.text_area if t.key == "enhancio_field_map_cols_input")
    targets_area = next(t for t in at2.text_area if t.key == "enhancio_field_map_targets_input")
    assert cols_area.value == consent_text
    assert targets_area.value == consent_text


def test_field_mapping_mismatched_line_counts_shows_error_and_keeps_existing_mapping(tmp_path, monkeypatch):
    # A typo (an extra/missing line in just one box) must be caught clearly
    # rather than silently zipping the wrong lines together -- and must not
    # destroy whatever mapping was already saved.
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    next(c for c in at.checkbox if c.label == "This client uploads to Enhancio").set_value(True).run()
    next(t for t in at.text_input if t.label == "Client name").set_value("Mismatch Client").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_area(key="enhancio_field_map_cols_input").set_value("Email\nFirst Name").run()
    at.text_area(key="enhancio_field_map_targets_input").set_value("Email Address").run()

    assert any("don't have the same number of lines" in e.value for e in at.error)

    next(b for b in at.button if "Save Client Profile" in b.label).click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Mismatch Client", get_clients_dir())
    assert loaded.enhancio.field_mapping == {}  # nothing was saved previously, so still empty


def test_convertr_field_mapping_handles_a_column_containing_a_comma(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    next(c for c in at.checkbox if c.label == "This client uploads to Convertr").set_value(True).run()
    next(t for t in at.text_input if t.label == "Client name").set_value("Comma Client").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_area(key="convertr_field_map_cols_input").set_value(
        "I'd like to receive news, and I agree to the policy").run()
    at.text_area(key="convertr_field_map_targets_input").set_value("optIn").run()

    next(b for b in at.button if "Save Client Profile" in b.label).click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Comma Client", get_clients_dir())
    assert loaded.convertr.field_mapping == {
        "I'd like to receive news, and I agree to the policy": "optIn"}


def test_enhancio_fetch_allocations_button_shows_client_id_prompt_when_unset(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    next(c for c in at.checkbox if c.label == "This client uploads to Enhancio").set_value(True).run()

    assert any("No Enhancio Client ID configured" in c.value for c in at.caption)
    assert not any("Fetch allocations from Enhancio" in b.label for b in at.button)


def test_enhancio_test_connection_flags_a_mandatory_field_with_no_mapping(tmp_path, monkeypatch):
    # This is the actual safeguard against the recurring "Missing mandatory
    # field(s)" failure at upload time: Test connection now cross-checks
    # Describe Fields' mandatory list against the field mapping above it,
    # so a gap is caught here -- before ever sending a single lead -- not
    # discovered later from Enhancio's own rejected-lead count.
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    next(c for c in at.checkbox if c.label == "This client uploads to Enhancio").set_value(True).run()
    at.text_area(key="enhancio_allocations_input").set_value("120022,L-22256").run()
    at.text_area(key="enhancio_field_map_cols_input").set_value("Email").run()
    at.text_area(key="enhancio_field_map_targets_input").set_value("Email Address").run()

    from core.app_settings import save_enhancio_client_id
    save_enhancio_client_id("CID123")
    at.run()

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.describe_fields", return_value=[
             {"fieldLabel": "Email Address", "mandatory": "Y"},
             {"fieldLabel": "First Name", "mandatory": "Y"},
         ]):
        next(b for b in at.button if b.label == "Test connection — allocation L-22256").click().run()

    assert not at.exception
    assert any("First Name" in e.value and "NO mapping entry" in e.value for e in at.error)
    assert any("mapped from leadfile column \"Email\"" in s.value for s in at.success)


def test_enhancio_test_connection_shows_a_picklist_fields_allowed_values(tmp_path, monkeypatch):
    # A field constrained to a fixed picklist rejects anything outside it
    # as "Invalid field value" at upload time, with no indication of what
    # IS valid -- Test connection surfaces Describe Fields' own allowed
    # values list so a rejected-batch mismatch can be diagnosed here.
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()
    next(c for c in at.checkbox if c.label == "This client uploads to Enhancio").set_value(True).run()
    at.text_area(key="enhancio_allocations_input").set_value("120022,L-22256").run()
    at.text_area(key="enhancio_field_map_cols_input").set_value("Email").run()
    at.text_area(key="enhancio_field_map_targets_input").set_value("Email Address").run()

    from core.app_settings import save_enhancio_client_id
    save_enhancio_client_id("CID123")
    at.run()

    with patch("core.enhancio_client.get_access_token", return_value={"access_token": "tok"}), \
         patch("core.enhancio_client.describe_fields", return_value=[
             {"fieldLabel": "Email Address", "mandatory": "Y"},
             {"fieldLabel": "Industry", "mandatory": "Y", "fieldValues": ["All"]},
         ]):
        next(b for b in at.button if b.label == "Test connection — allocation L-22256").click().run()

    assert not at.exception
    assert any("Allowed values: All" in c.value for c in at.caption)


def test_box_tracker_lead_template_map_and_pacing_skip_fields_save(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    box_tracker_checkbox = next(c for c in at.checkbox if c.label == "This client uses a Box Tracker")
    box_tracker_checkbox.set_value(True).run()

    next(t for t in at.text_input if t.label == "Client name").set_value("IBM APAC").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_input(key="box_tracker_mirror_path_input").set_value(str(tmp_path / "mirror.xlsx")).run()
    at.text_area(key="box_tracker_lead_template_map_input").set_value(
        f"118741,{tmp_path / 'bob_template.xlsx'}").run()
    at.text_area(key="box_tracker_pacing_skipped_input").set_value("CXO").run()

    save_button = next(b for b in at.button if "Save Client Profile" in b.label)
    save_button.click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("IBM APAC", get_clients_dir())
    assert loaded.box_tracker.cid_lead_template_path == {"118741": str(tmp_path / "bob_template.xlsx")}
    assert loaded.box_tracker.pacing_skipped_campaigns == ["CXO"]


def test_saving_integrate_section_persists_sid_and_field_mapping(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    next(c for c in at.checkbox if c.label == "This client uploads to Integrate").set_value(True).run()
    next(t for t in at.text_input if t.label == "Client name").set_value("Everpure EMEA").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_input(key="integrate_sid_input").set_value("d8a9deeb-7bb2-4832-b3d9-1df8a9fe5cab").run()
    at.text_area(key="integrate_field_map_cols_input").set_value("Email\nFirst Name").run()
    at.text_area(key="integrate_field_map_targets_input").set_value("email\nfirst_name").run()
    at.text_area(key="integrate_fixed_values_input").set_value("country,UK").run()

    next(b for b in at.button if "Save Client Profile" in b.label).click().run()
    assert not at.exception

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Everpure EMEA", get_clients_dir())
    assert loaded.integrate.enabled is True
    assert loaded.integrate.sid == "d8a9deeb-7bb2-4832-b3d9-1df8a9fe5cab"
    assert loaded.integrate.field_mapping == {"Email": "email", "First Name": "first_name"}
    assert loaded.integrate.fixed_field_values == {"country": "UK"}


def test_invalid_integrate_target_attribute_warns_but_still_saves(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    at = AppTest.from_file(_PAGE_PATH, default_timeout=15)
    at.run()

    next(c for c in at.checkbox if c.label == "This client uploads to Integrate").set_value(True).run()
    next(t for t in at.text_input if t.label == "Client name").set_value("Everpure EMEA").run()
    at.text_input(key="accumulated_path_input").set_value(str(tmp_path / "accumulated.xlsx")).run()
    at.text_input(key="integrate_sid_input").set_value("d8a9deeb-7bb2-4832-b3d9-1df8a9fe5cab").run()
    at.text_area(key="integrate_field_map_cols_input").set_value("Email").run()
    # "emial" is a typo -- not one of Integrate's 17 real attribute names.
    at.text_area(key="integrate_field_map_targets_input").set_value("emial").run()

    next(b for b in at.button if "Save Client Profile" in b.label).click().run()
    assert not at.exception
    assert any("emial" in w.value for w in at.warning)

    from core.app_settings import get_clients_dir
    from core.profile_store import load_profile

    loaded = load_profile("Everpure EMEA", get_clients_dir())
    assert loaded.integrate.field_mapping == {"Email": "emial"}
