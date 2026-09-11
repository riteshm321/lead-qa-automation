import datetime

import openpyxl
import pandas as pd

from core.box_tracker import (
    current_week_label, read_pacing_diffs, pick_leads_for_approval, sent_for_approval_label,
    append_mirror_rows, set_pacing_delivered, add_lead_template_columns,
    cleared_for_upload_label, uploaded_accepted_label, uploaded_rejected_label,
    read_lead_template_constants, parse_amal_id, project_code_for_cid, campaign_type_for_cid,
    uploaded_to_approval_sheet_label,
)


def test_current_week_label_finds_the_most_recent_monday():
    # 2026-09-09 is a Wednesday; that week's Monday is 2026-09-07.
    assert current_week_label(datetime.date(2026, 9, 9)) == "Week of 7"


def test_current_week_label_when_today_is_the_monday():
    assert current_week_label(datetime.date(2026, 9, 7)) == "Week of 7"


def _make_pacing_workbook(path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pacing"
    # Summary block starting at row 13, matching the real Box tracker's
    # layout: row 13 = campaign name headers, row 14 = Pending,
    # row 15 = Delivered, row 16 = Diff.
    ws["B13"] = "Bob"
    ws["C13"] = "wxO (AI Pod) IN"
    ws["D13"] = "wxO (AI Pod) AU"
    ws["A14"], ws["B14"], ws["C14"], ws["D14"] = "Pending", 333, 179, 74
    ws["A15"], ws["B15"], ws["C15"], ws["D15"] = "Delivered", 320, 166, 63
    ws["A16"], ws["B16"], ws["C16"], ws["D16"] = "Diff", 13, 13, 11
    wb.save(path)


def test_read_pacing_diffs_reads_every_campaign_column(tmp_path):
    path = str(tmp_path / "mirror.xlsx")
    _make_pacing_workbook(path)

    diffs = read_pacing_diffs(path)

    assert diffs == {"Bob": 13, "wxO (AI Pod) IN": 13, "wxO (AI Pod) AU": 11}


def test_read_pacing_diffs_is_dynamic_to_however_many_columns_exist(tmp_path):
    # Adding a 4th campaign column must be picked up with no code change --
    # this is a Global Constraint of the whole feature, not just a nicety.
    path = str(tmp_path / "mirror.xlsx")
    _make_pacing_workbook(path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Pacing"]
    ws["E13"] = "CXO"
    ws["E14"], ws["E15"], ws["E16"] = 50, 40, 10
    wb.save(path)

    diffs = read_pacing_diffs(path)

    assert diffs["CXO"] == 10
    assert len(diffs) == 4


def test_sent_for_approval_label_format():
    assert sent_for_approval_label(datetime.date(2026, 9, 7)) == "Sent for Approval - 07-Sep"


def test_status_label_formats_for_the_rest_of_the_lifecycle():
    d = datetime.date(2026, 9, 7)
    assert cleared_for_upload_label(d) == "Cleared for Upload - 07-Sep"
    assert uploaded_accepted_label(d) == "Accepted - Uploaded - 07-Sep"
    assert uploaded_rejected_label(d) == "Rejected - Refunded - 07-Sep"
    assert uploaded_to_approval_sheet_label(d) == "Uploaded to Approval Sheet - 07-Sep"


def _accumulated_df(rows):
    return pd.DataFrame(rows)


def test_pick_leads_for_approval_picks_diff_plus_five_per_campaign():
    accumulated = _accumulated_df([
        {"CID": "118741", "Status": ""} for _ in range(20)
    ] + [
        {"CID": "118743", "Status": ""} for _ in range(20)
    ])
    cid_campaign_map = {"118741": "Bob", "118743": "wxO (AI Pod) IN"}
    diffs = {"Bob": 13, "wxO (AI Pod) IN": 13}

    picked, shortfall = pick_leads_for_approval(
        accumulated, "CID", "Status", cid_campaign_map, diffs, buffer=5)

    assert len(picked[picked["CID"] == "118741"]) == 18  # 13 + 5
    assert len(picked[picked["CID"] == "118743"]) == 18
    assert shortfall == {}


def test_pick_leads_for_approval_picks_nothing_when_already_ahead_of_pace():
    # A negative diff (Delivered already exceeds Pending) plus the buffer
    # can still be negative -- e.g. diff=-7, buffer=5 -> -2. That must
    # clamp to 0 picks, never pandas' .head(-2) "all but the last 2" trap.
    accumulated = _accumulated_df([{"CID": "118741", "Status": ""} for _ in range(5)])

    picked, shortfall = pick_leads_for_approval(
        accumulated, "CID", "Status", {"118741": "Bob"}, {"Bob": -7}, buffer=5)

    assert picked.empty
    assert shortfall == {}


def test_pick_leads_for_approval_ignores_leads_with_a_non_blank_status():
    accumulated = _accumulated_df([
        {"CID": "118741", "Status": "Sent for Approval - 01-Sep"},
        {"CID": "118741", "Status": ""},
    ])
    picked, _ = pick_leads_for_approval(
        accumulated, "CID", "Status", {"118741": "Bob"}, {"Bob": 0}, buffer=5)

    assert len(picked) == 1


def test_pick_leads_for_approval_reports_shortfall_by_cid():
    accumulated = _accumulated_df([{"CID": "118741", "Status": ""} for _ in range(3)])
    diffs = {"Bob": 13}  # needs 13 + 5 = 18, only 3 available

    picked, shortfall = pick_leads_for_approval(
        accumulated, "CID", "Status", {"118741": "Bob"}, diffs, buffer=5)

    assert len(picked) == 3
    assert shortfall == {"118741": 15}  # needed 18, short by 15


def test_pick_leads_for_approval_ignores_cids_with_no_campaign_mapping():
    # A CID not in cid_campaign_map (999999 here) has no way to know its
    # target count -- must be left alone, never picked, never reported as
    # a shortfall. The mapped CID (118741) still correctly reports its own
    # shortfall since Accumulated has no matching leads for it at all.
    accumulated = _accumulated_df([{"CID": "999999", "Status": ""}])
    picked, shortfall = pick_leads_for_approval(
        accumulated, "CID", "Status", {"118741": "Bob"}, {"Bob": 13}, buffer=5)

    assert picked.empty
    assert shortfall == {"118741": 18}
    assert "999999" not in shortfall


def test_pick_leads_for_approval_takes_all_available_for_uncapped_campaigns():
    # A newly-live segment (e.g. CXO) with no established Pacing history
    # yet -- forcing it through the diff+buffer cap (0, or no entry at
    # all) isn't right; it should take every available blank-Status lead
    # instead, and never appear in shortfall (there's no target to fall
    # short of).
    accumulated = _accumulated_df([{"CID": "118742", "Status": ""} for _ in range(7)])

    picked, shortfall = pick_leads_for_approval(
        accumulated, "CID", "Status", {"118742": "CXO"}, diffs={}, buffer=5,
        uncapped_campaigns={"CXO"},
    )

    assert len(picked) == 7
    assert shortfall == {}


def test_pick_leads_for_approval_uncapped_still_respects_blank_status():
    accumulated = _accumulated_df([
        {"CID": "118742", "Status": "Sent for Approval - 01-Sep"},
        {"CID": "118742", "Status": ""},
    ])

    picked, _ = pick_leads_for_approval(
        accumulated, "CID", "Status", {"118742": "CXO"}, diffs={}, buffer=5,
        uncapped_campaigns={"CXO"},
    )

    assert len(picked) == 1


def test_append_mirror_rows_matches_by_header_and_appends_after_last_row(tmp_path):
    path = str(tmp_path / "mirror.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Approval Sheet"
    ws.append(["Company Name", "Market", "Date", "Segment"])
    ws.append(["Existing Co", "IN", "01-Sep", "SelectT"])
    wb.save(path)

    append_mirror_rows(path, "Approval Sheet", [
        {"Company Name": "New Co", "Market": "AU", "Date": "07-Sep", "Segment": "Named"},
    ])

    wb2 = openpyxl.load_workbook(path)
    ws2 = wb2["Approval Sheet"]
    assert ws2.cell(row=3, column=1).value == "New Co"
    assert ws2.cell(row=3, column=2).value == "AU"
    assert ws2.cell(row=3, column=4).value == "Named"


def test_append_mirror_rows_leaves_unmatched_dict_keys_out(tmp_path):
    path = str(tmp_path / "mirror.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["A", "B"])
    wb.save(path)

    append_mirror_rows(path, "Sheet1", [{"A": "value", "NotAColumn": "ignored"}])

    wb2 = openpyxl.load_workbook(path)
    ws2 = wb2["Sheet1"]
    assert ws2.cell(row=2, column=1).value == "value"
    assert ws2.cell(row=2, column=2).value is None


def test_set_pacing_delivered_writes_the_current_weeks_column(tmp_path):
    path = str(tmp_path / "mirror.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pacing"
    ws.append(["Funding Source", "Publisher", "Country", "Segment", "Campaign", "Week of 7", ""])
    ws.append([None, None, None, None, None, "P", "D"])
    ws.append(["Cash", "Madison Logic", "IN", "Select-T", "Bob", 18, 0])
    wb.save(path)

    set_pacing_delivered(path, "Bob", 18, week_label="Week of 7")

    wb2 = openpyxl.load_workbook(path)
    ws2 = wb2["Pacing"]
    assert ws2.cell(row=3, column=7).value == 18  # the "D" sub-column under "Week of 7"


def test_set_pacing_delivered_overwrites_not_adds(tmp_path):
    path = str(tmp_path / "mirror.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pacing"
    ws.append(["Funding Source", "Publisher", "Country", "Segment", "Campaign", "Week of 7", ""])
    ws.append([None, None, None, None, None, "P", "D"])
    ws.append(["Cash", "Madison Logic", "IN", "Select-T", "Bob", 18, 18])
    wb.save(path)

    set_pacing_delivered(path, "Bob", 15, week_label="Week of 7")

    wb2 = openpyxl.load_workbook(path)
    assert wb2["Pacing"].cell(row=3, column=7).value == 15


def test_set_pacing_delivered_finds_week_label_a_row_lower_than_the_static_columns(tmp_path):
    # The real Box file's static columns (Funding Source..Campaign) are
    # merged across rows 1-3 with the header text in row 1, but the week
    # labels live in row 2 (row 1 there holds a month label instead) and
    # the P/D sub-headers in row 3 -- one row lower than the simplest case.
    path = str(tmp_path / "mirror.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pacing"
    ws.append(["Funding Source", "Publisher", "Country", "Segment", "Campaign", "Tactic", "Live Date", "Total Planned", "July"])
    ws.append([None, None, None, None, None, None, None, None, "Week of 7"])
    ws.append([None, None, None, None, None, None, None, None, "P", "D"])
    ws.append(["Cash", "Madison Logic", "IN", "Select-T", "Bob", "2 Touch", None, 361, 0, 0])
    wb.save(path)

    set_pacing_delivered(path, "Bob", 18, week_label="Week of 7")

    wb2 = openpyxl.load_workbook(path)
    assert wb2["Pacing"].cell(row=4, column=10).value == 18  # the "D" sub-column under "Week of 7"


def test_set_pacing_delivered_disambiguates_same_campaign_name_by_country(tmp_path):
    # "wxO (AI Pod)" appears once per country (IN and AU), with the same
    # Campaign text -- only the Country column tells the rows apart. A
    # campaign key with a trailing " IN"/" AU" must match the row whose
    # Country column agrees, not just the first row with that text.
    path = str(tmp_path / "mirror.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pacing"
    ws.append(["Funding Source", "Publisher", "Country", "Segment", "Campaign", "Week of 7", ""])
    ws.append([None, None, None, None, None, "P", "D"])
    ws.append(["Cash", "Madison Logic", "IN", "Select-T", "wxO (AI Pod)", 15, 0])
    ws.append(["Cash", "Madison Logic", "AU", "Select-T", "wxO (AI Pod)", 8, 0])
    wb.save(path)

    set_pacing_delivered(path, "wxO (AI Pod) AU", 8, week_label="Week of 7")

    wb2 = openpyxl.load_workbook(path)
    ws2 = wb2["Pacing"]
    assert ws2.cell(row=3, column=7).value == 0   # IN row untouched
    assert ws2.cell(row=4, column=7).value == 8   # AU row updated


def test_add_lead_template_columns_uses_the_fixed_cid_map():
    leads_df = pd.DataFrame([
        {"CID": "118741", "LOB": "Software"},   # Bob
        {"CID": "120129", "LOB": "Finance"},    # AU CXO
        {"CID": "118743", "LOB": "Ops"},        # IN WXO
        {"CID": "118745", "LOB": "Ops"},        # AU WXO
        {"CID": "120130", "LOB": "Finance"},    # IN CXO
    ])

    result = add_lead_template_columns(leads_df, "CID")

    assert list(result["micro_audience"]) == ["Platform_SWE", "All", "AI Leaders", "AI Leaders", "All_CXO"]
    assert list(result["Industry"]) == ["All", "All", "All", "All", "All"]


def test_add_lead_template_columns_copies_lob_for_lob_sourced_cids():
    leads_df = pd.DataFrame([
        {"CID": "119750", "LOB": "Cloud Infra"},  # IN LOB (shares IN WXO template)
        {"CID": "119751", "LOB": "Data & AI"},    # AU LOB (shares AU WXO template)
    ])

    result = add_lead_template_columns(leads_df, "CID")

    assert list(result["micro_audience"]) == ["Cloud Infra", "Data & AI"]
    assert list(result["Industry"]) == ["All", "All"]


def test_add_lead_template_columns_copies_the_leadfiles_own_micro_audience_for_digisov():
    # IN DigiSov (120131) is neither a fixed value nor sourced from LOB --
    # the leadfile carries its own micro_audience column directly, which
    # passes straight through under the same header name.
    leads_df = pd.DataFrame([{"CID": "120131", "micro_audience": "Security Leaders"}])

    result = add_lead_template_columns(leads_df, "CID")

    assert result.loc[0, "micro_audience"] == "Security Leaders"
    assert result.loc[0, "Industry"] == "All"


def test_add_lead_template_columns_blank_for_unmapped_cid():
    leads_df = pd.DataFrame([{"CID": "999999", "LOB": "Anything"}])

    result = add_lead_template_columns(leads_df, "CID")

    assert result.loc[0, "micro_audience"] == ""
    assert result.loc[0, "Industry"] == "All"


def test_add_lead_template_columns_injects_template_constants_and_passthroughs():
    leads_df = pd.DataFrame([
        {"CID": "118741", "LOB": "Software", "Asset Title": "Omdia Universe", "Country": "IN",
         "Company Size": "1000-5000"},
    ])
    template_constants = {
        "AID": "L-22SD7", "NC_EMAIL_DETAIL": "UC", "NC_TELE_DETAIL": "UC", "campaign_code": "PVLAP",
    }

    result = add_lead_template_columns(
        leads_df, "CID", template_constants=template_constants,
        now=datetime.datetime(2026, 9, 7, 14, 30, 5),
    )

    assert result.loc[0, "AID"] == "L-22SD7"
    assert result.loc[0, "NC_EMAIL_DETAIL"] == "UC"
    assert result.loc[0, "NC_TELE_DETAIL"] == "UC"
    assert result.loc[0, "campaign_code"] == "PVLAP"
    assert result.loc[0, "asset_title"] == "Omdia Universe"
    assert result.loc[0, "country"] == "IN"
    assert result.loc[0, "Q_COMPS"] == "1000-5000"
    assert result.loc[0, "user_transaction_date"] == "2026-09-07 14:30:05"


def test_add_lead_template_columns_parses_the_leadfiles_own_timestamp_column():
    # The leadfile's Timestamp column is plain text in whatever format the
    # source system wrote it in -- it must be parsed into a real date and
    # reformatted to the Lead Template's exact "YYYY-MM-DD HH:MM:SS"
    # string, not just passed through as-is.
    leads_df = pd.DataFrame([
        {"CID": "118741", "Timestamp": "07/21/2026 3:45:00 PM"},
        {"CID": "118741", "Timestamp": "2026-08-02 09:05:00"},
    ])

    result = add_lead_template_columns(
        leads_df, "CID", template_constants={"AID": "L-22SD7"},
        now=datetime.datetime(2026, 9, 7, 14, 30, 5),
    )

    assert result.loc[0, "user_transaction_date"] == "2026-07-21 15:45:00"
    assert result.loc[1, "user_transaction_date"] == "2026-08-02 09:05:00"


def test_add_lead_template_columns_falls_back_to_now_when_timestamp_missing_or_unparseable():
    leads_df = pd.DataFrame([
        {"CID": "118741", "Timestamp": None},
        {"CID": "118741", "Timestamp": "not a date"},
    ])

    result = add_lead_template_columns(
        leads_df, "CID", template_constants={"AID": "L-22SD7"},
        now=datetime.datetime(2026, 9, 7, 14, 30, 5),
    )

    assert result.loc[0, "user_transaction_date"] == "2026-09-07 14:30:05"
    assert result.loc[1, "user_transaction_date"] == "2026-09-07 14:30:05"


def test_add_lead_template_columns_without_template_constants_adds_nothing_extra():
    # Backward-compatible default -- callers that don't pass
    # template_constants (or asset_title/country columns aren't present)
    # still just get micro_audience/Industry, same as before.
    leads_df = pd.DataFrame([{"CID": "118741"}])

    result = add_lead_template_columns(leads_df, "CID")

    assert "AID" not in result.columns
    assert "asset_title" not in result.columns
    assert "user_transaction_date" not in result.columns


def test_read_lead_template_constants_reads_the_first_row_with_a_non_blank_aid(tmp_path):
    path = str(tmp_path / "template.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LEAD_TEMPLATE"
    ws.append(["AID", "NC_EMAIL_DETAIL", "NC_TELE_DETAIL", "user_transaction_date", "campaign_code", "email"])
    ws.append(["L-22SD7", "UC", "UC", None, "PVLAP", "lead1@x.com"])
    ws.append(["L-22SD7", "UC", "UC", None, "PVLAP", "lead2@x.com"])
    wb.save(path)

    constants = read_lead_template_constants(path, "LEAD_TEMPLATE")

    assert constants == {"AID": "L-22SD7", "NC_EMAIL_DETAIL": "UC", "NC_TELE_DETAIL": "UC", "campaign_code": "PVLAP"}


def test_read_lead_template_constants_uses_a_template_only_row_with_no_real_lead_yet(tmp_path):
    path = str(tmp_path / "template.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LEAD_TEMPLATE"
    ws.append(["AID", "NC_EMAIL_DETAIL", "NC_TELE_DETAIL", "campaign_code", "email"])
    ws.append(["L-22UMP", "UC", "UC", "CXOAP", None])  # template row, no real lead
    ws.append([None, None, None, None, None])  # leftover dropdown-list debris row
    wb.save(path)

    constants = read_lead_template_constants(path, "LEAD_TEMPLATE")

    assert constants == {"AID": "L-22UMP", "NC_EMAIL_DETAIL": "UC", "NC_TELE_DETAIL": "UC", "campaign_code": "CXOAP"}


def test_read_lead_template_constants_returns_empty_when_no_row_has_an_aid(tmp_path):
    path = str(tmp_path / "template.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LEAD_TEMPLATE"
    ws.append(["AID", "NC_EMAIL_DETAIL", "NC_TELE_DETAIL", "campaign_code", "email"])
    wb.save(path)

    assert read_lead_template_constants(path, "LEAD_TEMPLATE") == {}


def test_parse_amal_id_takes_the_later_value_when_two_are_comma_separated():
    assert parse_amal_id("6a07821488e4d402cbbdece8, 6a0782b288e4d402cbbdece9") == "6a0782b288e4d402cbbdece9"


def test_parse_amal_id_passes_through_a_single_value_unchanged():
    assert parse_amal_id("6a07821488e4d402cbbdece8") == "6a07821488e4d402cbbdece8"


def test_parse_amal_id_blank_for_none_or_empty():
    assert parse_amal_id(None) == ""
    assert parse_amal_id("") == ""


def test_project_code_for_cid_uses_the_leadfile_value_by_default():
    assert project_code_for_cid("118743", "PAIAP") == "PAIAP"


def test_project_code_for_cid_overrides_newly_live_segments():
    # Newly live segments' Project Code is always the fixed value,
    # regardless of what (if anything) the leadfile carries.
    assert project_code_for_cid("120129", "") == "CXOAP"   # AU CXO
    assert project_code_for_cid("120130", "") == "CXOAP"   # IN CXO
    assert project_code_for_cid("120131", "") == "SNCAP"   # IN DigiSov
    assert project_code_for_cid("118743", "PAIAP") == "PAIAP"  # not overridden -- leadfile value passes through


def test_campaign_type_for_cid_known_campaigns():
    assert campaign_type_for_cid("118741") == "2T"  # Bob
    assert campaign_type_for_cid("118743") == "2T"  # IN WXO
    assert campaign_type_for_cid("118745") == "2T"  # AU WXO
    assert campaign_type_for_cid("120129") == "1T"  # AU CXO
    assert campaign_type_for_cid("120130") == "1T"  # IN CXO
    assert campaign_type_for_cid("120131") == "1T"  # IN DigiSov


def test_campaign_type_for_cid_unknown_cid_is_blank():
    assert campaign_type_for_cid("999999") == ""
