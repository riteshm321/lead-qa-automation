"""Synthetic PPRA decks built with python-pptx - shaped like the platform's
raw Program Performance Reports (titles in text boxes named "Title N", \\xa0 and
en dashes in titles, header-only tables for empty slides) without any real
client data."""
import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches, Pt

EN_DASH = chr(0x2013)
NBSP = chr(0xA0)
VT = chr(0x0B)

SLIDE_W, SLIDE_H = 12192000, 6858000
_BLANK = 6
_TITLE_SLIDE = 0
_SECTION = 2


def new_prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    # The platform's layout names; detection keys off them.
    prs.slide_layouts[_SECTION]._element.cSld.set("name", "Section Breaker 2")
    return prs


def to_bytes(prs) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def blank(prs, title=None, title_name="Title 2"):
    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK])
    if title is not None:
        textbox(slide, title_name, [title], Inches(0.3), Inches(0.2), Inches(12), Inches(0.9), size=28)
    return slide


def textbox(slide, name, paragraphs, left=Inches(1), top=Inches(1), width=Inches(4), height=Inches(1), size=16):
    """paragraphs: list of str or list of (text, bold) run tuples."""
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.name = name
    tf = shape.text_frame
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(para, str) and VT in para:
            p.text = para  # a vertical tab becomes <a:br/>, as in the platform's decks
            for r in p.runs:
                r.font.size = Pt(size)
            continue
        runs = [(para, None)] if isinstance(para, str) else para
        for text, bold in runs:
            if text == "":
                continue
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            if bold is not None:
                r.font.bold = bold
    return shape


def table(slide, rows, name="Table 1", left=Inches(0.5), top=Inches(1.2), width=Inches(10), height=None,
          col_widths=None):
    height = height or Inches(0.3) * len(rows)
    frame = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, height)
    frame.name = name
    tbl = frame.table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = value
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    if col_widths:
        for column, w in zip(tbl.columns, col_widths):
            column.width = Emu(w)
    return frame


def chart(slide, chart_type, categories, series, name="Chart 1", title=None, labels=False,
          left=Inches(1), top=Inches(1.3), width=Inches(6), height=Inches(3)):
    data = CategoryChartData()
    data.categories = categories
    for series_name, values in series.items():
        data.add_series(series_name, values)
    frame = slide.shapes.add_chart(chart_type, left, top, width, height, data)
    frame.name = name
    if title:
        frame.chart.has_title = True
        frame.chart.chart_title.text_frame.text = title
    if labels:
        plot = frame.chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(9)
        plot.data_labels.font.bold = False
    return frame


# ---------------------------------------------------------------- slide builders

def title_slide(prs, flight_lines):
    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK])
    textbox(slide, "object 9", ["Client X"], top=Inches(3))
    textbox(slide, "object 9", ["Flight Dates: " + VT + flight_lines[0], *flight_lines[1:]], top=Inches(5))
    return slide


def section(prs, text):
    slide = prs.slides.add_slide(prs.slide_layouts[_SECTION])
    textbox(slide, "Text Placeholder 1", text.split("\n"), top=Inches(4))
    return slide


def pacing(prs, rows, title="Campaign Overview and Pacing", col_widths=None):
    slide = blank(prs, title, title_name="Title 5")
    header = ["Campaign Name", "Budget Goal", "Budget Delivered", "Budget Remaining", "Units Goal",
              "Units Delivered", "Units Remaining", "% Delivered"]
    widths = col_widths or [4126000, 909000, 1050000, 1050000, 1050000, 1050000, 1050000, 1145000]
    return table(slide, [header] + rows, name="Table 5", left=Emu(381000), width=Emu(sum(widths)),
                 col_widths=widths)


def mlp_box(slide, program="139823"):
    return textbox(slide, "TextBox 14", [[("Click to view on ML Platform ", None),
                                          ("https://platform.madisonlogic.com/platform", None),
                                          (f"/3836/programs/{program} ", None)]],
                   top=Inches(6.3), width=Inches(7.4), height=Inches(0.27), size=10)


def audience_reach(prs, pct="26% ", title="Audience Reach", engagement=((550, 300, 217), (79, 60, 48)),
                   one_paragraph=False):
    """engagement: (All Accounts, Trending) values for Targeted / Reached /
    Engaged. one_paragraph puts the stat run in the sentence's paragraph, as
    the platform's real decks do."""
    slide = blank(prs, title + NBSP)
    chart(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, ["Targeted", "Reached", "Engaged"],
          {"'All Accounts'": engagement[0], "'Trending'": engagement[1]}, name="Chart 3",
          title="Account Engagement Summary")
    sentence = [("of leads delivered from accounts actively engaging with your ", None), ("intent topics", None)]
    body = [[(pct, True), *sentence]] if one_paragraph else [[(pct, True)], sentence]
    textbox(slide, "Rectangle 13", ["Key Takeaways:", *body],
            left=Inches(9.4), top=Inches(1), width=Inches(3.65), height=Inches(2.7))
    mlp_box(slide)
    return slide


def audience_insights(prs, topics, pct_runs=(("0% ", True), ("of accounts are trending on ", None),
                                             ("5+ ", True), ("intent topics", None))):
    slide = blank(prs, "Audience Insights")
    values = list(range(len(topics) * 10, 0, -10))
    chart(slide, XL_CHART_TYPE.BAR_CLUSTERED, topics, {"DOMAIN": values}, name="Chart 10",
          title="Top Intent Topics in your Campaign")
    chart(slide, XL_CHART_TYPE.BAR_CLUSTERED, ["Ai", "Cloud"], {"Acconts": (5, 4)}, name="Chart 2",
          title="Top Intent Topics Targeted Accounts are Surging on", top=Inches(4))
    textbox(slide, "Rectangle 1", ["Key Takeaways:", list(pct_runs), "",
                                   "[Topics trending across your audience that you are not targeting]", ""],
            left=Inches(9.4), top=Inches(1), width=Inches(3.65), height=Inches(2.7))
    mlp_box(slide)
    return slide


def content_insights(prs, assets=(("The Playbook", "217", "320", "100%"),)):
    slide = blank(prs, "Content Insights")
    table(slide, [["Asset Preview", "Asset Name", "# Accounts", "# of Leads", "% of Leads"]]
          + [["", *a] for a in assets], name="Table 14", width=Inches(9))
    textbox(slide, "Rectangle 14", ["Key Takeaways:", "",
                                    "Assets that generated the strongest engagement are closely aligned with: ",
                                    "X", "Y", "Z "],
            left=Inches(9.4), top=Inches(1), width=Inches(3.65), height=Inches(2.7))
    return slide


def industry_insights(prs, job_titles=(("Chief Operations Officer", "83", "26%"),
                                       ("Chief Technology Officer", "55", "17%"),
                                       ("Chief Information Officer", "28", "9%"),
                                       ("Founder", "28", "9%"))):
    slide = blank(prs, "Industry Insights")
    table(slide, [["Job Title", "# of Leads", "Percent"], *[list(j) for j in job_titles]], name="Table 7",
          left=Inches(0.3), width=Inches(3.8))
    table(slide, [["Industry", "# of Leads", "Percent"], ["Software", "100", "50%"]], name="Table 7",
          left=Inches(4.5), width=Inches(3.8))
    table(slide, [["Company Size", "# of Leads", "Percent"], ["Large", "100", "50%"]], name="Table 7",
          left=Inches(8.7), width=Inches(3.8))
    return slide


def country_insights(prs, countries=(("United States", "311", "97%"), ("Canada", "9", "3%"))):
    slide = blank(prs, "Country Insights")
    if countries:
        chart(slide, XL_CHART_TYPE.PIE, [c[0] for c in countries],
              {"Percents": [int(c[2].rstrip("%")) / 100 for c in countries]}, name="Chart 2", labels=True,
              left=Inches(4.9))
    else:
        chart(slide, XL_CHART_TYPE.PIE, ["1st Qtr", "2nd Qtr", "3rd Qtr", "4th Qtr"],
              {"Column1": (None, None, None, None)}, name="Chart 2", left=Inches(4.9))
    table(slide, [["Country", "# of Leads", "Percent"], *[list(c) for c in countries]], name="Table 7",
          left=Inches(0.27), width=Inches(4.5))
    textbox(slide, "Rectangle 1", ["Key Takeaways:", "", "To be filled by CSM team: "],
            left=Inches(9.4), top=Inches(1.2), width=Inches(3.65), height=Inches(2.7))
    return slide


def top_accounts_cs(prs, rows=()):
    slide = blank(prs, f"Top Accounts {EN_DASH} Content Syndication")
    textbox(slide, "TextBox 8", ["*CSM team to add in logos of top accounts*"], top=Inches(6.4))
    table(slide, [["Account Name", "Account Domain", "Leads", "Top Trending Topic (Over Last 7 Days)",
                   "# Trending Topics (Over Last 12 Weeks)"], *[list(r) for r in rows]], name="Table 4",
          left=Inches(0.4), width=Inches(12.0))
    return slide


def top_accounts_display(prs, rows, title=f"Top Accounts {EN_DASH} Display and CTV"):
    slide = blank(prs, title)
    table(slide, [["Account Name", "Account Domain", "Display Impressions", "CTV Impressions", "Site Visits",
                   "Clicks", "# Trending Topics (Over Last 12 Weeks)"], *[list(r) for r in rows]],
          name="Table 2", left=Inches(0.3), width=Inches(12.7))
    return slide


def custom_question(prs):
    slide = blank(prs, f"Custom Question {EN_DASH} Example")
    chart(slide, XL_CHART_TYPE.PIE, ["Fully ready", "Not ready"], {"Column1": (5, 3)}, name="Chart 5")
    return slide


def creative_sets(prs, prefix):
    slide = blank(prs, f"{prefix}{EN_DASH} Display Data, Creative Sets", title_name="Title 4")
    textbox(slide, "TextBox 7", ["This slide will not be populated -" + NBSP + "CSMs to pull this report"],
            top=Inches(5))
    return slide


HALO_EXPLAINER = ("Average number of website visits per account who engaged multi-channel, compared to "
                  "single-channel")


def halo(prs, categories=("Single-channel",), site_visits=(0,), accounts=None, value="X"):
    slide = blank(prs, "Multi-Channel Engagement (Halo Effect)" + NBSP)
    chart(slide, XL_CHART_TYPE.BAR_CLUSTERED, list(categories),
          {"ACCOUNTS": list(accounts or [173] * len(categories)), "Site Visits": list(site_visits)},
          name="Chart 15")
    textbox(slide, "Rectangle 13", [[("Key Takeaways:", True)], [(value, True)], [(HALO_EXPLAINER, None)], ""],
            left=Inches(9.4), top=Inches(1), width=Inches(3.65), height=Inches(2.7))
    return slide


def ctv_performance(prs):
    slide = blank(prs, "Connected TV Performance")
    table(slide, [["Asset Name", "CTV Impressions", "Accounts"], ["spot-30", "12,000", "5"]], name="Table 3")
    return slide


def layout_footer(prs, top=Inches(6.6), layout_index=_BLANK):
    """A confidential-line text box on the layout (inherited by every slide
    that uses it, but not one of the slide's own shapes)."""
    from pptx.oxml.shapes.autoshape import CT_Shape
    layout = prs.slide_layouts[layout_index]
    sp = CT_Shape.new_textbox_sp(900, "TextBox 900", Inches(0.5), top, Inches(5), Inches(0.27))
    layout.shapes._spTree.append(sp)
    return top


ASSET_NAME = ("417344473_Client Logic_BROAD_JOB_NAT_Awareness-Direct-Standard-ProBuyer.CTV ProBuyer Fixed "
              "Income_1x1_OLV_V_NA (ProBuyer_CTV) v{}")


def ctv_data_slide(prs, n_assets=3, n_accounts=3, assets_top=Inches(3.44), row_h=Inches(0.35)):
    """A CTV Advertising Data slide: KPI tiles, then 'Top Assets By
    Impression' (long names that wrap to two lines, stored one line high) and
    'Top Accounts By Impression' laid out for the stored heights."""
    slide = blank(prs, f"Client X Combined View{EN_DASH} CTV Advertising Data" + VT)
    textbox(slide, "TextBox 13", [[("Campaign Progress: Impressions", True)], [("96.7%", False)]],
            left=Inches(0.65), top=Inches(1.14), width=Inches(3), height=Inches(0.8), size=14)
    textbox(slide, "TextBox 12", [[("CTV Impressions Served", True)], [("112,090", False)]],
            left=Inches(3.94), top=Inches(1.14), width=Inches(3), height=Inches(0.8), size=14)
    textbox(slide, "Rectangle 6", [[("Top Assets By Impression", True)]],
            left=Inches(0.43), top=assets_top - Inches(0.43), width=Inches(2.89), height=Inches(0.34), size=14)
    assets = [["Asset Name", "CTV Impressions", "Accounts Reached", "VCR"]]
    assets += [[ASSET_NAME.format(i), f"{5000 + i:,}", "900", "99%"] for i in range(n_assets)]
    widths = [Inches(8.09), Inches(1.73), Inches(1.77), Inches(0.91)]
    frame = table(slide, assets, name="Table 21", left=Inches(0.42), top=assets_top, width=sum(widths),
                  height=row_h * len(assets), col_widths=widths)
    for row in frame.table.rows:
        row.height = row_h
    stored_bottom = assets_top + row_h * len(assets)
    heading_top = stored_bottom + Inches(0.16)
    textbox(slide, "Rectangle 7", [[("Top Accounts By Impression", True)]],
            left=Inches(0.43), top=heading_top, width=Inches(3.19), height=Inches(0.34), size=14)
    accounts = [["Domain", "CTV Impressions", "Site Visits", "# Trending Topics"]]
    accounts += [[f"account{i}.com", "9,000", "623", "24"] for i in range(n_accounts)]
    second = table(slide, accounts, name="Table 22", left=Inches(0.42), top=heading_top + Inches(0.37),
                   width=Inches(12.5), height=row_h * len(accounts))
    for row in second.table.rows:
        row.height = row_h
    return slide


def kpi_column(slide, stats, name="TextBox 7"):
    """A Program Performance stat column: big bold stat, its label, blank
    spacer, ... stats: [(stat, label), ...]."""
    shape = slide.shapes.add_textbox(Inches(10.1), Inches(1.2), Inches(2.94), Inches(3.74))
    shape.name = name
    tf = shape.text_frame
    first = True
    for i, (stat, label) in enumerate(stats):
        if i:
            spacer = tf.add_paragraph()
            spacer.text = ""
        for text, size, bold in ((stat, 26, True), (label, 14, None)):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            if bold:
                r.font.bold = True
    return shape


def account_engagement_slide(prs, takeaway=("{{2}}% of accounts were trending on 5+ intent topics.", "",
                                            "Top job titles for leads from  include  and ."), with_table=True):
    slide = blank(prs, f"Client X Combined View {EN_DASH} Account Engagement")
    textbox(slide, "TextBox 4", [[("Top Accounts: Display, and CTV", True)]], left=Inches(0.42),
            top=Inches(1.24), width=Inches(7.25), height=Inches(0.27), size=16)
    if with_table:
        rows = [["Account Domain", "Leads", "Display Impressions", "Clicks", "CTV Impressions", "Site Visits",
                 "# Trending Topics"],
                ["beta-corp.com", "0", "300,100", "328", "8,000", "668", "31"],
                ["alpha-corp.com", "0", "350,200", "288", "9,000", "623", "24"],
                ["gamma-corp.com", "0", "250,300", "237", "7,000", "402", "27"],
                ["delta-corp.com", "0", "200,400", "262", "6,000", "533", "23"]]
        table(slide, rows, name="Table 3", left=Inches(0.42), top=Inches(1.59), width=Inches(12.5))
    paragraphs = []
    for line in takeaway:
        if line.startswith("{{"):
            token, rest = line.split("%", 1)
            paragraphs.append([(token + "%", True), (rest, None)])
        else:
            paragraphs.append(line)
    textbox(slide, "TextBox 2", paragraphs, left=Inches(0.37), top=Inches(5.0), width=Inches(12.52),
            height=Inches(0.91))
    return slide


def key_takeaways_slide(prs):
    """The Key Takeaways summary slide with the generator's blank values."""
    slide = blank(prs, "Key Takeaways", title_name="Title 7")
    box = textbox(slide, "TextBox 8", [[("Top Industries", True)], [("Manufacturing ", True), (EN_DASH + " %", True)],
                                       [("Textiles ", True), (EN_DASH + " %", True)]],
                  left=Inches(6.67), top=Inches(1.19), width=Inches(6.33), height=Inches(1.45), size=14)
    titles = textbox(slide, "TextBox 9", [[("Top Job Titles", True)], [(" " + EN_DASH + " %", True)]],
                     left=Inches(0.33), top=Inches(5.22), width=Inches(6.07), height=Inches(1.0), size=14)
    sizes = textbox(slide, "TextBox 20", [[("Company Size", True)], [("% ", True), ("of companies had ", None),
                                                                      ("1-9", True), (" employees", None)]],
                    left=Inches(6.67), top=Inches(5.22), width=Inches(6.33), height=Inches(1.0), size=14)
    for shape in (box, titles, sizes):  # bulleted list items, as on the real slide
        for p in shape.text_frame.paragraphs[1:]:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", "228600")
            ppr.set("indent", "-228600")
            bu = ppr.makeelement(qn_("a:buChar"), {"char": "•"})
            ppr.append(bu)
    return slide


def qn_(tag):
    from pptx.oxml.ns import qn
    return qn(tag)


def long_pacing_deck(n_rows=18):
    """A CS + Display deck whose pacing table has long, wrapping campaign
    names: ContentSynd rows (Leads) and Display rows (Imps)."""
    prs = new_prs()
    title_slide(prs, ["ABM Content Syndication: 05/26/2026 - 10/01/2026", "ABM Display 05/26/2026 - 10/01/2026"])
    rows = []
    for i in range(n_rows):
        if i % 2:
            name = f"Int_X_ABM Display_ALLTAL_Global_Extended_Audience_Segment_{i:02d}_Q3-26"
            units = ["40,000", "40,000", "0"]
        else:
            name = f"Int_X_ABM ContentSynd_APACTAL_Singapore_Extended_Audience_Segment_{i:02d}_Q3-26"
            units = ["25", "25", "0"]
        rows.append([name, "$1,000", "$1,000", "$0", *units, "100%"])
    pacing(prs, rows)
    thank_you(prs)
    return prs


def thank_you(prs, lines=("Owner Name", "Owner Title", "Owner Email")):
    slide = blank(prs)
    textbox(slide, "Text Placeholder 1", ["Thank you!"], top=Inches(3))
    textbox(slide, "TextBox 7", [[(lines[0], True)], [(lines[1], False)], [(lines[2], False)]],
            left=Inches(1.6), top=Inches(5.6), width=Inches(7.2), height=Inches(1.0), size=18)
    return slide


def cs_deck(**overrides):
    """A raw CS report: every slide type the CS rules touch."""
    prs = new_prs()
    title_slide(prs, ["ABM Content Syndication: 08/04/2026 - 10/01/2026"])
    agenda = blank(prs)
    textbox(agenda, "Text Placeholder 2", ["AGENDA"])
    highlights = blank(prs, "Campaign Highlights", title_name="Title 7")
    textbox(highlights, "TextBox 7", ["Company Size" + NBSP,
                                      [("44% ", True), ("of companies had ", None), ("1000-4999", True),
                                       (" employees", None)],
                                      [("28% ", True), ("of companies had ", None), ("10000 ", True),
                                       ("employees", None)]])
    pacing(prs, overrides.get("pacing_rows") or [
        ["Client_NAMER_ABM_AV_2CQ_Direct_Q226", "$1,768", "$1,768", "$0", "26", "26", "0", "100%"],
        ["Client_NAMER_ABM_2C2CQ_Direct_Q226", "$19,992", "$19,992", "$0", "294", "294", "0", "100%"],
    ])
    section(prs, "Content Syndication:\nAudience, Content, and Industry Insights\nClient_NAMER_ABM")
    audience_reach(prs)
    audience_insights(prs, list(overrides.get("topics") or ["Device Security", "Device Enrollment",
                                                             "Endpoint Management", "Acme Widgets"]))
    content_insights(prs)
    industry_insights(prs, **({"job_titles": overrides["job_titles"]} if "job_titles" in overrides else {}))
    country_insights(prs, **({"countries": overrides["countries"]} if "countries" in overrides else {}))
    top_accounts_cs(prs)
    custom_question(prs)
    thank_you(prs)
    return prs


def combined_deck():
    """A raw CS + Display report with one Display and one CS section."""
    prs = new_prs()
    title_slide(prs, ["ABM Content Syndication: 05/26/2026 - 10/01/2026", "ABM Display 05/26/2026 - 10/01/2026"])
    pacing(prs, [
        ["Client_Fin Crime_ABM Display_Q1-26", "$8,000", "$8,000", "$0", "320,000", "320,000", "0", "100%"],
        ["Content Syndication - ABM", "$11,946", "$11,946", "$0", "181", "181", "0", "100%"],
    ])
    section(prs, "Audience, Content, and \nIndustry Insights: \nx12ab34: Display - ABM")
    halo(prs)
    top_accounts_display(prs, [["Wells", "wells.com", "7,294", "0", "0", "8", "19"],
                               ["State", "state.com", "7,193", "0", "0", "14", "15"]])
    creative_sets(prs, "x12ab34: Display - ABM")
    section(prs, "Audience, Content, and \nIndustry Insights: \nABM CS")
    halo(prs, categories=("Single-channel", "Two-channel"), site_visits=(0, 0))
    top_accounts_display(prs, [])
    creative_sets(prs, "ABM CS")
    thank_you(prs, ("<ML team member name>", "<title>", "<email>@madisonlogic.com"))
    return prs


def display_deck():
    """A raw Display report: program performance doughnut, float noise, empty KPI."""
    prs = new_prs()
    title_slide(prs, ["ABM Display Advertising: 09/20/2026 - 01/01/2027"])
    pacing(prs, [
        ["Client Telco Exec Digital Video - USA Q3-26", "$24,500", "$4,962", "$19,538", "1,633,333", "330,770",
         "1,302,563", "20%"],
        ["Client Telco Exec Digital Video - Other Q3-26", "$10,500", "$2,863", "$7,637", "700,000", "190,841",
         "509,159", "27%"],
    ], title=f"Pacing: Client / Telco Exec - Q3-26 Delivery Results")
    perf = blank(prs, f"Client Video Q3-26 {EN_DASH} Program Performance" + VT)
    chart(perf, XL_CHART_TYPE.DOUGHNUT, ["Reached", "Targeted"], {"DOMAINS": (19, 4)}, name="Chart 8",
          title="Target Penetration Rate", labels=True)
    textbox(perf, "TextBox 7", ["83%", "Penetration Rate", "", "97.16% ", "of Impressions delivered", "",
                                "+" + NBSP, "site visits generated"], left=Inches(9.2), width=Inches(2.9),
            height=Inches(3.2))
    data = blank(prs, f"Client Video Q3-26 {EN_DASH} Display Data" + VT)
    textbox(data, "TextBox 12", ["Campaign Progress: Impressions", "19.919999999999998%"])
    table(data, [["Asset Name", "Impressions", "Accounts", "Clicks", "CTR"],
                 ["video-30-sec", "325,342", "19", "56", "0.0172%"]], name="Table 21", left=Emu(219009),
          top=Inches(3.4), width=Emu(11642334))
    top_accounts_display(prs, [], title=f"Client Video Q3-26 {EN_DASH} Top Accounts" + NBSP)
    thank_you(prs)
    return prs
