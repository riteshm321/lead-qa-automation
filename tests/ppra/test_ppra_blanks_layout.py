"""Blank KPI stats (R12), channel columns (R27), broken takeaways (R28),
hidden shapes (R29), the usable bottom and the data-issue check."""
import io

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from core.ppra import engine
from core.ppra.blanks import find_data_issues
from core.ppra.detect import analyze, norm_text
from core.ppra.layout import has_hidden_shapes, table_row_heights
from core.ppra.rules import PPRA_BLUE_SCHEME, RULES_BY_ID, usable_bottom
from ppra import _decks as d

OWNER = {"owner_name": "Pat Lee", "owner_title": "CXM", "owner_email": "plee@madisonlogic.com"}


def run_rule(prs, rule_id, report_type="CS", jira=None):
    rule = RULES_BY_ID[rule_id]
    findings = rule.applies(analyze(prs), report_type, jira or {})
    for finding in findings:
        rule.apply(prs, finding)
    return findings


def shape_named(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def paragraphs(shape):
    return [norm_text(p.text) for p in shape.text_frame.paragraphs]


def reopen(prs):
    return Presentation(io.BytesIO(d.to_bytes(prs)))


# ---------------------------------------------------------------- usable bottom

def test_usable_bottom_sits_above_the_layout_footer_text_box():
    prs = d.new_prs()
    footer_top = d.layout_footer(prs, top=Inches(6.3))
    slide = d.blank(prs, "Anything")
    assert usable_bottom(slide, d.SLIDE_H) == footer_top - Inches(0.1)
    # Without the extra box, the layout's own footer placeholders count.
    plain = d.new_prs()
    assert usable_bottom(d.blank(plain, "x"), d.SLIDE_H) == 6356350 - Inches(0.1)


# ---------------------------------------------------------------- R12 blank KPI stats

def test_r12_removes_a_blank_percent_stat_with_its_label_and_one_spacer():
    prs = d.new_prs()
    slide = d.blank(prs, f"Client X {d.EN_DASH} Program Performance")
    box = d.kpi_column(slide, [("80%", "Penetration Rate"), ("%", "of leads delivered from trending accounts"),
                               ("17.83K+", "site visits generated"), ("98.77%", "CTV VCR")])
    findings = run_rule(prs, "R12", "CTV + Display")
    assert [f.short for f in findings] == ["Remove the blank '% of leads delivered from trending accounts' stat"]
    assert paragraphs(box) == ["80%", "Penetration Rate", "", "17.83K+", "site visits generated", "",
                               "98.77%", "CTV VCR"]
    assert run_rule(prs, "R12", "CTV + Display") == []  # idempotent


def test_r12_never_removes_a_stat_whose_label_has_a_digit_or_a_stat_with_a_number():
    prs = d.new_prs()
    slide = d.blank(prs, f"Client X {d.EN_DASH} Program Performance")
    box = d.kpi_column(slide, [("0%", "of accounts trending"), ("%", "of accounts on 5+ topics"),
                               ("12", "site visits")])
    assert run_rule(prs, "R12", "Display") == []
    assert paragraphs(box)[0] == "0%" and "%" in paragraphs(box)


def test_r12_removes_a_stat_box_and_the_label_box_under_it():
    prs = d.new_prs()
    slide = d.blank(prs, f"Client X {d.EN_DASH} Program Performance")
    d.textbox(slide, "TextBox 30", [[("%", True)]], left=Inches(10), top=Inches(2), width=Inches(2),
              height=Inches(0.5), size=26)
    d.textbox(slide, "TextBox 31", ["of leads delivered from trending accounts"], left=Inches(10),
              top=Inches(2.55), width=Inches(2.9), height=Inches(0.6), size=14)
    d.textbox(slide, "TextBox 32", [[("80%", True)]], left=Inches(10), top=Inches(4), width=Inches(2),
              height=Inches(0.5), size=26)
    findings = run_rule(prs, "R12", "Display")
    assert [f.short for f in findings] == ["Remove the blank '% of leads delivered from trending accounts' stat"]
    names = {s.name for s in slide.shapes}
    assert "TextBox 30" not in names and "TextBox 31" not in names and "TextBox 32" in names
    assert run_rule(prs, "R12", "Display") == []


def test_r12_removes_only_a_known_kpi_tile_with_no_value():
    prs = d.new_prs()
    slide = d.blank(prs, f"Client X {d.EN_DASH} Display Data")
    d.textbox(slide, "TextBox 18", [[("Exposure Time (Hours)", True)], ""], width=Inches(2.5),
              height=Inches(0.8), size=14)
    d.textbox(slide, "TextBox 20", [[("Avg CTR", True)], [("0.08%", False)]], left=Inches(4), width=Inches(2.5),
              height=Inches(0.8), size=14)
    d.textbox(slide, "TextBox 21", [[("Something New", True)], ""], left=Inches(7), width=Inches(2.5),
              height=Inches(0.8), size=14)
    d.table(slide, [["Asset Name", "Impressions"], ["banner", "1,000"]], top=Inches(3))
    findings = run_rule(prs, "R12", "Display")
    assert [f.short for f in findings] == ["Remove the blank 'Exposure Time (Hours)' tile"]
    names = {s.name for s in slide.shapes}
    assert "TextBox 18" not in names and {"TextBox 20", "TextBox 21"} <= names


def test_r12_leaves_the_halo_x_for_r23_and_the_audience_reach_percent_for_r25():
    prs = d.new_prs()
    d.halo(prs)
    d.audience_reach(prs, pct="% ")
    assert run_rule(prs, "R12", "CS + Display") == []


def test_r12_and_r28_leave_key_takeaways_list_items_alone():
    prs = d.new_prs()
    slide = d.key_takeaways_slide(prs)
    before = [paragraphs(s) for s in slide.shapes if s.has_text_frame]
    assert run_rule(prs, "R12", "Display") == []
    assert run_rule(prs, "R28", "Display") == []
    assert [paragraphs(s) for s in slide.shapes if s.has_text_frame] == before


# ---------------------------------------------------------------- R27 channel columns

def _engagement_table(prs):
    return shape_named(analyze(prs).slides[0].slide, "Table 3")


def test_r27_removes_the_leads_column_when_the_report_has_no_cs():
    prs = d.new_prs()
    d.account_engagement_slide(prs)
    findings = run_rule(prs, "R27", "CTV + Display")
    assert [f.short for f in findings] == ["Remove the Leads column (not a CS report)"]
    frame = _engagement_table(prs)
    assert [c.text for c in frame.table.rows[0].cells] == [
        "Account Domain", "Display Impressions", "Clicks", "CTV Impressions", "Site Visits", "# Trending Topics"]
    assert abs(sum(c.width for c in frame.table.columns) - Inches(12.5)) <= 10
    # CTV stays, so the heading keeps naming it.
    heading = shape_named(analyze(prs).slides[0].slide, "TextBox 4")
    assert norm_text(heading.text_frame.text) == "Top Accounts: Display, and CTV"
    assert run_rule(prs, "R27", "CTV + Display") == []


def test_r27_removes_ctv_columns_and_retitles_in_a_display_report():
    prs = d.new_prs()
    d.account_engagement_slide(prs)
    findings = run_rule(prs, "R27", "Display")
    assert len(findings) == 1 and "Leads, CTV Impressions columns" in findings[0].short
    frame = _engagement_table(prs)
    assert [c.text for c in frame.table.rows[0].cells] == [
        "Account Domain", "Display Impressions", "Clicks", "Site Visits", "# Trending Topics"]
    heading = shape_named(analyze(prs).slides[0].slide, "TextBox 4")
    assert norm_text(heading.text_frame.text) == "Top Accounts: Display"


def test_r27_keeps_leads_in_a_cs_report_and_never_touches_generic_headers():
    prs = d.new_prs()
    d.account_engagement_slide(prs)
    assert run_rule(prs, "R27", "CS + Display + CTV") == []
    prs = d.new_prs()
    slide = d.blank(prs, f"Client X {d.EN_DASH} Display Data")
    d.table(slide, [["Asset Name", "Impressions", "Clicks", "CTR", "VCR"], ["a", "1", "2", "3%", "4%"]])
    assert run_rule(prs, "R27", "CTV") == []  # generic headers are never a channel's


# ---------------------------------------------------------------- R28 broken takeaways

def test_r28_writes_the_standard_account_takeaway_when_every_sentence_is_broken():
    prs = d.new_prs()
    d.account_engagement_slide(prs)
    findings = run_rule(prs, "R28", "CTV + Display")
    assert [f.short for f in findings] == ["Fix the key takeaway (unfilled values)"]
    box = shape_named(analyze(prs).slides[0].slide, "TextBox 2")
    assert [t for t in paragraphs(box) if t] == [
        "alpha-corp.com, beta-corp.com and gamma-corp.com were the most engaged accounts, led by alpha-corp.com with "
        "350,200 display impressions."]
    number = next(r for p in box.text_frame.paragraphs for r in p.runs if r.text == "350,200")
    assert number.font.bold and number._r.find(qn("a:rPr")).find(qn("a:solidFill")).find(
        qn("a:schemeClr")).get("val") == PPRA_BLUE_SCHEME
    assert run_rule(prs, "R28", "CTV + Display") == []  # idempotent


def test_r28_fills_the_trending_value_from_the_deck_and_drops_the_empty_sentence():
    prs = d.new_prs()
    d.audience_insights(prs, ["Ai", "Cloud", "Data"], pct_runs=(("37% ", True), ("of accounts are trending on ", None),
                                                                ("5+ ", True), ("intent topics", None)))
    d.account_engagement_slide(prs)
    run_rule(prs, "R28", "CS + Display")
    box = shape_named(analyze(prs).slides[1].slide, "TextBox 2")
    assert [t for t in paragraphs(box) if t] == ["37% of accounts were trending on 5+ intent topics."]


def test_r28_removes_broken_text_and_asks_for_a_takeaway_when_there_is_no_table():
    prs = d.new_prs()
    d.title_slide(prs, ["ABM Display: 05/26/2026 - 10/01/2026"])
    d.account_engagement_slide(prs, with_table=False)
    out, log, attention = engine.format(d.to_bytes(prs), "Display", OWNER)
    deck = analyze(Presentation(io.BytesIO(out)))
    box = shape_named(deck.slides[1].slide, "TextBox 2")
    assert not any(paragraphs(box))
    assert "Slide 2 · Write the key takeaway" in engine.checklist_lines(attention)


def test_r28_keeps_good_sentences():
    prs = d.new_prs()
    d.account_engagement_slide(prs, takeaway=("Accounts engaged strongly with the video assets.", "",
                                              "Top job titles for leads from  include  and ."))
    run_rule(prs, "R28", "Display")
    box = shape_named(analyze(prs).slides[0].slide, "TextBox 2")
    assert [t for t in paragraphs(box) if t] == ["Accounts engaged strongly with the video assets."]


# ---------------------------------------------------------------- R29 hidden shapes

def test_r29_moves_the_heading_and_table_hidden_by_a_wrapping_asset_table():
    prs = d.new_prs()
    footer = d.layout_footer(prs, top=Inches(7.0))
    slide = d.ctv_data_slide(prs, assets_top=Inches(2.6))
    assets, heading, accounts = (shape_named(slide, n) for n in ("Table 21", "Rectangle 7", "Table 22"))
    assets_bottom = int(assets.top) + sum(table_row_heights(assets))
    assert int(heading.top) < assets_bottom  # hidden as built
    assert has_hidden_shapes(analyze(prs).slides[0], d.SLIDE_W, d.SLIDE_H)
    findings = run_rule(prs, "R29", "CTV")
    assert [f.short for f in findings] == [
        "Move 'Top Accounts By Impression' and 1 more shape below the 'Asset Name' table (they were hidden)"]
    assert int(heading.top) >= assets_bottom
    assert int(accounts.top) >= int(heading.top) + int(heading.height)
    assert int(accounts.top) + sum(table_row_heights(accounts)) <= footer - Inches(0.1)
    assert int(shape_named(slide, "Title 2").top) == Inches(0.2)  # the title never moves
    assert run_rule(prs, "R29", "CTV") == []  # idempotent
    assert not has_hidden_shapes(analyze(reopen(prs)).slides[0], d.SLIDE_W, d.SLIDE_H)


def test_r29_shrinks_table_text_when_moving_alone_would_cross_the_footer():
    prs = d.new_prs()
    # Moving alone ends about 6.12 in down; 9 pt text saves about 0.16 in.
    footer = d.layout_footer(prs, top=Inches(6.15))
    slide = d.ctv_data_slide(prs, assets_top=Inches(3.0), row_h=Inches(0.27))
    findings = run_rule(prs, "R29", "CTV")
    assert len(findings) == 1 and "pt" in findings[0].short
    assets, heading, accounts = (shape_named(slide, n) for n in ("Table 21", "Rectangle 7", "Table 22"))
    sizes = {int(r.get("sz")) for tr in assets.table._tbl.tr_lst[1:] for r in tr.iter(qn("a:rPr")) if r.get("sz")}
    assert max(sizes) < 1000 and min(sizes) >= 900  # stepped down from 10 pt, never below 9 pt
    assert int(heading.top) >= int(assets.top) + sum(table_row_heights(assets))
    for shape in (heading, accounts):
        height = sum(table_row_heights(shape)) if shape.has_table else int(shape.height)
        assert int(shape.top) + height <= footer - Inches(0.1)
    assert run_rule(prs, "R29", "CTV") == []


def test_r29_leaves_a_slide_that_cannot_fit_and_keeps_it_on_the_checklist():
    prs = d.new_prs()
    d.layout_footer(prs, top=Inches(6.0))
    slide = d.ctv_data_slide(prs, n_assets=6, n_accounts=6, assets_top=Inches(2.6))
    before = {s.name: int(s.top) for s in slide.shapes}
    assert run_rule(prs, "R29", "CTV") == []
    assert {s.name: int(s.top) for s in slide.shapes} == before
    items = engine.attention_items(analyze(prs), "CTV", OWNER)
    assert engine.MSG_OVERLAP in [a.message for a in items]


def test_r29_leaves_shapes_layered_inside_another_frame():
    prs = d.new_prs()
    slide = d.blank(prs, f"Client X {d.EN_DASH} Program Performance")
    d.chart(slide, d.XL_CHART_TYPE.DOUGHNUT, ["a", "b"], {"S": (1, 2)},
            top=Inches(1.2), height=Inches(2.8))
    d.textbox(slide, "TextBox 2", ["1.50K Total Accounts Targeted"], left=Inches(1.5), top=Inches(3.6),
              width=Inches(2.3), height=Inches(0.27), size=10)
    assert run_rule(prs, "R29", "Display") == []


# ---------------------------------------------------------------- data issues

def test_data_issues_are_found_before_formatting_and_listed_after():
    prs = d.new_prs()
    d.title_slide(prs, ["ABM Display: 05/26/2026 - 10/01/2026"])
    d.key_takeaways_slide(prs)
    raw = d.to_bytes(prs)
    scan = engine.scan(raw, "Display", OWNER)
    assert [(i.slide_number, i.count) for i in scan.data_issues] == [(2, 4)]
    assert scan.data_issues[0].line() == "Slide 2 · Key Takeaways · 4 blank values (e.g. 'Manufacturing " + \
        d.EN_DASH + " %')"
    out, _log, attention = engine.format(raw, "Display", OWNER)
    slide = Presentation(io.BytesIO(out)).slides[1]
    assert "Manufacturing " + d.EN_DASH + " %" in paragraphs(shape_named(slide, "TextBox 8"))  # flagged only
    assert "Slide 2 · Fill 4 blank % values" in engine.checklist_lines(attention)


def test_data_issues_cover_tokens_errors_and_empty_substitutions_but_not_header_cells():
    prs = d.new_prs()
    d.account_engagement_slide(prs)
    slide = d.blank(prs, "Other")
    d.textbox(slide, "TextBox 1", ["ERROR:Division by zero", "Fine 12% text"])
    d.table(slide, [["% Delivered", "Name"], ["100%", "Acme"]], top=Inches(4))
    issues = find_data_issues(analyze(prs))
    assert [(i.slide_number, i.count) for i in issues] == [(1, 2), (2, 1)]
    assert not issues[0].only_percent


def test_blank_detection_ignores_real_values():
    from core.ppra.blanks import blank_kind, broken_reason
    for text in ("26% of leads", "0.08%", "$1,000", "17.83K+", "Top job titles: CEO, CTO and CFO.",
                 "Revenue - $ 5"):
        assert blank_kind(text) is None, text
    for text in ("% of companies had 1-9 employees", "Textiles " + d.EN_DASH + " %", "{{2}} accounts", "$",
                 "K+ visits", "X%"):
        assert blank_kind(text) is not None, text
    assert broken_reason("Top job titles for leads from  include  and .") == "empty"
    assert broken_reason("A sentence with one  double space") is None


def test_formatting_a_ctv_display_deck_end_to_end_is_idempotent():
    prs = d.new_prs()
    d.layout_footer(prs, top=Inches(7.0))
    d.title_slide(prs, ["ABM Display 05/26/2026 - 10/01/2026", "ABM CTV 05/26/2026 - 10/01/2026"])
    perf = d.blank(prs, f"Client X Combined View {d.EN_DASH} Program Performance")
    d.kpi_column(perf, [("80%", "Penetration Rate"), ("%", "of leads delivered from trending accounts")])
    d.ctv_data_slide(prs, assets_top=Inches(2.6))
    d.account_engagement_slide(prs)
    d.thank_you(prs)
    out, log, attention = engine.format(d.to_bytes(prs), "CTV + Display", OWNER)
    ids = {line.split(":")[0] for line in log}
    assert {"R12", "R27", "R28", "R29"} <= ids
    again = engine.scan(out, "CTV + Display", OWNER)
    assert not [rid for rid, items in again.findings.items() if items and rid in ("R12", "R27", "R28", "R29")]
    assert engine.MSG_OVERLAP not in [a.message for a in attention]
