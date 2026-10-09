import io
import zipfile

from pptx import Presentation

from core.ppra import engine
from core.ppra.detect import analyze
from core.ppra.rules import RULES
from ppra import _decks as d

OWNER = {"owner_name": "Pat Lee", "owner_title": "CXM", "owner_email": "plee@madisonlogic.com"}


def test_scan_reports_every_rule():
    result = engine.scan(d.to_bytes(d.cs_deck()), "CS", OWNER)
    assert set(result.findings) == {r.id for r in RULES}
    assert result.detected_report_type == "CS"
    assert {"R1", "R2", "R4", "R7", "R15", "R20"} <= set(result.found_rule_ids())
    assert "R9" not in result.found_rule_ids()  # CS report: CTV rule never runs


def test_format_deletes_slides_cleanly_and_file_reopens():
    out, log, _attention = engine.format(d.to_bytes(d.cs_deck()), "CS", OWNER)
    prs = Presentation(io.BytesIO(out))
    kinds = [s.kind for s in analyze(prs).slides]
    assert "top_accounts_cs" not in kinds and "custom_question" not in kinds
    assert len(prs.slides) == 11
    # The deleted slides' parts are gone from the package, not just hidden.
    names = zipfile.ZipFile(io.BytesIO(out)).namelist()
    assert len([n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]) == 11
    rels = prs.part.rels
    assert len([r for r in rels.values() if r.reltype.endswith("/slide")]) == 11
    # Saving the re-opened file again works too.
    Presentation(io.BytesIO(engine.save(prs)))
    assert any(line.startswith("R4:") for line in log) and any(line.startswith("R7:") for line in log)


def test_format_is_idempotent():
    out, _log, _att = engine.format(d.to_bytes(d.cs_deck()), "CS", OWNER)
    again, log, _att = engine.format(out, "CS", OWNER)
    assert log == []
    assert len(Presentation(io.BytesIO(again)).slides) == len(Presentation(io.BytesIO(out)).slides)


def test_format_combined_deck_is_idempotent():
    out, log, _att = engine.format(d.to_bytes(d.combined_deck()), "CS + Display", OWNER)
    assert any(line.startswith("R9:") for line in log)
    assert engine.scan(out, "CS + Display", OWNER).found_rule_ids() == []


def test_only_enabled_rules_run():
    out, log, _att = engine.format(d.to_bytes(d.cs_deck()), "CS", OWNER, enabled_rule_ids=["R1"])
    assert log and all(line.startswith("R1:") for line in log)
    assert len(Presentation(io.BytesIO(out)).slides) == len(d.cs_deck().slides)


def test_deletions_skip_other_changes_on_deleted_slides():
    prs = d.cs_deck(countries=(("United States", "181", "100%"),))
    _out, log, _att = engine.format(d.to_bytes(prs), "CS", OWNER)
    assert not any(line.startswith("R18:") or line.startswith("R14:") for line in log)
    assert any(line.startswith("R5:") for line in log)


def test_attention_items_flag_what_the_tool_cannot_do():
    _out, _log, attention = engine.format(d.to_bytes(d.combined_deck()), "CS + Display", {})
    messages = [str(a) for a in attention]
    assert any("Creative Sets" in m for m in messages)
    assert any("placeholder owner details" in m for m in messages)


def test_delete_slides_removes_section_list_references():
    prs = d.cs_deck()
    ext_lst = prs.part._element.makeelement(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}extLst", {})
    ns = engine._P14_NS
    ext = ext_lst.makeelement("{http://schemas.openxmlformats.org/presentationml/2006/main}ext",
                              {"uri": "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"})
    section_lst = ext.makeelement(f"{{{ns}}}sectionLst", {})
    section = section_lst.makeelement(f"{{{ns}}}section", {"name": "All", "id": "{00000000-0000-0000-0000-000000000001}"})
    ids = section.makeelement(f"{{{ns}}}sldIdLst", {})
    for entry in prs.slides._sldIdLst:
        ids.append(ids.makeelement(f"{{{ns}}}sldId", {"id": entry.get("id")}))
    section.append(ids)
    section_lst.append(section)
    ext.append(section_lst)
    ext_lst.append(ext)
    prs.part._element.append(ext_lst)

    engine.delete_slides(prs, [10, 11])
    remaining = {e.get("id") for e in prs.slides._sldIdLst}
    referenced = {e.get("id") for e in prs.part._element.iter(f"{{{ns}}}sldId")}
    assert referenced == remaining
    Presentation(io.BytesIO(engine.save(prs)))
