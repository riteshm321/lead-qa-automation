import io

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from core.ppra.detect import analyze, iter_shapes, norm_text
from core.ppra.rules import RULES_BY_ID, TABLE_WIDTH, TAKEAWAY_TEMPLATES, estimate_text_height, wrapped_lines
from ppra import _decks as d

OWNER = {"owner_name": "Pat Lee", "owner_title": "CXM", "owner_email": "plee@madisonlogic.com"}


def run_rule(prs, rule_id, report_type="CS", jira=None):
    rule = RULES_BY_ID[rule_id]
    findings = rule.applies(analyze(prs), report_type, jira or {})
    for finding in findings:
        rule.apply(prs, finding)
    return findings


def reopen(prs):
    return Presentation(io.BytesIO(d.to_bytes(prs)))


def shape_named(slide, name):
    return next(s for s in iter_shapes(slide.shapes) if s.name == name)


def slide_of(prs, kind):
    return analyze(prs).of_kind(kind)[0].slide


def paragraphs(shape):
    return [norm_text(p.text) for p in shape.text_frame.paragraphs]


def last(rows):
    return rows[len(rows) - 1]


def pacing_table(prs):
    return shape_named(slide_of(prs, "pacing"), "Table 5")


# ---------------------------------------------------------------- R1-R3 pacing

def test_r1_adds_leads_to_unit_values_only():
    prs = d.cs_deck()
    assert len(run_rule(prs, "R1")) == 1
    rows = [[c.text for c in r.cells] for r in pacing_table(prs).table.rows]
    assert rows[1] == ["Client_NAMER_ABM_AV_2CQ_Direct_Q226", "$1,768", "$1,768", "$0", "26 Leads", "26 Leads",
                       "0 Leads", "100%"]
    assert run_rule(prs, "R1") == []  # idempotent: no "26 Leads Leads"


def test_r1_uses_campaign_keywords_for_imps_and_leads():
    prs = d.combined_deck()
    run_rule(prs, "R1", "CS + Display")
    rows = [[c.text for c in r.cells] for r in pacing_table(prs).table.rows]
    assert rows[1][4:7] == ["320,000 Imps", "320,000 Imps", "0 Imps"]
    assert rows[2][4:7] == ["181 Leads", "181 Leads", "0 Leads"]


def test_r2_total_row_values_and_format():
    prs = d.cs_deck()
    run_rule(prs, "R1")
    frame = pacing_table(prs)
    height_before = frame.height
    assert len(run_rule(prs, "R2")) == 1
    table = frame.table
    total = [c.text for c in last(table.rows).cells]
    assert total == ["Total", "$21,760", "$21,760", "$0", "320 Leads", "320 Leads", "0 Leads", "100%"]
    assert last(table.rows).height == table.rows[0].height
    assert frame.height == height_before + table.rows[0].height
    cell = table.cell(len(table.rows) - 1, 1)
    run = cell.text_frame.paragraphs[0].runs[0]
    assert run.font.bold and run.font.name == "Montserrat"
    tcpr = cell._tc.tcPr
    assert tcpr.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val") == "EBEEF5"
    assert (tcpr.get("marL"), tcpr.get("marB"), tcpr.get("anchor")) == ("6350", "0", "ctr")
    assert run_rule(prs, "R2") == []  # never a second Total row


def test_r2_mixed_units_are_two_paragraphs_in_one_cell():
    prs = d.combined_deck()
    run_rule(prs, "R1", "CS + Display")
    run_rule(prs, "R2", "CS + Display")
    total = last(pacing_table(prs).table.rows).cells
    assert [p.text for p in total[4].text_frame.paragraphs] == ["320,000 Imps", "181 Leads"]
    assert total[1].text == "$19,946"
    # % Delivered per unit type, Imps then Leads, never blended.
    assert [p.text for p in total[7].text_frame.paragraphs] == ["100%", "100%"]


def test_r2_percent_is_sum_delivered_over_sum_goal():
    prs = d.display_deck()
    run_rule(prs, "R2", "Display")
    total = [c.text for c in last(pacing_table(prs).table.rows).cells]
    assert total[1:4] == ["$35,000", "$7,825", "$27,175"]
    assert total[4:] == ["2,333,333 Imps", "521,611 Imps", "1,811,722 Imps", "22%"]


def test_r3_rebalances_then_fits_12_5_inches_centred():
    prs = d.cs_deck()
    before = [c.width for c in pacing_table(prs).table.columns]
    run_rule(prs, "R3")
    frame = pacing_table(prs)
    widths = [c.width for c in frame.table.columns]
    assert sum(widths) == TABLE_WIDTH and frame.width == TABLE_WIDTH
    assert frame.left == (d.SLIDE_W - TABLE_WIDTH) // 2
    assert widths[0] < before[0]
    assert widths[4] > widths[2]  # Units wider than Budget after the rebalance
    assert run_rule(prs, "R3") == []  # the -600000/+200000 shift happens once


# ---------------------------------------------------------------- R4-R8 deletions

def _deleted_titles(prs, rule_id, report_type="CS"):
    deck = analyze(prs)
    return [deck.slides[f.slide_index].title for f in RULES_BY_ID[rule_id].applies(deck, report_type, {})]


def test_r4_marks_header_only_and_default_chart_slides_and_keeps_data_slides():
    prs = d.cs_deck(countries=())
    titles = _deleted_titles(prs, "R4")
    assert f"Top Accounts {d.EN_DASH} Content Syndication" in titles
    assert "Country Insights" in titles
    assert "Content Insights" not in titles and "Industry Insights" not in titles


def test_r4_never_deletes_a_slide_with_real_data():
    prs = d.new_prs()
    d.top_accounts_cs(prs, rows=[["Acme", "acme.com", "3", "uem", "5"]])
    assert _deleted_titles(prs, "R4") == []


def test_r5_single_country_slide():
    assert _deleted_titles(d.cs_deck(countries=(("United States", "181", "100%"),)), "R5") == ["Country Insights"]
    assert _deleted_titles(d.cs_deck(), "R5") == []


def test_r6_deletes_halo_only_in_single_channel_reports():
    assert _deleted_titles(d.combined_deck(), "R6", "CS + Display") == []
    assert len(_deleted_titles(d.combined_deck(), "R6", "Display")) == 2


def test_r7_custom_question_example():
    assert _deleted_titles(d.cs_deck(), "R7") == [f"Custom Question {d.EN_DASH} Example"]


def test_r8_creative_sets_only_in_cs_sections():
    assert _deleted_titles(d.combined_deck(), "R8", "CS + Display") == [
        f"ABM CS{d.EN_DASH} Display Data, Creative Sets"]


# ---------------------------------------------------------------- R9 / R21 tables

def test_r9_drops_zero_ctv_column_and_retitles_for_reports_with_display():
    prs = d.combined_deck()
    assert run_rule(prs, "R9", "CS") == []  # CS-only reports never touch it
    findings = run_rule(prs, "R9", "CS + Display")
    assert len(findings) == 1
    slide = prs.slides[findings[0].slide_index]
    frame = shape_named(slide, "Table 2")
    assert "CTV Impressions" not in [c.text for c in frame.table.rows[0].cells]
    assert len(frame.table.columns) == 6
    assert norm_text(shape_named(slide, "Title 2").text_frame.text) == f"Top Accounts {d.EN_DASH} Display"
    assert run_rule(prs, "R9", "CS + Display") == []


def test_r21_top_accounts_table_12_5_inches_centred():
    prs = d.combined_deck()
    findings = run_rule(prs, "R21", "CS + Display")
    assert findings
    frame = shape_named(prs.slides[findings[0].slide_index], "Table 2")
    assert frame.width == TABLE_WIDTH and sum(c.width for c in frame.table.columns) == TABLE_WIDTH
    assert frame.left == (d.SLIDE_W - TABLE_WIDTH) // 2
    assert run_rule(prs, "R21", "CS + Display") == []


# ---------------------------------------------------------------- text rules

def test_r10_hyperlinks_the_split_platform_url():
    prs = d.cs_deck()
    assert len(run_rule(prs, "R10")) == 2
    prs = reopen(prs)
    box = shape_named(slide_of(prs, "audience_reach"), "TextBox 14")
    runs = box.text_frame.paragraphs[0].runs
    linked = [r for r in runs if r.hyperlink.address]
    url = "https://platform.madisonlogic.com/platform/3836/programs/139823"
    assert "".join(r.text for r in linked) == url
    assert all(r.hyperlink.address == url for r in linked)
    assert norm_text(box.text_frame.text) == f"Click to view on ML Platform {url}"
    assert run_rule(prs, "R10") == []


def test_r11_rounds_float_noise_to_two_places():
    prs = d.display_deck()
    run_rule(prs, "R11", "Display")
    slide = slide_of(prs, "display_performance")
    assert paragraphs(shape_named(slide, "TextBox 12"))[1] == "19.92%"
    assert shape_named(slide, "Table 21").table.cell(1, 4).text == "0.02%"
    assert run_rule(prs, "R11", "Display") == []


def test_r12_removes_empty_site_visit_kpi_lines():
    prs = d.display_deck()
    run_rule(prs, "R12", "Display")
    box = shape_named(slide_of(prs, "program_performance"), "TextBox 7")
    assert paragraphs(box) == ["83%", "Penetration Rate", "", "97.16%", "of Impressions delivered"]


def test_r13_company_size_thousands_and_open_top_band():
    prs = d.cs_deck()
    run_rule(prs, "R13")
    box = shape_named(slide_of(prs, "highlights"), "TextBox 7")
    assert paragraphs(box)[1:] == ["44% of companies had 1,000-4,999 employees",
                                   "28% of companies had 10,000+ employees"]
    assert run_rule(prs, "R13") == []


def _label_rprs(chart_shape):
    return list(chart_shape.chart._chartSpace.iter(qn("a:defRPr")))


def test_r14_country_pie_labels_bold_and_one_point_bigger_once():
    prs = d.cs_deck()
    run_rule(prs, "R14")
    run_rule(prs, "R14")
    rprs = [r for r in _label_rprs(shape_named(slide_of(prs, "country_insights"), "Chart 2"))
            if r.getparent().getparent().getparent().getparent().tag == qn("c:dLbls")]
    assert rprs and all(r.get("b") == "1" and r.get("sz") == "1000" for r in rprs)


def test_r14_doughnut_labels_bold_12pt():
    prs = d.display_deck()
    run_rule(prs, "R14", "Display")
    chart = shape_named(slide_of(prs, "program_performance"), "Chart 8")
    rprs = [r for r in _label_rprs(chart) if r.getparent().getparent().getparent().getparent().tag == qn("c:dLbls")]
    assert rprs and all(r.get("b") == "1" and r.get("sz") == "1200" for r in rprs)
    assert run_rule(prs, "R14", "Display") == []


def _takeaway(prs, kind, name):
    return paragraphs(shape_named(slide_of(prs, kind), name))


def test_r15_audience_reach_top_three_topics():
    prs = d.cs_deck()
    run_rule(prs, "R15")
    text = _takeaway(prs, "audience_reach", "Rectangle 13")
    assert text == ["Key Takeaways:", "26%",
                    "of leads delivered from accounts actively engaging with the following top 3 intent topics:",
                    "", "Device Security", "Device Enrollment", "Endpoint Management"]
    box = shape_named(slide_of(prs, "audience_reach"), "Rectangle 13")
    numbered = [p for p in box.text_frame.paragraphs if p._p.find(".//" + qn("a:buAutoNum")) is not None]
    assert len(numbered) == 3
    assert run_rule(prs, "R15") == []


@pytest.mark.parametrize("topics, sentence", [
    (["Device Security", "Device Enrollment"], "engaging with the following top 2 intent topics:"),
    (["Device Security"], "engaging with the following top intent topic:"),
])
def test_r15_fewer_than_three_topics(topics, sentence):
    prs = d.cs_deck(topics=topics)
    run_rule(prs, "R15")
    text = _takeaway(prs, "audience_reach", "Rectangle 13")
    assert text[2].endswith(sentence)
    assert text[4:] == topics


def test_r15_skips_when_the_lead_percentage_is_blank_and_not_computable():
    prs = d.new_prs()
    d.audience_reach(prs, pct="% ", engagement=((10, 5, 0), (5, 2, 0)))
    d.audience_insights(prs, ["A", "B", "C"])
    assert run_rule(prs, "R15") == []


def test_r16_content_insights_top_job_titles():
    prs = d.cs_deck()
    run_rule(prs, "R16")
    assert _takeaway(prs, "content_insights", "Rectangle 14") == [
        "Key Takeaways:",
        "Assets that generated the strongest engagement are closely aligned with the following top 3 job titles:",
        "", "Chief Operations Officer", "Chief Technology Officer", "Chief Information Officer"]
    box = shape_named(slide_of(prs, "content_insights"), "Rectangle 14")
    assert box.text_frame.paragraphs[1].runs[0].font.size.pt == 16  # the old line's formatting
    assert run_rule(prs, "R16") == []


@pytest.mark.parametrize("titles, sentence", [
    ((("CEO", "5", "50%"), ("CTO", "5", "50%")), "with the following top 2 job titles:"),
    ((("CEO", "5", "100%"),), "with the following top job title:"),
])
def test_r16_fewer_job_titles(titles, sentence):
    prs = d.cs_deck(job_titles=titles)
    run_rule(prs, "R16")
    text = _takeaway(prs, "content_insights", "Rectangle 14")
    assert text[1].endswith(sentence) and text[3:] == [t[0] for t in titles]


def test_r17_percentage_template_and_placeholder_removed():
    prs = d.new_prs()
    d.audience_insights(prs, ["A", "B", "C"], pct_runs=(("79% ", True), ("of accounts are trending on ", None),
                                                        ("5+ ", True), ("intent topics", None)))
    run_rule(prs, "R17")
    text = _takeaway(prs, "audience_insights", "Rectangle 1")
    assert text[:2] == ["Key Takeaways:", TAKEAWAY_TEMPLATES["audience_insights"].format(pct="79")]
    assert not any(t.startswith("[Topics trending") for t in text)
    runs = shape_named(slide_of(prs, "audience_insights"), "Rectangle 1").text_frame.paragraphs[1].runs
    assert runs[0].text == "79% " and runs[0].font.bold  # original formatting kept
    assert run_rule(prs, "R17") == []


@pytest.mark.parametrize("first_run", ["ERROR:Division by zero% ", "0% "])
def test_r17_error_or_zero_falls_back_to_top_topics(first_run):
    prs = d.new_prs()
    d.audience_insights(prs, ["Fraud", "Forensics"], pct_runs=((first_run, True),
                                                               ("of accounts are trending on 5+ intent topics", None)))
    run_rule(prs, "R17")
    text = _takeaway(prs, "audience_insights", "Rectangle 1")
    assert text[1:5] == ["The following are the top 2 intent topics in your campaign:", "", "Fraud", "Forensics"]
    assert run_rule(prs, "R17") == []


def test_r18_country_takeaway_three_countries():
    prs = d.cs_deck(countries=(("United States", "300", "75%"), ("Canada", "60", "15%"), ("India", "40", "10%")))
    run_rule(prs, "R18")
    box = shape_named(slide_of(prs, "country_insights"), "Rectangle 1")
    p = box.text_frame.paragraphs[2]
    assert norm_text(p.text) == (
        "The United States led lead delivery with 75% (300 of 400 leads), followed by Canada at 15% and India "
        "at 10%. Together, the top three markets contributed 100% of all leads.")
    numbers = ["75%", "300", "400", "15%", "10%", "100%"]
    styled = [r.text for r in p.runs if r.font.bold]
    assert styled == numbers  # every number its own bold run, nothing else bold
    for run in p.runs:
        fill = run._r.find(qn("a:rPr")).find(qn("a:solidFill"))
        if run.text in numbers:
            assert fill.find(qn("a:schemeClr")).get("val") == "bg2"  # PPRA blue, theme reference
        else:
            assert run.font.bold is False and (fill is None or fill.find(qn("a:schemeClr")) is None)
    assert run_rule(prs, "R18") == []


def test_r18_two_countries():
    prs = d.cs_deck()
    run_rule(prs, "R18")
    assert _takeaway(prs, "country_insights", "Rectangle 1")[2] == (
        "The United States led lead delivery with 97% (311 of 320 leads), followed by Canada at 3%. "
        "Together, these two markets contributed 100% of all leads.")


def test_r18_top_pct_counts_only_listed_countries():
    prs = d.cs_deck(countries=(("India", "50", "50%"), ("Canada", "30", "30%"), ("Brazil", "10", "10%"),
                               ("Chile", "10", "10%")))
    run_rule(prs, "R18")
    text = _takeaway(prs, "country_insights", "Rectangle 1")[2]
    assert text.startswith("India led lead delivery with 50% (50 of 100 leads)")
    assert text.endswith("the top three markets contributed 90% of all leads.")


def test_r19_autofit_never_shrinks_and_fits_estimate():
    prs = d.cs_deck()
    run_rule(prs, "R16")
    box = shape_named(slide_of(prs, "content_insights"), "Rectangle 14")
    original = box.height
    run_rule(prs, "R19")
    body = box.text_frame._txBody.find(qn("a:bodyPr"))
    assert body.find(qn("a:spAutoFit")) is not None
    assert box.height >= original and box.height >= estimate_text_height(box)
    assert run_rule(prs, "R19") == []


def test_wrapped_lines_breaks_on_words():
    assert wrapped_lines("aaa bbb ccc", 7) == 2
    assert wrapped_lines("x" * 25, 10) == 3
    assert wrapped_lines("", 10) == 1


def test_r20_thank_you_owner_and_mailto():
    prs = d.cs_deck()
    assert run_rule(prs, "R20", jira={}) == []  # nothing to fill without ticket data
    run_rule(prs, "R20", jira=OWNER)
    prs = reopen(prs)
    box = shape_named(slide_of(prs, "thank_you"), "TextBox 7")
    assert paragraphs(box) == ["Pat Lee", "CXM", "plee@madisonlogic.com"]
    assert box.text_frame.paragraphs[0].runs[0].font.bold  # run formatting kept
    assert box.text_frame.paragraphs[2].runs[0].hyperlink.address == "mailto:plee@madisonlogic.com"
    assert run_rule(prs, "R20", jira=OWNER) == []


def test_r20_replaces_ml_team_member_placeholders():
    prs = d.combined_deck()
    run_rule(prs, "R20", "CS + Display", jira=OWNER)
    assert paragraphs(shape_named(slide_of(prs, "thank_you"), "TextBox 7")) == [
        "Pat Lee", "CXM", "plee@madisonlogic.com"]


# ---------------------------------------------------------------- report types / channels

def _pacing_only(names, units=("100", "100", "0")):
    prs = d.new_prs()
    d.pacing(prs, [[name, "$1,000", "$1,000", "$0", *units, "100%"] for name in names])
    return prs


def _unit_cells(prs):
    return [[c.text for c in r.cells][4] for r in list(pacing_table(prs).table.rows)[1:]]


def test_r1_unit_per_row_from_campaign_keywords():
    prs = _pacing_only(["Client Display Q1", "Client Banner Q1", "Client CTV Q1", "Client OTT Q1",
                        "Client Spotify Q1", "Client Podcast Q1", "Client LinkedIn Q1", "Client Lead Gen Q1",
                        "Client Content Syndication Q1", "Client CS Q1"])
    run_rule(prs, "R1", "CS + Display + Audio + CTV")
    assert _unit_cells(prs) == ["100 Imps"] * 7 + ["100 Leads"] * 3


@pytest.mark.parametrize("report_type, unit", [
    ("CS", "Leads"), ("Display", "Imps"), ("CTV", "Imps"), ("Audio", "Imps"), ("LinkedIn", "Imps"),
    ("CTV + Display", "Imps"),
])
def test_r1_falls_back_to_a_single_unit_report_type(report_type, unit):
    prs = _pacing_only(["Acme_NAMER_ABM_Q226"])
    run_rule(prs, "R1", report_type)
    assert _unit_cells(prs) == [f"100 {unit}"]


@pytest.mark.parametrize("report_type", ["CS + Display", "CS + Audio", "CS + Display + CTV",
                                         "CS + Display + Audio + CTV"])
def test_r1_mixed_report_type_leaves_unknown_rows_for_attention(report_type):
    from core.ppra import engine
    prs = _pacing_only(["Acme_NAMER_ABM_Q226", "Acme Display Q1"])
    run_rule(prs, "R1", report_type)
    assert _unit_cells(prs) == ["100", "100 Imps"]
    messages = [str(a) for a in engine.attention_items(analyze(prs), report_type)]
    assert any("Acme_NAMER_ABM_Q226" in m and "Imps or Leads" in m for m in messages)


def test_r2_all_impression_channels_total_as_one_imps_line():
    prs = _pacing_only(["Client Display Q1", "Client CTV Q1", "Client Audio Q1", "Client CS Q1"])
    run_rule(prs, "R1", "CS + Display + Audio + CTV")
    findings = run_rule(prs, "R2", "CS + Display + Audio + CTV")
    assert findings[0].summary == "Add Total row: 300 Imps + 100 Leads, $4,000"
    total = last(pacing_table(prs).table.rows).cells
    assert [p.text for p in total[4].text_frame.paragraphs] == ["300 Imps", "100 Leads"]
    assert [p.text for p in total[7].text_frame.paragraphs] == ["100%", "100%"]


@pytest.mark.parametrize("report_type, expected", [
    ("CS", 0), ("Display", 1), ("CTV", 1), ("Audio", 1), ("LinkedIn", 1), ("CTV + Display", 1),
    ("CS + Audio", 1), ("CS + Display + Audio + CTV", 1),
])
def test_r9_runs_for_any_report_with_an_impression_channel(report_type, expected):
    assert len(RULES_BY_ID["R9"].applies(analyze(d.combined_deck()), report_type, {})) == expected


def test_r9_keeps_a_ctv_column_with_data_and_its_title():
    prs = d.new_prs()
    d.top_accounts_display(prs, [["Wells", "wells.com", "7,294", "12", "0", "8", "19"]])
    assert run_rule(prs, "R9", "CTV + Display") == []
    assert analyze(prs).slides[0].title == f"Top Accounts {d.EN_DASH} Display and CTV"


def test_every_finding_has_a_short_summary():
    from core.ppra.rules import RULES
    for deck_fn, report_type in ((d.cs_deck, "CS"), (d.combined_deck, "CS + Display"), (d.display_deck, "Display")):
        deck = analyze(deck_fn())
        for rule in RULES:
            for finding in rule.applies(deck, report_type, OWNER):
                assert finding.summary and "slide" not in finding.summary.lower().split(" ")[:1]


# ---------------------------------------------------------------- pacing CS keywords

@pytest.mark.parametrize("name, unit", [
    ("Int_X_ABM ContentSynd_APACTAL_Singapore", "100 Leads"),
    ("Int_X_ABM Display_ALLTAL", "100 Imps"),
    ("Acme Content Q3", "100 Leads"),
    ("Acme_CS_Q3", "100 Leads"),
    ("Acme CS-Q3", "100 Leads"),
    ("Acme cs Q3", "100 Leads"),
    ("CS_Acme_Q3", "100 Leads"),
    ("Acme Display - Content Hub", "100 Imps"),  # "Content" alone never outranks a named channel
    ("Acme CSM Q3", "100"),  # not a CS token: left for the checklist in a mixed report
    ("Acme ECS Q3", "100"),
])
def test_r1_pacing_cs_keywords(name, unit):
    prs = _pacing_only([name])
    run_rule(prs, "R1", "CS + Display")
    assert _unit_cells(prs) == [unit]


def test_r1_contentsynd_rows_get_leads_next_to_display_imps():
    prs = d.long_pacing_deck(4)
    run_rule(prs, "R1", "CS + Display")
    assert _unit_cells(prs) == ["25 Leads", "40,000 Imps", "25 Leads", "40,000 Imps"]


# ---------------------------------------------------------------- R22 channel deletion

def _channel_deck():
    """CS, Display and CTV slides in one deck, sectioned per channel."""
    prs = d.new_prs()
    d.title_slide(prs, ["ABM Content Syndication: x", "ABM Display x"])
    d.pacing(prs, [["Acme CS Q3", "$1", "$1", "$0", "5", "5", "0", "100%"]])
    d.section(prs, "Audience, Content, and \nIndustry Insights: \nAcme Display - ABM")
    d.audience_reach(prs)
    d.ctv_performance(prs)
    d.top_accounts_display(prs, [["Wells", "wells.com", "7,294", "5", "0", "8", "19"]])
    d.creative_sets(prs, "Acme Display - ABM")
    d.section(prs, "Audience, Content, and \nIndustry Insights: \nAcme CS")
    d.top_accounts_cs(prs, [["Acme", "acme.com", "3", "uem", "5"]])
    d.thank_you(prs)
    return prs


def _removed(prs, report_type):
    deck = analyze(prs)
    return {f.slide_index + 1: f.summary for f in RULES_BY_ID["R22"].applies(deck, report_type, {})}


def test_r22_removes_ctv_slides_from_reports_without_ctv():
    assert _removed(_channel_deck(), "CS + Display") == {5: "Remove CTV slide (not in a CS + Display report)"}


def test_r22_removes_a_whole_display_section_from_a_cs_report():
    removed = _removed(_channel_deck(), "CS")
    # Display-section breaker, its insight slide, CTV, Top Accounts Display, Creative Sets.
    assert sorted(removed) == [3, 4, 5, 6, 7]
    assert removed[3] == "Remove Display section breaker (not in a CS report)"
    assert removed[6] == "Remove Display / CTV slide (not in a CS report)"


def test_r22_removes_cs_slides_from_a_display_report_but_keeps_the_frame():
    removed = _removed(_channel_deck(), "Display")
    # CTV and Top Accounts CS; the CS breaker stays because the Thank You
    # page (never removed) sits in its section.
    assert sorted(removed) == [5, 9]


def test_r22_slide_is_listed_once_with_its_channel_reason():
    prs = d.cs_deck()
    d.ctv_performance(prs)
    deck = analyze(prs)
    ctv = deck.of_kind("ctv_performance")[0].index
    assert ctv in {f.slide_index for f in RULES_BY_ID["R22"].applies(deck, "CS", {})}
    for rule_id in ("R4", "R5", "R6", "R7", "R8"):
        assert ctv not in {f.slide_index for f in RULES_BY_ID[rule_id].applies(deck, "CS", {})}


def test_r22_does_nothing_without_a_report_type():
    assert _removed(_channel_deck(), None) == {}


# ---------------------------------------------------------------- R23 Halo

def _halo_text(prs):
    return paragraphs(shape_named(slide_of(prs, "halo"), "Rectangle 13"))


def test_r23_halo_ratio_from_chart():
    prs = d.new_prs()
    # single: 100 accounts, 50 visits (0.5 each); multi: 20 accounts, 40 visits (2.0 each) -> 4.0x
    d.halo(prs, categories=("Single-channel", "Two-channel", "Three-channel"), accounts=(100, 10, 10),
           site_visits=(50, 30, 10))
    assert len(run_rule(prs, "R23", "CS + Display")) == 1
    assert _halo_text(prs)[1:3] == ["4.0x", d.HALO_EXPLAINER]
    run = shape_named(slide_of(prs, "halo"), "Rectangle 13").text_frame.paragraphs[1].runs[0]
    assert run.font.bold and run.font.size.pt == 16  # the placeholder run's formatting
    assert run_rule(prs, "R23", "CS + Display") == []


@pytest.mark.parametrize("value", ["X", "ERROR:Division by zero", "%"])
def test_r23_single_channel_only_gets_the_fallback_sentence(value):
    prs = d.new_prs()
    d.halo(prs, value=value)
    run_rule(prs, "R23", "CS + Display")
    assert _halo_text(prs)[1:] == [TAKEAWAY_TEMPLATES["halo_single_channel"], ""]
    run = shape_named(slide_of(prs, "halo"), "Rectangle 13").text_frame.paragraphs[1].runs[0]
    assert run.font.bold and run.font.size.pt == 16
    assert run_rule(prs, "R23", "CS + Display") == []


def test_r23_leaves_an_uncomputable_value_for_the_checklist():
    from core.ppra import engine
    prs = d.new_prs()
    d.halo(prs, categories=("Single-channel", "Two-channel"), site_visits=(0, 4))
    assert run_rule(prs, "R23", "CS + Display") == []
    messages = [a.message for a in engine.attention_items(analyze(prs), "CS + Display", OWNER)]
    assert messages == [engine.MSG_HALO]


# ---------------------------------------------------------------- R25 Audience Reach stat

@pytest.mark.parametrize("stat", ["0% ", "0", "% ", "X% ", "ERROR:Division by zero% "])
def test_r25_fills_a_blank_stat_from_the_engagement_chart(stat):
    prs = d.new_prs()
    d.audience_reach(prs, pct=stat, engagement=((4866, 818, 784), (3227, 501, 478)))
    assert run_rule(prs, "R25")[0].after == "61%"  # 478 / 784
    run = shape_named(slide_of(prs, "audience_reach"), "Rectangle 13").text_frame.paragraphs[1].runs[0]
    assert run.text == "61% " and run.font.bold and run.font.size.pt == 16
    assert run_rule(prs, "R25") == []


def test_r25_never_overwrites_a_real_value_and_flags_a_zero_denominator():
    from core.ppra import engine
    prs = d.new_prs()
    d.audience_reach(prs, pct="26% ", engagement=((10, 5, 4), (5, 2, 1)))
    assert run_rule(prs, "R25") == []
    prs = d.new_prs()
    d.audience_reach(prs, pct="0% ", engagement=((10, 5, 0), (5, 2, 0)))
    assert run_rule(prs, "R25") == []
    assert engine.MSG_REACH in [a.message for a in engine.attention_items(analyze(prs), "CS", OWNER)]


def test_r15_keeps_a_stat_run_in_the_same_paragraph():
    prs = d.new_prs()
    d.audience_reach(prs, pct="% ", one_paragraph=True)
    d.audience_insights(prs, ["A", "B", "C"])
    run_rule(prs, "R25")
    run_rule(prs, "R15")
    runs = shape_named(slide_of(prs, "audience_reach"), "Rectangle 13").text_frame.paragraphs[1].runs
    assert runs[0].text == "22% " and runs[0].font.bold  # 48 / 217, still its own bold run
    assert norm_text("".join(r.text for r in runs[1:])) == (
        "of leads delivered from accounts actively engaging with the following top 3 intent topics:")
    assert not runs[1].font.bold


# ---------------------------------------------------------------- R19 padding

def test_r19_adds_breathing_room_below_the_text():
    from core.ppra.rules import BREATHING_ROOM, MIN_BOTTOM_INSET
    prs = d.cs_deck()
    run_rule(prs, "R16")
    run_rule(prs, "R19")
    for kind, name in (("content_insights", "Rectangle 14"), ("country_insights", "Rectangle 1")):
        box = shape_named(slide_of(prs, kind), name)
        body = box.text_frame._txBody.find(qn("a:bodyPr"))
        assert int(body.get("bIns")) >= MIN_BOTTOM_INSET + BREATHING_ROOM
        assert box.height >= estimate_text_height(box)
    assert run_rule(prs, "R19") == []


def test_r19_never_grows_past_the_bottom_margin():
    from core.ppra import engine
    from core.ppra.rules import BOTTOM_MARGIN
    prs = d.new_prs()
    d.content_insights(prs)
    d.industry_insights(prs, job_titles=tuple((f"A very long job title number {i} " * 6, "1", "1%")
                                              for i in range(3)))
    run_rule(prs, "R16")
    run_rule(prs, "R19")
    box = shape_named(slide_of(prs, "content_insights"), "Rectangle 14")
    assert box.top + box.height <= d.SLIDE_H - BOTTOM_MARGIN
    assert engine.MSG_OVERFLOW in [a.message for a in engine.attention_items(analyze(prs), "CS", OWNER)]


# ---------------------------------------------------------------- R20 always replaces

def test_r20_replaces_someone_elses_details():
    prs = d.new_prs()
    d.thank_you(prs, ("Old Person", "Old Title", "Old Team"))
    run_rule(prs, "R20", jira=OWNER)
    assert paragraphs(shape_named(slide_of(prs, "thank_you"), "TextBox 7")) == [
        "Pat Lee", "CXM", "plee@madisonlogic.com"]


# ---------------------------------------------------------------- R24 pacing split

def test_r24_splits_a_long_pacing_table_with_header_on_every_part_and_total_last():
    from core.ppra import engine
    out, log, _att = engine.format(d.to_bytes(d.long_pacing_deck()), "CS + Display", OWNER)
    deck = analyze(Presentation(io.BytesIO(out)))
    parts = deck.of_kind("pacing")
    n = len(parts)
    assert n >= 2
    assert [p.title for p in parts] == [f"[{i}/{n}] Campaign Overview and Pacing" for i in range(1, n + 1)]
    assert [p.index for p in parts] == list(range(parts[0].index, parts[0].index + n))  # right after it
    body = []
    for i, info in enumerate(parts):
        frame, rows = info.tables[0]
        assert norm_text(rows[0][0]) == "Campaign Name"
        totals = [r for r in rows[1:] if norm_text(r[0]) == "Total"]
        assert len(totals) == (1 if i == n - 1 else 0)
        body += [norm_text(r[0]) for r in rows[1:] if norm_text(r[0]) != "Total"]
        assert frame.top == parts[0].tables[0][0].top  # same top on every part
    assert len(body) == 18 and len(set(body)) == 18
    assert any(line.startswith("R24:") for line in log)
    # Idempotent: an already split deck is not split again, and gets no second Total row.
    again, log2, _att = engine.format(out, "CS + Display", OWNER)
    assert log2 == []
    assert len(Presentation(io.BytesIO(again)).slides) == len(deck.slides)


def test_r24_summary_and_short_tables_are_left_alone():
    from core.ppra.rules import split_part
    findings = RULES_BY_ID["R24"].applies(analyze(d.long_pacing_deck()), "CS + Display", {})
    assert findings[0].summary.startswith("Split the pacing table across ")
    assert findings[0].summary.endswith(" slides (it runs off the slide)")
    assert RULES_BY_ID["R24"].applies(analyze(d.cs_deck()), "CS", {}) == []
    assert split_part("[2/3] Campaign Overview") == (2, 3) and split_part("Campaign Overview") is None


def test_duplicate_slide_copies_charts_and_pictures(tmp_path):
    from PIL import Image
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    from core.ppra.slides import duplicate_slide
    prs = d.new_prs()
    slide = d.blank(prs, "Source")
    d.chart(slide, XL_CHART_TYPE.PIE, ["A", "B"], {"S": (1, 2)}, name="Chart 9")
    img = tmp_path / "px.png"
    Image.new("RGB", (4, 4), (200, 0, 0)).save(img)
    slide.shapes.add_picture(str(img), Inches(1), Inches(1))
    d.blank(prs, "After")
    duplicate_slide(prs, 0, 1)
    prs = reopen(prs)
    assert [s.title for s in analyze(prs).slides] == ["Source", "Source", "After"]
    chart_parts = [next(r.target_part for r in s.part.rels.values() if r.reltype.endswith("/chart"))
                   for s in list(prs.slides)[:2]]
    assert chart_parts[0] is not chart_parts[1]  # each slide owns its chart
    copy_slide = list(prs.slides)[1]
    assert [s.name for s in copy_slide.shapes] == [s.name for s in list(prs.slides)[0].shapes]
    assert list(copy_slide.shapes)[1].chart.plots[0].categories[0] == "A"
    assert any(r.reltype.endswith("/image") for r in copy_slide.part.rels.values())
