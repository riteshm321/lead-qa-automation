import openpyxl
import pandas as pd

from core.fuzzy_match import compare_columns, apply_match_column_colors, MATCH_COLUMN


def test_compare_columns_scores_an_exact_match_as_100():
    df = pd.DataFrame([{"Job Title": "VP Engineering", "LinkedIn Job Title": "VP Engineering"}])

    result = compare_columns(df, "Job Title", "LinkedIn Job Title")

    assert result.loc[0, MATCH_COLUMN] == "100%"


def test_compare_columns_is_case_and_order_insensitive():
    df = pd.DataFrame([{"Job Title": "vp engineering", "LinkedIn Job Title": "Engineering, VP"}])

    result = compare_columns(df, "Job Title", "LinkedIn Job Title")

    # token_set_ratio ignores case and word order -- both sides share the
    # same words, just reordered/punctuated differently, so this should
    # score very high even though it isn't a literal exact match.
    score = int(result.loc[0, MATCH_COLUMN].rstrip("%"))
    assert score >= 90


def test_compare_columns_scores_a_weak_match_low():
    df = pd.DataFrame([{"Job Title": "Chief Financial Officer", "LinkedIn Job Title": "Warehouse Associate"}])

    result = compare_columns(df, "Job Title", "LinkedIn Job Title")

    score = int(result.loc[0, MATCH_COLUMN].rstrip("%"))
    assert score < 50


def test_compare_columns_treats_blank_as_empty_string_not_a_crash():
    df = pd.DataFrame([{"Job Title": None, "LinkedIn Job Title": "VP Engineering"}])

    result = compare_columns(df, "Job Title", "LinkedIn Job Title")

    assert result.loc[0, MATCH_COLUMN] == "0%"


def test_compare_columns_does_not_mutate_the_original_dataframe():
    df = pd.DataFrame([{"Job Title": "A", "LinkedIn Job Title": "A"}])

    compare_columns(df, "Job Title", "LinkedIn Job Title")

    assert MATCH_COLUMN not in df.columns


def test_apply_match_column_colors_bands_by_score(tmp_path):
    path = str(tmp_path / "output.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", MATCH_COLUMN])
    ws.append(["Green", "95%"])
    ws.append(["Yellow", "80%"])
    ws.append(["Red", "40%"])
    wb.save(path)

    apply_match_column_colors(path)

    wb2 = openpyxl.load_workbook(path)
    ws2 = wb2.active
    assert ws2.cell(row=2, column=2).fill.fgColor.rgb == "00C6EFCE"
    assert ws2.cell(row=3, column=2).fill.fgColor.rgb == "00FFEB9C"
    assert ws2.cell(row=4, column=2).fill.fgColor.rgb == "00FFC7CE"
