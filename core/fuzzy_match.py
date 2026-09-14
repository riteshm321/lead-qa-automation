import pandas as pd
from openpyxl.styles import PatternFill
from rapidfuzz import fuzz

# Score bands match the standalone Fuzzy Match tool's own coloring exactly,
# so a file processed by either one looks the same to the client reviewer.
_GREEN_FILL = PatternFill(start_color="C6EFCE", fill_type="solid")
_YELLOW_FILL = PatternFill(start_color="FFEB9C", fill_type="solid")
_RED_FILL = PatternFill(start_color="FFC7CE", fill_type="solid")

MATCH_COLUMN = "Match %"


def _clean_text(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip().lower()


def compare_columns(df: pd.DataFrame, column_a: str, column_b: str) -> pd.DataFrame:
    """Adds a "Match %" column scoring how similar column_a and column_b
    are on each row (e.g. the leadfile's own Job Title vs. a LinkedIn Job
    Title column), via rapidfuzz's token_set_ratio -- case/whitespace-
    insensitive and order-independent, so "VP Engineering" and
    "Engineering, VP" score as a strong match. A blank value on either
    side is treated as an empty string, not skipped, so it still gets a
    real (low) score rather than leaving the row unscored.
    """
    result = df.copy()
    result[MATCH_COLUMN] = result.apply(
        lambda row: f"{round(fuzz.token_set_ratio(_clean_text(row[column_a]), _clean_text(row[column_b])))}%",
        axis=1,
    )
    return result


def _fill_for_score(score: int) -> PatternFill:
    if score >= 90:
        return _GREEN_FILL
    if score >= 75:
        return _YELLOW_FILL
    return _RED_FILL


def apply_match_column_colors(path: str, sheet_name: str | None = None) -> None:
    """Color-codes an already-saved workbook's "Match %" column: green
    (>=90%), yellow (>=75%), red (below) -- run this after writing the
    DataFrame from compare_columns to a real .xlsx file, since colors
    can't be set on a DataFrame/Series itself.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        if MATCH_COLUMN not in headers:
            return
        col_idx = headers.index(MATCH_COLUMN) + 1
        for row in ws.iter_rows(min_row=2):
            cell = row[col_idx - 1]
            try:
                score = int(str(cell.value).rstrip("%"))
            except (TypeError, ValueError):
                continue
            cell.fill = _fill_for_score(score)
        wb.save(path)
    finally:
        wb.close()
