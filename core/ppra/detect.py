"""Slide and report-type detection for raw PPRA decks.

The platform-generated decks are not built from title placeholders: slide
titles are usually plain text boxes named "Title N" (sometimes inside a
group), and their text carries non-breaking spaces, vertical tabs and en
dashes. Everything here works on that normalised title text plus secondary
signals (table header rows, chart titles, slide layout names).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

# Channels a report can cover. Every report type is a set of these, and the
# rules ask questions of that set ("has CS?", "any impression channel?")
# instead of comparing report-type strings.
CS = "CS"
DISPLAY = "Display"
CTV = "CTV"
AUDIO = "Audio"
LINKEDIN = "LinkedIn"
CHANNELS = (CS, DISPLAY, CTV, AUDIO, LINKEDIN)

IMPS = "Imps"
LEADS = "Leads"
# The pacing unit each channel is bought in. All impression channels sum
# together as one "Imps" total.
UNIT_BY_CHANNEL = {
    CS: LEADS,
    DISPLAY: IMPS,
    CTV: IMPS,
    AUDIO: IMPS,
    LINKEDIN: IMPS,  # LinkedIn unit unconfirmed
}

REPORT_CS = "CS"
REPORT_DISPLAY = "Display"
REPORT_CTV = "CTV"
REPORT_AUDIO = "Audio"
REPORT_LINKEDIN = "LinkedIn"
REPORT_CTV_DISPLAY = "CTV + Display"
REPORT_CS_DISPLAY = "CS + Display"
REPORT_CS_AUDIO = "CS + Audio"
REPORT_CS_DISPLAY_CTV = "CS + Display + CTV"
REPORT_CS_DISPLAY_AUDIO_CTV = "CS + Display + Audio + CTV"

# Selectbox order on the page.
REPORT_TYPE_CHANNELS: dict[str, frozenset[str]] = {
    REPORT_CS: frozenset({CS}),
    REPORT_DISPLAY: frozenset({DISPLAY}),
    REPORT_CTV: frozenset({CTV}),
    REPORT_AUDIO: frozenset({AUDIO}),
    REPORT_LINKEDIN: frozenset({LINKEDIN}),
    REPORT_CTV_DISPLAY: frozenset({CTV, DISPLAY}),
    REPORT_CS_DISPLAY: frozenset({CS, DISPLAY}),
    REPORT_CS_AUDIO: frozenset({CS, AUDIO}),
    REPORT_CS_DISPLAY_CTV: frozenset({CS, DISPLAY, CTV}),
    REPORT_CS_DISPLAY_AUDIO_CTV: frozenset({CS, DISPLAY, AUDIO, CTV}),
}
REPORT_TYPES = tuple(REPORT_TYPE_CHANNELS)

DEFAULT_CHART_CATEGORIES = ["1st Qtr", "2nd Qtr", "3rd Qtr", "4th Qtr"]

_TITLE_NAME_RE = re.compile(r"^Title \d+$")
_WS_RE = re.compile(r"\s+")
_CS_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9])CS(?![A-Za-z0-9])")
_CHANNEL_RES = {
    CS: re.compile(r"content syndication|(?<![a-z])lead(?:s|gen|\s*gen(?:eration)?)?(?![a-z])", re.IGNORECASE),
    DISPLAY: re.compile(r"display|video|banner", re.IGNORECASE),
    CTV: re.compile(r"connected tv|(?<![a-z])(?:ctv|ott)(?![a-z])", re.IGNORECASE),
    AUDIO: re.compile(r"(?<![a-z])audio(?![a-z])|podcast|spotify", re.IGNORECASE),
    LINKEDIN: re.compile(r"linked\s?in", re.IGNORECASE),
}


def norm_text(value: str | None) -> str:
    """Collapse the platform's \\xa0 / \\x0b / newlines into single spaces."""
    if not value:
        return ""
    value = value.replace("\xa0", " ").replace("\x0b", " ")
    return _WS_RE.sub(" ", value).strip()


def channels_of(text: str) -> frozenset[str]:
    """Channel keywords in a campaign / section / flight-date / Jira label."""
    found = {channel for channel, pattern in _CHANNEL_RES.items() if pattern.search(text)}
    if _CS_TOKEN_RE.search(text):
        found.add(CS)
    return frozenset(found)


def report_channels(report_type: str | None) -> frozenset[str]:
    """The channels a report type covers (empty when unknown)."""
    return REPORT_TYPE_CHANNELS.get(report_type or "", frozenset())


def has_cs(report_type: str | None) -> bool:
    return CS in report_channels(report_type)


def has_impressions(report_type: str | None) -> bool:
    """True when the report includes any impression channel (Display, CTV, Audio, LinkedIn)."""
    return any(UNIT_BY_CHANNEL[c] == IMPS for c in report_channels(report_type))


def single_unit(channels) -> str | None:
    """'Imps' or 'Leads' when every channel in the set uses the same unit."""
    units = {UNIT_BY_CHANNEL[c] for c in channels}
    return units.pop() if len(units) == 1 else None


def report_type_for(channels) -> str | None:
    """The listed report type with exactly these channels, or None."""
    channels = frozenset(channels)
    return next((name for name, chans in REPORT_TYPE_CHANNELS.items() if chans == channels), None)


def iter_shapes(shapes):
    """Every shape on a slide, descending into groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)
        else:
            yield shape


def find_shape(slide, shape_id: int):
    for shape in iter_shapes(slide.shapes):
        if shape.shape_id == shape_id:
            return shape
    return None


def is_top_level(shape) -> bool:
    """True when the shape sits directly on the slide (not inside a group),
    so its offsets are slide coordinates."""
    return shape._element.getparent().tag == qn("p:spTree")


def table_rows(table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def chart_data(chart) -> tuple[list[str], list[tuple[str, list]]]:
    """(category labels, [(series name, values), ...]) - empty on any
    chart python-pptx can't read."""
    try:
        categories = [str(c) for c in chart.plots[0].categories]
    except Exception:  # noqa: BLE001 - unusual chart XML; treat as no data
        categories = []
    series = []
    try:
        for plot in chart.plots:
            for s in plot.series:
                series.append((str(s.name), list(s.values)))
    except Exception:  # noqa: BLE001
        pass
    return categories, series


def chart_title(chart) -> str:
    try:
        if chart.has_title:
            return norm_text(chart.chart_title.text_frame.text)
    except Exception:  # noqa: BLE001
        pass
    return ""


def chart_has_data(chart) -> bool:
    categories, series = chart_data(chart)
    if [norm_text(c) for c in categories] == DEFAULT_CHART_CATEGORIES:
        return False
    for _name, values in series:
        if any(v not in (None, 0, 0.0) for v in values):
            return True
    return False


@dataclass
class SlideInfo:
    index: int  # 0-based position in the deck
    kind: str
    title: str
    layout: str
    section_name: str = ""
    section_channels: frozenset = frozenset()
    title_shape_id: int | None = None
    tables: list = field(default_factory=list, repr=False)  # [(shape, rows)]
    charts: list = field(default_factory=list, repr=False)  # [shape]
    texts: list = field(default_factory=list, repr=False)  # [(shape, normalised text)]
    slide: object = field(default=None, repr=False)

    @property
    def number(self) -> int:
        return self.index + 1

    @property
    def label(self) -> str:
        return self.title or self.kind.replace("_", " ").title()

    def text_shape(self, prefix: str):
        """First text shape whose normalised text starts with prefix (case-insensitive)."""
        prefix = prefix.lower()
        for shape, text in self.texts:
            if text.lower().startswith(prefix):
                return shape
        return None

    def has_text(self, needle: str) -> bool:
        needle = needle.lower()
        return any(needle in text.lower() for _shape, text in self.texts)


@dataclass
class DeckInfo:
    slides: list[SlideInfo]
    flight_lines: list[str]
    flight_channels: frozenset
    detected_report_type: str | None
    slide_width: int
    slide_height: int
    prs: object = field(default=None, repr=False)

    def of_kind(self, *kinds: str) -> list[SlideInfo]:
        return [s for s in self.slides if s.kind in kinds]

    def next_slide(self, info: SlideInfo) -> SlideInfo | None:
        nxt = info.index + 1
        return self.slides[nxt] if nxt < len(self.slides) else None


def _slide_title(slide) -> tuple[str, int | None]:
    for shape in iter_shapes(slide.shapes):
        if _TITLE_NAME_RE.match(shape.name or "") and shape.has_text_frame:
            text = norm_text(shape.text_frame.text)
            if text:
                return text, shape.shape_id
    return "", None


def _classify(title: str, layout: str, tables, charts, texts) -> str:
    t = title.lower()
    lay = layout.lower()
    headers = [[norm_text(c) for c in rows[0]] for _shape, rows in tables if rows]
    first_cells = {h[0].lower() for h in headers if h}
    all_texts = [text for _shape, text in texts]
    lowered = [x.lower() for x in all_texts]

    if lay.startswith("title slide") or any(x.startswith("flight dates") for x in lowered):
        return "title"
    if lay.startswith("team slide") or any(x.startswith("thank you") for x in lowered):
        return "thank_you"
    if lay.startswith("section breaker"):
        return "section"
    if lay.startswith("agenda") or "agenda" in lowered:
        return "agenda"
    for h in headers:
        low = [c.lower() for c in h]
        if "campaign name" in low and "units goal" in low:
            return "pacing"
    if "creative sets" in t:
        return "creative_sets"
    if "halo effect" in t or "multi-channel engagement" in t:
        return "halo"
    if t.startswith("custom question"):
        return "custom_question"
    if t.startswith("audience reach"):
        return "audience_reach"
    if t.startswith("audience insights"):
        return "audience_insights"
    if t.startswith("content insights") or "asset preview" in first_cells:
        return "content_insights"
    if t.startswith("industry insights"):
        return "industry_insights"
    if t.startswith("country insights") or "country" in first_cells:
        return "country_insights"
    if "top accounts" in t:
        header_cells = " | ".join(c.lower() for h in headers for c in h)
        if re.search(r"(?<![a-z])leads(?![a-z])", header_cells):
            return "top_accounts_cs"
        if "impressions" in header_cells:
            return "top_accounts_display"
        if "content syndication" in t:
            return "top_accounts_cs"
        title_channels = channels_of(title)
        return "top_accounts_display" if title_channels and CS not in title_channels else "top_accounts"
    if "connected tv performance" in t or any(x.startswith("top assets by ctv impression") for x in lowered):
        return "ctv_performance"
    if ("display performance" in t or "display data" in t
            or any(x.startswith("top assets by display impression") or x.startswith("top assets by impression")
                   for x in lowered)):
        return "display_performance"
    if "program performance" in t:
        return "program_performance"
    if t.startswith("campaign highlights") or "what topics are your accounts researching?" in first_cells:
        return "highlights"
    if "key call outs" in t:
        return "key_call_outs"
    if "recommended actions" in t:
        return "recommended_actions"
    return "other"


def _flight_lines(slides: list[SlideInfo]) -> list[str]:
    for info in slides:
        for shape, _text in info.texts:
            raw = shape.text_frame.text
            if "flight dates" not in raw.lower():
                continue
            after = re.split(r"flight dates\s*:?", raw, maxsplit=1, flags=re.IGNORECASE)[-1]
            lines = [norm_text(x) for x in re.split(r"[\n\x0b\r]", after)]
            return [x for x in lines if x]
    return []


def report_type_from_flight_lines(lines: list[str]) -> str | None:
    """The report type whose channels match the title slide's Flight Dates
    lines exactly. None when a line names no known channel (in a multi-line
    list) or the combination isn't one of REPORT_TYPES."""
    if not lines:
        return None
    per_line = [channels_of(line) for line in lines]
    if any(not ch for ch in per_line) and len(lines) > 1:
        return None
    return report_type_for(frozenset().union(*per_line))


def map_report_type(ppra_format_value: str | None, products: list[str] | None) -> str | None:
    """Report type from the ticket's PPRA Report Format + Products fields.

    Channel keywords in either field (Lead Gen / Content Syndication -> CS,
    Display / Video, CTV / Connected TV, Audio, LinkedIn) are collected and
    the listed report type with exactly that channel set is returned; any
    other combination -> None (the caller falls back to deck detection).
    """
    found: set[str] = set()
    for part in [ppra_format_value or ""] + list(products or []):
        found |= channels_of(part)
    return report_type_for(found)


def parse_slides(prs) -> list[SlideInfo]:
    slides = []
    section_name, section_channels = "", frozenset()
    for index, slide in enumerate(prs.slides):
        title, title_id = _slide_title(slide)
        tables, charts, texts = [], [], []
        for shape in iter_shapes(slide.shapes):
            if getattr(shape, "has_table", False) and shape.has_table:
                tables.append((shape, table_rows(shape.table)))
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                charts.append(shape)
            if shape.has_text_frame:
                text = norm_text(shape.text_frame.text)
                if text:
                    texts.append((shape, text))
        layout = slide.slide_layout.name or ""
        kind = _classify(title, layout, tables, charts, texts)
        if kind == "section":
            section_name = " ".join(text for _s, text in texts)
            section_channels = channels_of(section_name)
        slides.append(SlideInfo(
            index=index, kind=kind, title=title, layout=layout,
            section_name=section_name if kind != "section" else "",
            section_channels=section_channels if kind != "section" else frozenset(),
            title_shape_id=title_id, tables=tables, charts=charts, texts=texts, slide=slide,
        ))
    return slides


def analyze(prs) -> DeckInfo:
    slides = parse_slides(prs)
    lines = _flight_lines([s for s in slides if s.kind == "title"] or slides[:1])
    channels = frozenset().union(*[channels_of(x) for x in lines])
    return DeckInfo(
        slides=slides, flight_lines=lines, flight_channels=channels,
        detected_report_type=report_type_from_flight_lines(lines),
        slide_width=int(prs.slide_width), slide_height=int(prs.slide_height), prs=prs,
    )
