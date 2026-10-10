"""One function pair per formatting rule (R1-R26 of the PPRA spec).

Every rule has applies(deck, report_type, jira_info) -> [Finding] and
apply(prs, finding). applies() only reports work that is still needed, so a
rule run on an already-formatted deck finds nothing (idempotency): no second
Total row, no "26 Leads Leads", no double hyperlink, and so on.

apply() re-reads the current slide state rather than trusting values captured
at scan time, so rules can run one after another on the same Presentation.
Slide deletions (R4-R8, R22) only mark slides; engine.py removes them last.
"""
from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from typing import Callable

from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.text.text import _Run

from core.ppra.detect import (
    AUDIO, CHANNELS, CS, CTV, DISPLAY, IMPS, LEADS, LINKEDIN, DeckInfo, SlideInfo, analyze, chart_data,
    chart_has_data, chart_title, channels_of, find_shape, has_impressions, is_top_level, iter_shapes, norm_text,
    pacing_channels_of, parse_slides, report_channels, single_unit, table_rows,
)

# Wording for every generated Key Takeaway. Edit here to change what the tool
# writes; {n} is the number of list items, the *_one variants are used when
# only one item exists.
TAKEAWAY_TEMPLATES = {
    # Audience Reach: keeps the original sentence start up to "engaging with".
    "audience_reach": "engaging with the following top {n} intent topics:",
    "audience_reach_one": "engaging with the following top intent topic:",
    "content_insights": ("Assets that generated the strongest engagement are closely aligned with the following "
                         "top {n} job titles:"),
    "content_insights_one": ("Assets that generated the strongest engagement are closely aligned with the "
                             "following top job title:"),
    "audience_insights": ("{pct}% of accounts are trending on 5+ intent topics, showing strong research "
                          "activity around your campaign topics."),
    "audience_insights_fallback": "The following are the top {n} intent topics in your campaign:",
    "audience_insights_fallback_one": "The following is the top intent topic in your campaign:",
    # Country Insights: {p1}% / {n1} / {total} / {p2}% / {p3}% / {top_pct}% are
    # written as bold PPRA-blue runs, the rest in the box's body style.
    "country_three": ("{c1} led lead delivery with {p1}% ({n1} of {total} leads), followed by {c2} at {p2}% "
                      "and {c3} at {p3}%. Together, the top three markets contributed {top_pct}% of all leads."),
    "country_two": ("{c1} led lead delivery with {p1}% ({n1} of {total} leads), followed by {c2} at {p2}%. "
                    "Together, these two markets contributed {top_pct}% of all leads."),
    # Only used when R5 (one-country slide deletion) is switched off.
    "country_one": "{c1} delivered all leads: {p1}% ({n1} of {total} leads).",
    # Halo Effect: the number before "Average number of website visits per
    # account who engaged multi-channel, compared to single-channel".
    "halo_ratio": "{ratio}x",
    # Halo Effect with no multi-channel accounts (the platform's division by zero).
    "halo_single_channel": "All engaged accounts engaged through a single channel during this campaign.",
}

TOP_N = 3
TABLE_WIDTH = 11430000  # 12.5 in
PACING_NAME_DELTA = -600000
PACING_UNITS_DELTA = 200000
TOTAL_FILL = "EBEEF5"
FONT = "Montserrat"

_DELETABLE_WHEN_EMPTY = (
    "top_accounts_cs", "top_accounts_display", "top_accounts", "content_insights", "industry_insights",
    "country_insights", "display_performance", "ctv_performance",
)
_WIDE_TABLE_KINDS = ("top_accounts_cs", "top_accounts_display", "top_accounts", "display_performance",
                     "ctv_performance")
_TAKEAWAY_KINDS = ("audience_reach", "audience_insights", "content_insights", "country_insights", "halo")
_COUNTRIES_WITH_THE = {
    "united states", "united kingdom", "netherlands", "philippines", "united arab emirates",
    "czech republic", "dominican republic", "bahamas", "maldives", "gambia", "central african republic",
}
_FLOAT_NOISE_RE = re.compile(r"(\d+\.\d{3,})%")
_NUMBER_RE = re.compile(r"^-?[\d,]+(?:\.\d+)?$")
_SUFFIXED_RE = re.compile(r"^(-?[\d,]+(?:\.\d+)?)\s+(Imps|Leads)$", re.IGNORECASE)
_MONEY_RE = re.compile(r"^(-?)\$?\s*(-?[\d,]+(?:\.\d+)?)$")
_MLP_PREFIX = "click to view on ml platform"
_MLP_URL_RE = re.compile(r"https?://platform\.madisonlogic\.com/\S*")
_PCT_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*%")


@dataclass
class Finding:
    rule_id: str
    slide_index: int
    description: str
    before: str = ""
    after: str = ""
    data: dict = field(default_factory=dict)
    # One short plain-language line for the review list (no slide number -
    # the page prefixes it). Falls back to the description.
    summary: str = ""

    @property
    def short(self) -> str:
        return self.summary or self.description


@dataclass(frozen=True)
class Rule:
    id: str
    label: str
    finder: Callable
    applier: Callable
    deletes: bool = False

    def applies(self, deck: DeckInfo, report_type: str | None, jira_info: dict | None = None) -> list[Finding]:
        return self.finder(deck, report_type, jira_info or {})

    def apply(self, prs, finding: Finding) -> None:
        self.applier(prs, finding)


# ---------------------------------------------------------------- helpers

def _shape(prs, finding: Finding):
    shape = find_shape(prs.slides[finding.slide_index], finding.data["shape_id"])
    if shape is None:
        raise LookupError(f"Shape {finding.data['shape_id']} not found on slide {finding.slide_index + 1}")
    return shape


def _p_text(p_el) -> str:
    return "".join(t.text or "" for t in p_el.iter(qn("a:t")))


def _runs(p_el) -> list:
    return p_el.findall(qn("a:r"))


def _run_text(r_el) -> str:
    t = r_el.find(qn("a:t"))
    return (t.text or "") if t is not None else ""


def _set_run_text(r_el, text: str) -> None:
    t = r_el.find(qn("a:t"))
    if t is None:
        t = OxmlElement("a:t")
        r_el.append(t)
    t.text = text


def _rpr_from(source) -> object:
    """A clean a:rPr copied from a run's rPr or a paragraph's endParaRPr,
    without hyperlinks or spell-check noise."""
    rpr = OxmlElement("a:rPr")
    if source is not None:
        for key, value in source.attrib.items():
            if key not in ("err", "dirty"):
                rpr.set(key, value)
        for child in source:
            if child.tag not in (qn("a:hlinkClick"), qn("a:hlinkMouseOver")):
                rpr.append(copy.deepcopy(child))
    return rpr


def _first_rpr(p_el):
    for r in _runs(p_el):
        if _run_text(r).strip():
            return r.find(qn("a:rPr"))
    runs = _runs(p_el)
    if runs:
        return runs[0].find(qn("a:rPr"))
    return p_el.find(qn("a:endParaRPr"))


def _make_run(text: str, rpr_source=None, bold: bool | None = None, font: str | None = None):
    r = OxmlElement("a:r")
    rpr = _rpr_from(rpr_source)
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if font:
        _set_latin(rpr, font)
    r.append(rpr)
    t = OxmlElement("a:t")
    t.text = text
    r.append(t)
    return r


_RPR_AFTER_LATIN = ("a:ea", "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


def _set_latin(rpr, typeface: str) -> None:
    latin = rpr.find(qn("a:latin"))
    if latin is None:
        latin = OxmlElement("a:latin")
        anchor = next((rpr.find(qn(tag)) for tag in _RPR_AFTER_LATIN if rpr.find(qn(tag)) is not None), None)
        if anchor is not None:
            anchor.addprevious(latin)
        else:
            rpr.append(latin)
    latin.set("typeface", typeface)


def _make_paragraph(runs=(), ppr=None, end_rpr_source=None):
    p = OxmlElement("a:p")
    if ppr is not None:
        p.append(ppr)
    for r in runs:
        p.append(r)
    if end_rpr_source is not None:
        end = _rpr_from(end_rpr_source)
        end.tag = qn("a:endParaRPr")
        p.append(end)
    return p


def _numbered_ppr():
    ppr = OxmlElement("a:pPr")
    ppr.set("marL", "342900")
    ppr.set("indent", "-342900")
    bu = OxmlElement("a:buAutoNum")
    bu.set("type", "arabicPeriod")
    ppr.append(bu)
    return ppr


def _bullet_ppr():
    ppr = OxmlElement("a:pPr")
    ppr.set("marL", "228600")
    ppr.set("indent", "-228600")
    font = OxmlElement("a:buFont")
    font.set("typeface", "Arial")
    ppr.append(font)
    char = OxmlElement("a:buChar")
    char.set("char", "\u2022")
    ppr.append(char)
    return ppr


def _replace_paragraph_runs(p_el, runs) -> None:
    for child in list(p_el):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p_el.remove(child)
    end = p_el.find(qn("a:endParaRPr"))
    for r in runs:
        if end is not None:
            end.addprevious(r)
        else:
            p_el.append(r)


def _insert_list_after(p_el, items: list[str], rpr_source) -> None:
    """Blank line then a 1./2./3. list after p_el, styled like rpr_source."""
    blank = _make_paragraph(end_rpr_source=rpr_source)
    p_el.addnext(blank)
    anchor = blank
    for item in items:
        para = _make_paragraph([_make_run(item, rpr_source)], ppr=_numbered_ppr())
        anchor.addnext(para)
        anchor = para


def _top_n(items: list[str]) -> list[str]:
    return [x for x in items if x][:TOP_N]


def _parse_number(text: str) -> float | None:
    text = norm_text(text)
    m = _SUFFIXED_RE.match(text)
    if m:
        text = m.group(1)
    if not _NUMBER_RE.match(text):
        return None
    return float(text.replace(",", ""))


def _parse_money(text: str) -> float | None:
    m = _MONEY_RE.match(norm_text(text))
    if not m:
        return None
    value = float(m.group(2).replace(",", ""))
    return -abs(value) if m.group(1) == "-" else value


def _fmt_int_or_float(value: float) -> str:
    if abs(value - round(value)) < 1e-9:
        return f"{int(round(value)):,}"
    return f"{value:,.2f}"


def _fmt_money(value: float) -> str:
    sign = "-" if value < 0 else ""
    return f"{sign}${_fmt_int_or_float(abs(value))}"


def _round_half_up(value: float, places: str = "1") -> Decimal:
    return Decimal(str(value)).quantize(Decimal(places), rounding=ROUND_HALF_UP)


def _scale_widths(widths: list[int], total: int) -> list[int]:
    current = sum(widths)
    if current <= 0:
        return widths
    scaled = [int(round(w * total / current)) for w in widths]
    scaled[scaled.index(max(scaled))] += total - sum(scaled)
    return scaled


def _fit_frame(slide_width: int, frame, table, widths: list[int] | None = None) -> None:
    widths = _scale_widths(widths or [c.width for c in table.columns], TABLE_WIDTH)
    for column, width in zip(table.columns, widths):
        column.width = width
    frame.width = TABLE_WIDTH
    frame.left = (slide_width - TABLE_WIDTH) // 2


def _frame_fits(slide_width: int, frame, table) -> bool:
    centred = (slide_width - TABLE_WIDTH) // 2
    return (abs(int(frame.width) - TABLE_WIDTH) <= 1000 and abs(int(frame.left) - centred) <= 1000
            and abs(sum(c.width for c in table.columns) - TABLE_WIDTH) <= 1000)


def _slide_label(info: SlideInfo) -> str:
    return f"slide {info.number} ({info.label})"


# ---------------------------------------------------------------- pacing (R1-R3)

_PACING_HEADERS = {
    "campaign": "campaign name", "bg": "budget goal", "bd": "budget delivered", "br": "budget remaining",
    "ug": "units goal", "ud": "units delivered", "ur": "units remaining", "pct": "% delivered",
}
_UNIT_KEYS = ("ug", "ud", "ur")
_BUDGET_KEYS = ("bg", "bd", "br")


def pacing_tables(deck: DeckInfo):
    for info in deck.of_kind("pacing"):
        for shape, rows in info.tables:
            if not rows:
                continue
            header = [norm_text(c).lower() for c in rows[0]]
            cols = {key: header.index(name) for key, name in _PACING_HEADERS.items() if name in header}
            if "campaign" in cols and "ug" in cols:
                yield info, shape, rows, cols


def _is_total_row(row: list[str]) -> bool:
    return bool(row) and norm_text(row[0]).lower() == "total"


def data_rows(rows) -> list[int]:
    return [i for i in range(1, len(rows)) if not _is_total_row(rows[i]) and any(norm_text(c) for c in rows[i])]


def row_unit(campaign_name: str, row: list[str], cols: dict, deck: DeckInfo, report_type: str | None) -> str | None:
    """'Imps' or 'Leads' for one pacing row: an existing suffix, then campaign
    name keywords (display/video/banner, ctv/ott, audio/podcast, linkedin ->
    Imps; cs/content syndication/contentsynd/content/lead -> Leads), then the
    title slide's flight-date channels, then the report type - each only when
    it points to exactly one unit. None means mixed and undetermined (flagged
    for attention)."""
    for key in _UNIT_KEYS:
        if key in cols:
            m = _SUFFIXED_RE.match(norm_text(row[cols[key]]))
            if m:
                return IMPS if m.group(2).lower() == "imps" else LEADS
    for channels in (pacing_channels_of(campaign_name), deck.flight_channels, report_channels(report_type)):
        unit = single_unit(channels)
        if unit:
            return unit
    return None


def _find_r1(deck, report_type, jira_info):
    findings = []
    for info, shape, rows, cols in pacing_tables(deck):
        cells, sample = [], None
        for r in data_rows(rows):
            unit = row_unit(rows[r][cols["campaign"]], rows[r], cols, deck, report_type)
            if unit is None:
                continue
            for key in _UNIT_KEYS:
                if key not in cols:
                    continue
                text = norm_text(rows[r][cols[key]])
                if _NUMBER_RE.match(text):
                    cells.append([r, cols[key], unit])
                    sample = sample or (text, f"{text} {unit}")
        if cells:
            units = sorted({c[2] for c in cells})
            findings.append(Finding(
                "R1", info.index, f"Add {' / '.join(units)} after {len(cells)} unit values on {_slide_label(info)}",
                before=sample[0], after=sample[1], data={"shape_id": shape.shape_id, "cells": cells},
                summary=f"Add {' / '.join(units)} after {len(cells)} unit values"))
    return findings


def _apply_r1(prs, finding):
    table = _shape(prs, finding).table
    for r, c, unit in finding.data["cells"]:
        cell = table.cell(r, c)
        if not _NUMBER_RE.match(norm_text(cell.text)):
            continue
        p = cell.text_frame.paragraphs[0]
        target = None
        for run in p.runs:
            if run.text.strip():
                target = run
        if target is None:
            target = p.add_run()
        target.text = target.text.rstrip() + f" {unit}"
        target.font.name = FONT


def _pacing_totals(rows, cols, deck, report_type) -> list[list[str]]:
    data = data_rows(rows)
    out = [[""] for _ in rows[0]]
    out[cols["campaign"]] = ["Total"]
    for key in _BUDGET_KEYS:
        if key in cols:
            values = [_parse_money(rows[r][cols[key]]) for r in data]
            if values and all(v is not None for v in values):
                out[cols[key]] = [_fmt_money(sum(values))]
    units = {}
    order = []
    for r in data:
        unit = row_unit(rows[r][cols["campaign"]], rows[r], cols, deck, report_type) or ""
        if unit not in units:
            units[unit] = {k: 0.0 for k in _UNIT_KEYS}
            order.append(unit)
        for key in _UNIT_KEYS:
            if key in cols:
                units[unit][key] += _parse_number(rows[r][cols[key]]) or 0.0
    for key in _UNIT_KEYS:
        if key in cols:
            out[cols[key]] = [f"{_fmt_int_or_float(units[u][key])} {u}".strip() for u in order]
    # % Delivered per unit type (Imps and Leads each against their own goal),
    # one line per type like the units cells - never a blend of the two.
    if "pct" in cols and "ud" in cols:
        pcts = [f"{_round_half_up(units[u]['ud'] * 100 / units[u]['ug'])}%"
                for u in order if units[u].get("ug")]
        if pcts:
            out[cols["pct"]] = pcts
    return out


def _find_r2(deck, report_type, jira_info):
    findings = []
    for info, shape, rows, cols in pacing_tables(deck):
        if any(_is_total_row(r) for r in rows[1:]) or not data_rows(rows):
            continue
        part = split_part(info.title)
        if part and part[0] < part[1]:
            continue  # an earlier part of a split table: its Total row is on the last part
        totals = _pacing_totals(rows, cols, deck, report_type)
        headline = [" + ".join(x for x in totals[cols["ug"]] if x)]
        if "bg" in cols:
            headline.append(totals[cols["bg"]][0])
        findings.append(Finding(
            "R2", info.index, f"Add a bold Total row to the pacing table on {_slide_label(info)}",
            before="(no Total row)", after=" | ".join(" / ".join(c) for c in totals),
            data={"shape_id": shape.shape_id, "report_type": report_type},
            summary="Add Total row: " + ", ".join(x for x in headline if x)))
    return findings


def _total_cell(lines: list[str], ppr_source, size: str | None):
    tc = OxmlElement("a:tc")
    body = OxmlElement("a:txBody")
    body.append(OxmlElement("a:bodyPr"))
    body.append(OxmlElement("a:lstStyle"))
    for line in lines or [""]:
        ppr = copy.deepcopy(ppr_source) if ppr_source is not None else OxmlElement("a:pPr")
        ppr.set("algn", "ctr")
        runs = []
        if line:
            rpr_src = OxmlElement("a:rPr")
            rpr_src.set("lang", "en-US")
            if size:
                rpr_src.set("sz", size)
            runs = [_make_run(line, rpr_src, bold=True, font=FONT)]
        body.append(_make_paragraph(runs, ppr=ppr))
    tc.append(body)
    tcpr = OxmlElement("a:tcPr")
    for attr, value in (("marL", "6350"), ("marR", "6350"), ("marT", "6350"), ("marB", "0"), ("anchor", "ctr")):
        tcpr.set(attr, value)
    fill = OxmlElement("a:solidFill")
    clr = OxmlElement("a:srgbClr")
    clr.set("val", TOTAL_FILL)
    fill.append(clr)
    tcpr.append(fill)
    tc.append(tcpr)
    return tc


def _apply_r2(prs, finding):
    frame = _shape(prs, finding)
    table = frame.table
    rows = table_rows(table)
    if any(_is_total_row(r) for r in rows[1:]):
        return
    header = [norm_text(c).lower() for c in rows[0]]
    cols = {key: header.index(name) for key, name in _PACING_HEADERS.items() if name in header}
    deck = analyze(prs)
    totals = _pacing_totals(rows, cols, deck, finding.data.get("report_type"))
    tbl = table._tbl
    trs = tbl.tr_lst
    data_tr = trs[data_rows(rows)[0]]
    ppr_source = data_tr.find(".//" + qn("a:pPr"))
    size_rpr = data_tr.find(".//" + qn("a:rPr"))
    size = size_rpr.get("sz") if size_rpr is not None else None
    tr = OxmlElement("a:tr")
    height = int(trs[0].get("h", "0")) or 370840
    tr.set("h", str(height))
    for lines in totals:
        tr.append(_total_cell(lines, ppr_source, size))
    trs[-1].addnext(tr)
    frame.height = int(frame.height) + height


def _pacing_rebalanced(widths: list[int], cols: dict) -> bool:
    ref = cols.get("bd", cols.get("bg"))
    if ref is None or "ug" not in cols:
        return True
    return widths[cols["ug"]] > widths[ref] * 1.05


def _find_r3(deck, report_type, jira_info):
    findings = []
    for info, shape, rows, cols in pacing_tables(deck):
        widths = [c.width for c in shape.table.columns]
        needs_rebalance = not _pacing_rebalanced(widths, cols)
        if not needs_rebalance and _frame_fits(deck.slide_width, shape, shape.table):
            continue
        what = "narrow Campaign Name, widen Units columns, " if needs_rebalance else ""
        findings.append(Finding(
            "R3", info.index, f"Pacing table on {_slide_label(info)}: {what}12.5 in wide, centred",
            before=f"width {int(shape.width) / 914400:.2f} in, left {int(shape.left) / 914400:.2f} in",
            after=f"width 12.50 in, left {(deck.slide_width - TABLE_WIDTH) / 2 / 914400:.2f} in",
            data={"shape_id": shape.shape_id},
            summary=("Narrow Campaign Name, widen the Units columns, " if needs_rebalance else "Make the table ")
            + "12.5 in wide and centred"))
    return findings


def _apply_r3(prs, finding):
    frame = _shape(prs, finding)
    table = frame.table
    header = [norm_text(c).lower() for c in table_rows(table)[0]]
    cols = {key: header.index(name) for key, name in _PACING_HEADERS.items() if name in header}
    _fit_frame(int(prs.slide_width), frame, table, _planned_pacing_widths([c.width for c in table.columns], cols))


def _planned_pacing_widths(widths: list[int], cols: dict) -> list[int]:
    """Column widths once R3 has run: Campaign Name narrowed and the Units
    columns widened (once), then scaled to 12.5 in."""
    widths = list(widths)
    if not _pacing_rebalanced(widths, cols) and widths[cols["campaign"]] > 2 * abs(PACING_NAME_DELTA):
        widths[cols["campaign"]] += PACING_NAME_DELTA
        for key in _UNIT_KEYS:
            if key in cols:
                widths[cols[key]] += PACING_UNITS_DELTA
    return _scale_widths(widths, TABLE_WIDTH)


# ---------------------------------------------------------------- R24 pacing split

_SPLIT_RE = re.compile(r"\s*\[(\d+)/(\d+)\]\s*$")
_OLD_SPLIT_RE = re.compile(r"^\s*\[(\d+)/(\d+)\]\s*")  # earlier leading form
_DEFAULT_TABLE_PT = 12.0
_FOOTER_GAP = 45720  # 0.05 in clear of the footer bar
_NO_FOOTER_MARGIN = 457200  # 0.5 in when the layout has no footer shape


def split_part(title: str) -> tuple[int, int] | None:
    """(i, n) when a pacing title already carries an "[i/n]" split marker:
    trailing ("Title [1/2]"), or the earlier leading form ("[1/2] Title")."""
    m = _SPLIT_RE.search(title or "") or _OLD_SPLIT_RE.match(title or "")
    return (int(m.group(1)), int(m.group(2))) if m else None


def _cell_font_pt(tc) -> float:
    sizes = [int(r.get("sz")) for r in tc.iter(qn("a:rPr")) if r.get("sz")]
    sizes += [int(r.get("sz")) for r in tc.iter(qn("a:endParaRPr")) if r.get("sz")]
    return max(sizes) / 100 if sizes else _DEFAULT_TABLE_PT


def _cell_margin(tc, name: str, default: int) -> int:
    tcpr = tc.find(qn("a:tcPr"))
    value = tcpr.get(name) if tcpr is not None else None
    return int(value) if value is not None else default


def _cell_height(tc, lines_text: list[str], width: int) -> int:
    """Conservative rendered height (EMU) of one cell's text."""
    size = _cell_font_pt(tc)
    usable = width - _cell_margin(tc, "marL", 91440) - _cell_margin(tc, "marR", 91440)
    per_line = max(1, int(usable / (size * CHAR_WIDTH_EM * _EMU_PER_PT)))
    lines = sum(wrapped_lines(text, per_line) for text in lines_text or [""])
    return (int(lines * size * _LINE_SPACING * _EMU_PER_PT)
            + _cell_margin(tc, "marT", 45720) + _cell_margin(tc, "marB", 45720))


def _row_height(tr, widths: list[int], suffix_cols=()) -> int:
    height = int(tr.get("h", "0"))
    for c, (tc, width) in enumerate(zip(tr.findall(qn("a:tc")), widths)):
        texts = [_p_text(p) for p in tc.iter(qn("a:p"))]
        if c in suffix_cols:  # " Imps" / " Leads" still to come from R1
            texts = [t + " Leads" if _NUMBER_RE.match(norm_text(t)) else t for t in texts]
        height = max(height, _cell_height(tc, texts, width))
    return height


def pacing_bottom_limit(slide, slide_height: int) -> int:
    """Lowest EMU a table may reach: just above the layout's footer bar (the
    highest non-placeholder layout/master shape in the bottom half)."""
    tops = []
    layout = slide.slide_layout
    for shape in list(layout.shapes) + list(layout.slide_master.shapes):
        if shape.is_placeholder or shape.top is None:
            continue
        if slide_height // 2 < int(shape.top) < slide_height:
            tops.append(int(shape.top))
    return (min(tops) if tops else slide_height - _NO_FOOTER_MARGIN) - _FOOTER_GAP


def _split_plan(info: SlideInfo, shape, rows, cols, slide_height: int) -> list[list[int]]:
    """Row indices (header excluded) for each part, Total row last on the
    last part; one part means the table fits. Planned like the deck will be
    after R1-R3: final widths, unit suffixes and a Total row."""
    tbl = shape.table._tbl
    trs = tbl.tr_lst
    widths = _planned_pacing_widths([c.width for c in shape.table.columns], cols)
    suffix_cols = {cols[k] for k in _UNIT_KEYS if k in cols}
    heights = [_row_height(tr, widths, suffix_cols if r else ()) for r, tr in enumerate(trs)]
    body = [r for r in range(1, len(rows)) if not _is_total_row(rows[r])]
    total = next((r for r in range(1, len(rows)) if _is_total_row(rows[r])), None)
    total_height = heights[total] if total is not None else 0
    if total is None and data_rows(rows):
        # R2 will add one: header-row height, one line per unit type.
        total_height = max(heights[0], _cell_height(trs[1].findall(qn("a:tc"))[0], ["Total", "Total"], widths[0]))
    available = pacing_bottom_limit(info.slide, slide_height) - int(shape.top)
    parts, current, used = [], [], heights[0]
    for r in body:
        if current and used + heights[r] > available:
            parts.append(current)
            current, used = [], heights[0]
        current.append(r)
        used += heights[r]
    if current and used + total_height > available and len(current) > 1:
        parts.append(current[:-1])
        current = current[-1:]
    parts.append(current + ([total] if total is not None else []))
    return parts


def _find_r24(deck, report_type, jira_info):
    findings = []
    for info, shape, rows, cols in pacing_tables(deck):
        if split_part(info.title) or not is_top_level(shape):
            continue
        parts = _split_plan(info, shape, rows, cols, deck.slide_height)
        if len(parts) < 2:
            continue
        findings.append(Finding(
            "R24", info.index, f"Split the pacing table on {_slide_label(info)} across {len(parts)} slides",
            before=info.title, after=f"{info.title} [1/{len(parts)}]",
            data={"shape_id": shape.shape_id, "parts": len(parts)},
            summary=f"Split the pacing table across {len(parts)} slides (it runs off the slide)"))
    return findings


def _mark_title(slide, title_shape_id: int | None, i: int, n: int) -> None:
    """Put " [i/n]" after the title, dropping any marker it already has
    (trailing, or the earlier leading form)."""
    shape = find_shape(slide, title_shape_id) if title_shape_id else None
    if shape is None:
        return
    for p in shape.text_frame.paragraphs:
        runs = [r for r in p.runs if r.text.strip()]
        if runs:
            runs[0].text = _OLD_SPLIT_RE.sub("", runs[0].text)
            runs[-1].text = _SPLIT_RE.sub("", runs[-1].text.rstrip()) + f" [{i}/{n}]"
            return


def _keep_rows(frame, keep: set[int]) -> None:
    tbl = frame.table._tbl
    for r, tr in reversed(list(enumerate(tbl.tr_lst))):
        if r not in keep:
            tbl.remove(tr)
    frame.height = sum(int(tr.get("h", "0")) for tr in tbl.tr_lst)


def _apply_r24(prs, finding):
    from core.ppra.slides import duplicate_slide  # local: slides.py imports nothing from here

    deck = analyze(prs)
    info = deck.slides[finding.slide_index]
    match = next(((shape, rows, cols) for table_info, shape, rows, cols in pacing_tables(deck)
                  if table_info.index == info.index and shape.shape_id == finding.data["shape_id"]), None)
    if match is None or split_part(info.title):
        return
    shape, rows, cols = match
    parts = _split_plan(info, shape, rows, cols, deck.slide_height)
    n = len(parts)
    if n < 2:
        return
    slides = [info.slide] + [duplicate_slide(prs, info.index, info.index + k) for k in range(1, n)]
    for k, (slide, part) in enumerate(zip(slides, parts), start=1):
        _keep_rows(find_shape(slide, finding.data["shape_id"]), {0, *part})
        _mark_title(slide, info.title_shape_id, k, n)


# ---------------------------------------------------------------- deletions (R4-R8)

def _slide_has_data(info: SlideInfo) -> bool:
    containers = 0
    for _shape_, rows in info.tables:
        containers += 1
        if any(any(norm_text(c) for c in row) for row in rows[1:]):
            return True
    for chart_shape in info.charts:
        containers += 1
        if chart_has_data(chart_shape.chart):
            return True
    if info.kind in ("display_performance", "ctv_performance"):
        for _s, text in info.texts:
            low = text.lower()
            if "impressions served" in low and re.search(r"[1-9]", low.split("served", 1)[1]):
                return True
    if containers:
        return False
    placeholders = ("error:division by zero", "avg ctr %")
    return not any(info.has_text(p) for p in placeholders)


def _delete_finding(rule_id: str, info: SlideInfo, why: str) -> Finding:
    return Finding(rule_id, info.index, f"Delete {_slide_label(info)}: {why}",
                   before=info.label, after="(slide removed)", data={}, summary=f"Remove '{info.label}' ({why})")


def _find_r4(deck, report_type, jira_info):
    return [_delete_finding("R4", info, "no data in its tables/charts")
            for info in deck.of_kind(*_DELETABLE_WHEN_EMPTY) if not _slide_has_data(info)]


def _country_rows(info: SlideInfo) -> list[list[str]]:
    for _shape_, rows in info.tables:
        if rows and norm_text(rows[0][0]).lower() == "country":
            return [r for r in rows[1:] if norm_text(r[0])]
    return []


def _find_r5(deck, report_type, jira_info):
    return [_delete_finding("R5", info, f"only one country ({norm_text(_country_rows(info)[0][0])})")
            for info in deck.of_kind("country_insights") if len(_country_rows(info)) == 1]


def _find_r6(deck, report_type, jira_info):
    # A one-channel report has no multi-channel engagement to show; in a
    # multi-channel report the slide stays (R23 fills its takeaway).
    if len(report_channels(report_type)) != 1:
        return []
    return [_delete_finding("R6", info, f"a {report_type} report has only one channel")
            for info in deck.of_kind("halo")]


def _find_r7(deck, report_type, jira_info):
    return [_delete_finding("R7", info, "template example slide")
            for info in deck.of_kind("custom_question") if "example" in info.title.lower()]


def creative_sets_channels(info: SlideInfo) -> frozenset:
    prefix = re.split(r"[\u2013\u2014-]\s*display data", info.title, flags=re.IGNORECASE)[0]
    return channels_of(prefix) or info.section_channels


def _find_r8(deck, report_type, jira_info):
    findings = []
    for info in deck.of_kind("creative_sets"):
        channels = creative_sets_channels(info)
        if channels == {CS} or (not channels and report_channels(report_type) == {CS}):
            findings.append(_delete_finding("R8", info, "Creative Sets slide in a Content Syndication section"))
    return findings


def _apply_delete(prs, finding):
    """Deletions are carried out by engine.delete_slides after every other rule."""


# ---------------------------------------------------------------- R22 channel deletion

_IMPRESSION_CHANNELS = frozenset({DISPLAY, CTV, AUDIO, LINKEDIN})
# Never removed for their channel: the deck frame and the pages every report
# keeps. Halo Effect is R6's call.
_CHANNEL_NEUTRAL_KINDS = ("title", "agenda", "section", "pacing", "recommended_actions", "thank_you", "halo")
# Display-family slides: their impression channel comes from the title or
# section ("CTV Insights" -> CTV), Display when neither names one.
_DISPLAY_FAMILY_KINDS = ("program_performance", "key_call_outs", "creative_sets")


def slide_channels(info: SlideInfo) -> frozenset:
    """The channel(s) a slide reports on, from its kind first and then its
    title / section keywords. Empty means it is not tied to a channel."""
    kind = info.kind
    if kind in _CHANNEL_NEUTRAL_KINDS or info.title.lower().startswith("campaign highlights"):
        return frozenset()
    title_channels = channels_of(info.title)
    if kind == "ctv_performance":
        return frozenset({CTV})
    if kind == "display_performance":
        return frozenset({DISPLAY})
    if kind == "top_accounts_cs":
        return frozenset({CS})
    if kind == "top_accounts_display":
        # "Top Accounts - Display and CTV" serves both; a plain "- Top Accounts" is Display.
        return (title_channels & _IMPRESSION_CHANNELS) or frozenset({DISPLAY})
    if kind in _DISPLAY_FAMILY_KINDS:
        return ((title_channels | info.section_channels) & _IMPRESSION_CHANNELS) or frozenset({DISPLAY})
    # Insight slides (Audience Reach, Content / Industry / Country Insights,
    # Top Accounts, Custom Question, ...) belong to their section's channel.
    return title_channels or info.section_channels


def _channel_names(channels) -> str:
    return " / ".join(c for c in CHANNELS if c in channels)


def _section_ranges(deck: DeckInfo):
    """(section breaker, [slides up to the next breaker]) pairs."""
    breakers = [s.index for s in deck.slides if s.kind == "section"]
    for i, start in enumerate(breakers):
        end = breakers[i + 1] if i + 1 < len(breakers) else len(deck.slides)
        yield deck.slides[start], deck.slides[start + 1:end]


def _find_r22(deck, report_type, jira_info):
    wanted = report_channels(report_type)
    if not wanted:
        return []
    findings = {}
    for info in deck.slides:
        channels = slide_channels(info)
        if channels and not channels & wanted:
            findings[info.index] = Finding(
                "R22", info.index, f"Delete {_slide_label(info)}: {_channel_names(channels)} slide, not in a "
                f"{report_type} report", before=info.label, after="(slide removed)", data={},
                summary=f"Remove {_channel_names(channels)} slide (not in a {report_type} report)")
    # A section breaker goes too once every slide under it is removed.
    for breaker, members in _section_ranges(deck):
        if members and all(m.index in findings for m in members):
            channels = channels_of(" ".join(text for _s, text in breaker.texts))
            what = f"{_channel_names(channels)} section breaker" if channels else "Section breaker"
            findings[breaker.index] = Finding(
                "R22", breaker.index, f"Delete {_slide_label(breaker)}: every slide in its section is removed",
                before=breaker.label, after="(slide removed)", data={},
                summary=f"Remove {what} (not in a {report_type} report)")
    return [findings[i] for i in sorted(findings)]


def _skip_off_channel(finder):
    """The other deleting rules leave alone slides R22 already removes, so a
    slide is listed once, with its channel reason."""
    def wrapped(deck, report_type, jira_info):
        off = {f.slide_index for f in _find_r22(deck, report_type, jira_info)}
        return [f for f in finder(deck, report_type, jira_info) if f.slide_index not in off]
    return wrapped


# ---------------------------------------------------------------- R9 CTV column

def _find_r9(deck, report_type, jira_info):
    # Any report with an impression channel (Display, CTV, Audio, LinkedIn);
    # never CS-only. The title loses "and CTV" only with the removed column.
    if not has_impressions(report_type):
        return []
    findings = []
    for info in deck.of_kind("top_accounts_display", "top_accounts"):
        for shape, rows in info.tables:
            if not rows:
                continue
            header = [norm_text(c).lower() for c in rows[0]]
            if "ctv impressions" not in header:
                continue
            col = header.index("ctv impressions")
            data = [r for r in rows[1:] if any(norm_text(c) for c in r)]
            if data and all((_parse_number(r[col]) or 0) == 0 for r in data):
                new_title = re.sub(r"\s+and CTV\b", "", info.title)
                findings.append(Finding(
                    "R9", info.index, f"Remove the all-zero CTV Impressions column on {_slide_label(info)}",
                    before=info.title, after=new_title,
                    data={"shape_id": shape.shape_id, "title_shape_id": info.title_shape_id},
                    summary="Drop the all-zero CTV Impressions column"
                    + (" and retitle" if new_title != info.title else "")))
    return findings


def _apply_r9(prs, finding):
    frame = _shape(prs, finding)
    table = frame.table
    header = [norm_text(c).lower() for c in table_rows(table)[0]]
    if "ctv impressions" not in header:
        return
    col = header.index("ctv impressions")
    old_total = sum(c.width for c in table.columns)
    grid = table._tbl.tblGrid
    grid.remove(grid.findall(qn("a:gridCol"))[col])
    for tr in table._tbl.tr_lst:
        tr.remove(tr.findall(qn("a:tc"))[col])
    widths = _scale_widths([c.width for c in table.columns], old_total)
    for column, width in zip(table.columns, widths):
        column.width = width
    title_id = finding.data.get("title_shape_id")
    title = find_shape(prs.slides[finding.slide_index], title_id) if title_id else None
    if title is not None:
        _retitle_without_ctv(title)


def _retitle_without_ctv(title_shape) -> None:
    pattern = re.compile(r"\s+and CTV\b")
    for p in title_shape.text_frame.paragraphs:
        runs = p.runs
        if any(pattern.search(r.text) for r in runs):
            for r in runs:
                r.text = pattern.sub("", r.text)
            return
        joined = "".join(r.text for r in runs)
        if pattern.search(joined) and runs:
            runs[0].text = pattern.sub("", joined)
            for r in runs[1:]:
                r.text = ""
            return


# ---------------------------------------------------------------- R21 wide tables

def _find_r21(deck, report_type, jira_info):
    findings = []
    for info in deck.of_kind(*_WIDE_TABLE_KINDS):
        for shape, rows in info.tables:
            if not is_top_level(shape) or _frame_fits(deck.slide_width, shape, shape.table):
                continue
            findings.append(Finding(
                "R21", info.index, f"Make the '{norm_text(rows[0][0]) if rows else ''}' table on "
                f"{_slide_label(info)} 12.5 in wide and centred",
                before=f"width {int(shape.width) / 914400:.2f} in", after="width 12.50 in, centred",
                data={"shape_id": shape.shape_id},
                summary=f"Make the '{norm_text(rows[0][0]) if rows else ''}' table 12.5 in wide and centred"))
    return findings


def _apply_r21(prs, finding):
    frame = _shape(prs, finding)
    _fit_frame(int(prs.slide_width), frame, frame.table)


# ---------------------------------------------------------------- R10 MLP link

def _mlp_paragraphs(info: SlideInfo):
    for shape, text in info.texts:
        if _MLP_PREFIX not in text.lower():
            continue
        for p_index, p in enumerate(shape.text_frame.paragraphs):
            joined = "".join(r.text for r in p.runs)
            if _MLP_PREFIX not in joined.lower():
                continue
            m = _MLP_URL_RE.search(joined)
            if m:
                yield shape, p_index, p, joined, m


def _url_span(m) -> tuple[int, int, str]:
    url = m.group(0).rstrip(".,;)")
    return m.start(), m.start() + len(url), url


def _runs_in_span(p, start: int, end: int):
    pos = 0
    for run in p.runs:
        length = len(run.text)
        if pos < end and pos + length > start and run.text.strip():
            yield run
        pos += length


def _find_r10(deck, report_type, jira_info):
    findings = []
    for info in deck.slides:
        for shape, p_index, p, joined, m in _mlp_paragraphs(info):
            start, end, url = _url_span(m)
            if all(r.hyperlink.address == url for r in _runs_in_span(p, start, end)):
                continue
            findings.append(Finding(
                "R10", info.index, f"Hyperlink the ML Platform URL on {_slide_label(info)}",
                before=url, after=f"{url} (clickable)",
                data={"shape_id": shape.shape_id, "paragraph": p_index},
                summary="Make the ML Platform link clickable"))
    return findings


def _split_run(r_el, at: int):
    """Split one a:r at character offset `at`; returns the new right-hand run."""
    text = _run_text(r_el)
    right = copy.deepcopy(r_el)
    _set_run_text(r_el, text[:at])
    _set_run_text(right, text[at:])
    r_el.addnext(right)
    return right


def _isolate_span(p_el, start: int, end: int) -> list:
    pos = 0
    for r in list(_runs(p_el)):
        text = _run_text(r)
        r_start, r_end = pos, pos + len(text)
        pos = r_end
        if r_end <= start or r_start >= end:
            continue
        if r_start < start:
            r = _split_run(r, start - r_start)
            r_start = start
        if r_end > end:
            _split_run(r, end - r_start)
    out, pos = [], 0
    for r in _runs(p_el):
        length = len(_run_text(r))
        if pos >= start and pos + length <= end and length:
            out.append(r)
        pos += length
    return out


def _apply_r10(prs, finding):
    shape = _shape(prs, finding)
    p = shape.text_frame.paragraphs[finding.data["paragraph"]]
    joined = "".join(r.text for r in p.runs)
    m = _MLP_URL_RE.search(joined)
    if not m:
        return
    start, end, url = _url_span(m)
    for r_el in _isolate_span(p._p, start, end):
        _Run(r_el, p).hyperlink.address = url


# ---------------------------------------------------------------- R11 float noise

def _all_text_frames(slide):
    for shape in iter_shapes(slide.shapes):
        if shape.has_text_frame:
            yield shape.text_frame
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def _fix_float(text: str) -> str:
    return _FLOAT_NOISE_RE.sub(lambda m: f"{_round_half_up(float(m.group(1)), '0.01')}%", text)


def _find_r11(deck, report_type, jira_info):
    findings = []
    for info in deck.slides:
        hits = []
        for tf in _all_text_frames(info.slide):
            for p in tf.paragraphs:
                for run in p.runs:
                    for m in _FLOAT_NOISE_RE.finditer(run.text):
                        hits.append(m.group(0))
        if hits:
            findings.append(Finding(
                "R11", info.index, f"Round {len(hits)} long decimal value(s) to 2 places on {_slide_label(info)}",
                before=hits[0], after=_fix_float(hits[0]), data={},
                summary=f"Round {len(hits)} long decimal{'s' if len(hits) != 1 else ''} to 2 places "
                f"({hits[0]} -> {_fix_float(hits[0])})"))
    return findings


def _apply_r11(prs, finding):
    for tf in _all_text_frames(prs.slides[finding.slide_index]):
        for p in tf.paragraphs:
            for run in p.runs:
                if _FLOAT_NOISE_RE.search(run.text):
                    run.text = _fix_float(run.text)


# ---------------------------------------------------------------- R12 empty KPI lines

def _empty_kpi_paragraphs(tf) -> list[int]:
    paras = [p._p for p in tf.paragraphs]
    remove = []
    for i in range(len(paras) - 1):
        stat = norm_text(_p_text(paras[i])).replace(" ", "")
        if stat in ("+", "0+") and norm_text(_p_text(paras[i + 1])).lower().startswith("site visits"):
            if i > 0 and not norm_text(_p_text(paras[i - 1])) and (i - 1) not in remove:
                remove.append(i - 1)
            remove += [i, i + 1]
    return remove


def _find_r12(deck, report_type, jira_info):
    findings = []
    for info in deck.slides:
        for shape, text in info.texts:
            if _empty_kpi_paragraphs(shape.text_frame):
                findings.append(Finding(
                    "R12", info.index, f"Remove the empty 'site visits generated' line on {_slide_label(info)}",
                    before="+ site visits generated", after="(removed)", data={"shape_id": shape.shape_id},
                    summary="Remove the empty '+ site visits generated' line"))
    return findings


def _apply_r12(prs, finding):
    tf = _shape(prs, finding).text_frame
    paras = [p._p for p in tf.paragraphs]
    for i in sorted(_empty_kpi_paragraphs(tf), reverse=True):
        if len(tf.paragraphs) > 1:
            paras[i].getparent().remove(paras[i])


# ---------------------------------------------------------------- R13 company sizes

def _fix_size_band(text: str) -> str:
    if not re.search(r"\d{4,}", text):
        return text
    lone = re.fullmatch(r"\s*\d+\s*", text) is not None
    new = re.sub(r"\d{4,}", lambda m: f"{int(m.group()):,}", text)
    if lone and "+" not in new:
        new = re.sub(r"([\d,]+)", r"\1+", new, count=1)
    return new


def _find_r13(deck, report_type, jira_info):
    findings = []
    for info in deck.of_kind("highlights"):
        shape = info.text_shape("company size")
        if shape is None:
            continue
        hits = [r.text for p in shape.text_frame.paragraphs for r in p.runs if _fix_size_band(r.text) != r.text]
        if hits:
            findings.append(Finding(
                "R13", info.index, f"Format company sizes on {_slide_label(info)}",
                before=", ".join(h.strip() for h in hits), after=", ".join(_fix_size_band(h).strip() for h in hits),
                data={"shape_id": shape.shape_id}, summary="Add thousands separators to company sizes"))
    return findings


def _apply_r13(prs, finding):
    for p in _shape(prs, finding).text_frame.paragraphs:
        for run in p.runs:
            run.text = _fix_size_band(run.text)


# ---------------------------------------------------------------- R26 duplicate highlight items

def _item_key(text: str) -> str:
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    return re.sub(r"\s+", " ", text).strip().lower()


def _duplicate_items(shape) -> list:
    """Sub-list paragraphs (level > 0) repeating an earlier item under the
    same parent bullet; headings (level 0) are never picked."""
    dupes, seen = [], {}
    for p in shape.text_frame.paragraphs:
        level, key = p.level, _item_key(p.text)
        for deeper in [lv for lv in seen if lv > level]:
            del seen[deeper]
        if not key:
            continue
        if level == 0:
            seen.clear()
            continue
        keys = seen.setdefault(level, set())
        if key in keys:
            dupes.append(p)
        else:
            keys.add(key)
    return dupes


def _find_r26(deck, report_type, jira_info):
    findings = []
    for info in deck.of_kind("highlights"):
        for shape in info.slide.shapes:
            if not shape.has_text_frame:
                continue
            dupes = _duplicate_items(shape)
            if dupes:
                n = len(dupes)
                items = "item" if n == 1 else "items"
                findings.append(Finding(
                    "R26", info.index, f"Remove {n} duplicate {items} on {_slide_label(info)}",
                    before=", ".join(norm_text(p.text) for p in dupes), after="(removed)",
                    data={"shape_id": shape.shape_id},
                    summary=f"Remove {n} duplicate {items} on Campaign Highlights"))
    return findings


def _apply_r26(prs, finding):
    for p in _duplicate_items(_shape(prs, finding)):
        p._p.getparent().remove(p._p)


# ---------------------------------------------------------------- R14 chart labels

_DLBLS_TXPR_BEFORE = ("c:dLblPos", "c:showLegendKey", "c:showVal", "c:showCatName", "c:showSerName",
                      "c:showPercent", "c:showBubbleSize", "c:separator", "c:showLeaderLines",
                      "c:leaderLines", "c:extLst")


def _has_labels(chart_space) -> bool:
    return any(d.find(qn("c:delete")) is None for d in chart_space.iter(qn("c:dLbls")))


def _label_def_rprs(chart_space, create: bool) -> list:
    out = []
    for dlbls in chart_space.iter(qn("c:dLbls")):
        if dlbls.find(qn("c:delete")) is not None:
            continue
        txprs = [dlbls.find(qn("c:txPr"))] + [d.find(qn("c:txPr")) for d in dlbls.findall(qn("c:dLbl"))]
        txprs = [t for t in txprs if t is not None]
        if not txprs and create:
            txpr = OxmlElement("c:txPr")
            txpr.append(OxmlElement("a:bodyPr"))
            txpr.append(OxmlElement("a:lstStyle"))
            p = OxmlElement("a:p")
            ppr = OxmlElement("a:pPr")
            ppr.append(OxmlElement("a:defRPr"))
            p.append(ppr)
            txpr.append(p)
            anchor = next((dlbls.find(qn(t)) for t in _DLBLS_TXPR_BEFORE if dlbls.find(qn(t)) is not None), None)
            if anchor is not None:
                anchor.addprevious(txpr)
            else:
                dlbls.append(txpr)
            txprs = [txpr]
        for txpr in txprs:
            for defrpr in txpr.iter(qn("a:defRPr")):
                out.append(defrpr)
    return out


def _chart_targets(deck: DeckInfo):
    for info in deck.of_kind("country_insights"):
        for shape in info.charts:
            if shape.chart.chart_type in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED):
                yield info, shape, "country"
    for info in deck.of_kind("program_performance"):
        for shape in info.charts:
            if shape.chart.chart_type in (XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.DOUGHNUT_EXPLODED):
                yield info, shape, "doughnut"


def _find_r14(deck, report_type, jira_info):
    findings = []
    for info, shape, target in _chart_targets(deck):
        if not _has_labels(shape.chart._chartSpace):
            continue
        rprs = _label_def_rprs(shape.chart._chartSpace, create=False)
        if target == "country":
            needed = not rprs or any(r.get("b") != "1" for r in rprs)
            after = "bold, +1 pt"
        else:
            needed = not rprs or any(r.get("b") != "1" or r.get("sz") != "1200" for r in rprs)
            after = "bold, 12 pt"
        if needed:
            findings.append(Finding(
                "R14", info.index, f"Make the {'pie' if target == 'country' else 'doughnut'} chart labels "
                f"on {_slide_label(info)} {after}", before="regular labels", after=after,
                data={"shape_id": shape.shape_id, "target": target},
                summary=f"Make the {'pie' if target == 'country' else 'doughnut'} chart labels {after}"))
    return findings


def _apply_r14(prs, finding):
    chart_space = _shape(prs, finding).chart._chartSpace
    for rpr in _label_def_rprs(chart_space, create=True):
        if finding.data["target"] == "country":
            if rpr.get("b") == "1":
                continue
            rpr.set("b", "1")
            size = int(rpr.get("sz") or 900)
            rpr.set("sz", str(size + 100))
        else:
            rpr.set("b", "1")
            rpr.set("sz", "1200")


# ---------------------------------------------------------------- takeaways (R15-R18)

def _takeaway_shape(info: SlideInfo):
    return info.text_shape("key takeaways")


def _sorted_categories(chart) -> list[str]:
    categories, series = chart_data(chart)
    if not series:
        return categories
    values = series[0][1]
    pairs = [(c, v if v is not None else float("-inf")) for c, v in zip(categories, values)]
    pairs.sort(key=lambda cv: cv[1], reverse=True)
    return [norm_text(c) for c, v in pairs if v != float("-inf")]


def campaign_topics(info: SlideInfo | None) -> list[str]:
    """Top intent topics from an Audience Insights slide's 'Top Intent
    Topics in your Campaign' chart (Chart 10)."""
    if info is None:
        return []
    chosen = None
    for shape in info.charts:
        if "top intent topics in your campaign" in chart_title(shape.chart).lower():
            chosen = shape
            break
    if chosen is None:
        chosen = next((s for s in info.charts if s.name == "Chart 10"), None)
    if chosen is None or not chart_has_data(chosen.chart):
        return []
    return _top_n(_sorted_categories(chosen.chart))


def job_titles(info: SlideInfo | None) -> list[str]:
    if info is None:
        return []
    for _shape_, rows in info.tables:
        if rows and norm_text(rows[0][0]).lower() == "job title":
            return _top_n([norm_text(r[0]) for r in rows[1:]])
    return []


def _template(key: str, n: int, **kw) -> str:
    if n == 1 and f"{key}_one" in TAKEAWAY_TEMPLATES:
        return TAKEAWAY_TEMPLATES[f"{key}_one"].format(n=n, **kw)
    return TAKEAWAY_TEMPLATES[key].format(n=n, **kw)


def _paragraph_index(tf, predicate) -> int | None:
    for i, p in enumerate(tf.paragraphs):
        if predicate(norm_text(_p_text(p._p))):
            return i
    return None


def _already_listed(tf, start: int, items: list[str]) -> bool:
    texts = [norm_text(_p_text(p._p)) for p in tf.paragraphs[start + 1:]]
    return all(item in texts for item in items)


# R15 Audience Reach

def _reach_sentence_index(tf) -> int | None:
    return _paragraph_index(tf, lambda t: "engaging with" in t.lower())


def _find_r15(deck, report_type, jira_info):
    findings = []
    for info in deck.of_kind("audience_reach"):
        shape = _takeaway_shape(info)
        if shape is None:
            continue
        tf = shape.text_frame
        idx = _reach_sentence_index(tf)
        if idx is None:
            continue
        pct_ok = (any(_PCT_RE.match(norm_text(_p_text(p._p))) for p in tf.paragraphs[:idx + 1])
                  or reach_value(info) is not None)  # R25 fills a blank stat first
        topics = campaign_topics(deck.next_slide(info))
        if not topics or not pct_ok:
            continue
        sentence = norm_text(_p_text(tf.paragraphs[idx]._p))
        new = sentence.split("engaging with")[0] + _template("audience_reach", len(topics))
        if sentence == new and _already_listed(tf, idx, topics):
            continue
        findings.append(Finding(
            "R15", info.index, f"Rewrite the Audience Reach takeaway on {_slide_label(info)} with the top "
            f"{len(topics)} intent topic(s)", before=sentence, after=new + " " + "; ".join(
                f"{i}. {t}" for i, t in enumerate(topics, 1)),
            data={"shape_id": shape.shape_id, "topics": topics},
            summary=f"Rewrite the Audience Reach takeaway with the top {len(topics)} intent topics"))
    return findings


def _apply_r15(prs, finding):
    tf = _shape(prs, finding).text_frame
    idx = _reach_sentence_index(tf)
    if idx is None:
        return
    p_el = tf.paragraphs[idx]._p
    topics = finding.data["topics"]
    # Only the text from "engaging with" on is rewritten, so a stat run in
    # the same paragraph (the blue "26% ") keeps its own formatting.
    joined = "".join(_run_text(r) for r in _runs(p_el))
    start = joined.find("engaging with")
    span = _isolate_span(p_el, start, len(joined)) if start >= 0 else []
    if span:
        rpr = span[0].find(qn("a:rPr"))
        _set_run_text(span[0], _template("audience_reach", len(topics)))
        for extra in span[1:]:
            p_el.remove(extra)
    else:
        sentence = norm_text(_p_text(p_el))
        rpr = _first_rpr(p_el)
        new = sentence.split("engaging with")[0] + _template("audience_reach", len(topics))
        _replace_paragraph_runs(p_el, [_make_run(new, rpr)])
    _remove_old_list(tf, idx)
    _insert_list_after(p_el, topics, rpr)


def _remove_old_list(tf, idx: int) -> None:
    """Drop a numbered list (and blank lines) previously written after paragraph idx."""
    paras = [p._p for p in tf.paragraphs]
    for p_el in paras[idx + 1:]:
        ppr = p_el.find(qn("a:pPr"))
        numbered = ppr is not None and ppr.find(qn("a:buAutoNum")) is not None
        if numbered or not norm_text(_p_text(p_el)):
            p_el.getparent().remove(p_el)
        else:
            break


# R16 Content Insights

def _heading_count(tf) -> int:
    """1 when the box starts with its "Key Takeaways:" heading paragraph."""
    paras = tf.paragraphs
    return 1 if paras and norm_text(_p_text(paras[0]._p)).lower().startswith("key takeaways") else 0


def _content_body(tf) -> list[str]:
    return [norm_text(_p_text(p._p)) for p in tf.paragraphs[_heading_count(tf):]]


def _content_target(titles: list[str]) -> list[str]:
    return [_template("content_insights", len(titles)), "", *titles]


def _find_r16(deck, report_type, jira_info):
    findings = []
    for info in deck.of_kind("content_insights"):
        shape = _takeaway_shape(info)
        if shape is None:
            continue
        tf = shape.text_frame
        titles = job_titles(deck.next_slide(info))
        if not titles:
            continue
        body = _content_body(tf)
        while body and not body[-1]:
            body.pop()
        if body == _content_target(titles):
            continue
        findings.append(Finding(
            "R16", info.index, f"Rewrite the Content Insights takeaway on {_slide_label(info)} with the top "
            f"{len(titles)} job title(s)", before=" ".join(x for x in body if x),
            after=_template("content_insights", len(titles)) + " " + "; ".join(
                f"{i}. {t}" for i, t in enumerate(titles, 1)),
            data={"shape_id": shape.shape_id, "titles": titles},
            summary=f"Rewrite the Content Insights takeaway with the top {len(titles)} job titles"))
    return findings


def _apply_r16(prs, finding):
    """The whole takeaway after the heading becomes the fixed sentence, a
    blank line and the numbered job titles, styled like the old first line."""
    tf = _shape(prs, finding).text_frame
    start = _heading_count(tf)
    body = [p._p for p in tf.paragraphs[start:]]
    source = next((p for p in body if norm_text(_p_text(p))), body[0] if body else None)
    if source is None:
        return
    rpr = _first_rpr(source)
    ppr = source.find(qn("a:pPr"))
    sentence_ppr = copy.deepcopy(ppr) if ppr is not None and ppr.find(qn("a:buAutoNum")) is None else None
    titles = finding.data["titles"]
    sentence = _make_paragraph([_make_run(_template("content_insights", len(titles)), rpr)], ppr=sentence_ppr)
    source.addprevious(sentence)
    for p_el in body:
        p_el.getparent().remove(p_el)
    _insert_list_after(sentence, titles, rpr)


# R17 Audience Insights

def _trending_index(tf) -> int | None:
    return _paragraph_index(tf, lambda t: "of accounts are trending on" in t.lower()
                            or "intent topics in your campaign" in t.lower() or "intent topic in your campaign" in t.lower())


def _placeholder_index(tf) -> int | None:
    return _paragraph_index(tf, lambda t: t.startswith("[Topics trending"))


def _trending_pct(text: str) -> str | None:
    m = _PCT_RE.match(text)
    if not m or float(m.group(1)) == 0:
        return None
    return m.group(1)


def _find_r17(deck, report_type, jira_info):
    findings = []
    for info in deck.of_kind("audience_insights"):
        shape = _takeaway_shape(info)
        if shape is None:
            continue
        tf = shape.text_frame
        idx = _trending_index(tf)
        if idx is None:
            continue
        text = norm_text(_p_text(tf.paragraphs[idx]._p))
        has_placeholder = _placeholder_index(tf) is not None
        pct = _trending_pct(text)
        if pct is not None:
            new = TAKEAWAY_TEMPLATES["audience_insights"].format(pct=pct)
            topics = []
        else:
            topics = campaign_topics(info)
            if not topics:
                continue
            new = _template("audience_insights_fallback", len(topics))
        done = text == new and (pct is not None or _already_listed(tf, idx, topics))
        if done and not has_placeholder:
            continue
        after = new + ("" if not topics else " " + "; ".join(f"{i}. {t}" for i, t in enumerate(topics, 1)))
        findings.append(Finding(
            "R17", info.index, f"Rewrite the Audience Insights takeaway on {_slide_label(info)}",
            before=text, after=after, data={"shape_id": shape.shape_id, "pct": pct, "topics": topics},
            summary="Rewrite the Audience Insights takeaway"))
    return findings


def _apply_r17(prs, finding):
    tf = _shape(prs, finding).text_frame
    placeholder = _placeholder_index(tf)
    if placeholder is not None:
        p_el = tf.paragraphs[placeholder]._p
        prev = p_el.getprevious()
        p_el.getparent().remove(p_el)
        # the single blank spacer line directly above the placeholder goes too
        if prev is not None and prev.tag == qn("a:p") and not norm_text(_p_text(prev)):
            prev.getparent().remove(prev)
    idx = _trending_index(tf)
    if idx is None:
        return
    p_el = tf.paragraphs[idx]._p
    text = norm_text(_p_text(p_el))
    pct = finding.data.get("pct")
    if pct is not None:
        new = TAKEAWAY_TEMPLATES["audience_insights"].format(pct=pct)
        if text == new:
            return
        suffix = new[len(text):] if new.startswith(text) else None
        runs = [r for r in _runs(p_el) if _run_text(r)]
        if suffix is not None and runs:
            _set_run_text(runs[-1], _run_text(runs[-1]).rstrip() + suffix)
        else:
            _replace_paragraph_runs(p_el, [_make_run(new, _first_rpr(p_el))])
        return
    topics = finding.data["topics"]
    runs = [r for r in _runs(p_el) if _run_text(r).strip()]
    plain = next((r.find(qn("a:rPr")) for r in runs if r.find(qn("a:rPr")) is not None
                  and r.find(qn("a:rPr")).get("b") != "1"), _first_rpr(p_el))
    _replace_paragraph_runs(p_el, [_make_run(_template("audience_insights_fallback", len(topics)), plain)])
    _remove_old_list(tf, idx)
    _insert_list_after(p_el, topics, plain)


# R18 Country

# Template placeholders written as bold PPRA-blue number runs; the rest of
# the sentence stays in the box's regular body style.
_COUNTRY_NUMBER_KEYS = ("p1", "p2", "p3", "n1", "total", "top_pct")
_TEMPLATE_TOKEN_RE = re.compile(r"(\{\w+\}%?)")
# The PPRA blue as a theme reference: bg2 maps to lt2 (1C6BFF) in the
# platform's "Madison Logic" theme, the colour of the slide titles.
PPRA_BLUE_SCHEME = "bg2"


def _with_the(country: str) -> str:
    return f"the {country}" if country.lower() in _COUNTRIES_WITH_THE else country


def _pct_text(value) -> str:
    text = str(value)
    return text[:-2] if text.endswith(".0") else text


def _country_values(rows: list[list[str]], header: list[str]) -> tuple[str, dict]:
    """(template key, values) for the takeaway sentence."""
    low = [norm_text(h).lower() for h in header]
    leads_col = next((i for i, h in enumerate(low) if "lead" in h), 1)
    pct_col = next((i for i, h in enumerate(low) if "percent" in h or "%" in h), 2)
    data = []
    for r in rows:
        n = _parse_number(r[leads_col]) if leads_col < len(r) else None
        pct = norm_text(r[pct_col]).rstrip("%").strip() if pct_col < len(r) else ""
        data.append((norm_text(r[0]), int(n or 0), pct))
    total = sum(n for _c, n, _p in data)
    top = data[:TOP_N]
    values = {"total": f"{total:,}", "n1": f"{top[0][1]:,}",
              "top_pct": _pct_text(_round_half_up(sum(n for _c, n, _p in top) * 100 / total)) if total else "0"}
    for i, (country, n, pct) in enumerate(top, 1):
        values[f"c{i}"] = _with_the(country)
        values[f"p{i}"] = pct or (_pct_text(_round_half_up(n * 100 / total)) if total else "0")
    key = {1: "country_one", 2: "country_two"}.get(len(top), "country_three")
    return key, values


def _country_parts(rows, header) -> list[tuple[str, bool]]:
    """The sentence as (text, is_number) pieces, numbers as their own runs."""
    key, values = _country_values(rows, header)
    parts = []
    for piece in _TEMPLATE_TOKEN_RE.split(TAKEAWAY_TEMPLATES[key]):
        if not piece:
            continue
        m = re.fullmatch(r"\{(\w+)\}(%?)", piece)
        if m:
            parts.append((values[m.group(1)] + m.group(2), m.group(1) in _COUNTRY_NUMBER_KEYS))
        else:
            parts.append((piece, False))
    if parts and parts[0][0][:1].islower():  # "the United States led ..." starts the sentence
        parts[0] = (parts[0][0][:1].upper() + parts[0][0][1:], parts[0][1])
    merged = []
    for text, is_number in parts:  # neighbouring plain pieces become one run
        if merged and not is_number and not merged[-1][1]:
            merged[-1] = (merged[-1][0] + text, False)
        else:
            merged.append((text, is_number))
    return merged


def _blue_number_run(text: str, rpr_source):
    r = _make_run(text, rpr_source, bold=True)
    rpr = r.find(qn("a:rPr"))
    for fill in rpr.findall(qn("a:solidFill")):
        rpr.remove(fill)
    fill = OxmlElement("a:solidFill")
    clr = OxmlElement("a:schemeClr")
    clr.set("val", PPRA_BLUE_SCHEME)
    fill.append(clr)
    # a:solidFill sits before the effect / font children of a:rPr.
    anchor = next((rpr.find(qn(t)) for t in ("a:effectLst", "a:effectDag", "a:highlight", "a:uLnTx", "a:uLn",
                                             "a:uFillTx", "a:uFill", "a:latin", *_RPR_AFTER_LATIN)
                   if rpr.find(qn(t)) is not None), None)
    if anchor is not None:
        anchor.addprevious(fill)
    else:
        rpr.append(fill)
    return r


def _country_table(info: SlideInfo):
    for _shape_, rows in info.tables:
        if rows and norm_text(rows[0][0]).lower() == "country":
            return rows[0], [r for r in rows[1:] if norm_text(r[0])]
    return None, []


def _find_r18(deck, report_type, jira_info):
    findings = []
    for info in deck.of_kind("country_insights"):
        shape = _takeaway_shape(info)
        header, rows = _country_table(info)
        if shape is None or not rows:
            continue
        idx = _paragraph_index(shape.text_frame, lambda t: t.lower().startswith("to be filled by csm"))
        if idx is None:
            continue
        parts = _country_parts(rows, header)
        findings.append(Finding(
            "R18", info.index, f"Write the Country Insights takeaway on {_slide_label(info)}",
            before=norm_text(_p_text(shape.text_frame.paragraphs[idx]._p)),
            after="".join(text for text, _number in parts), data={"shape_id": shape.shape_id},
            summary="Write the Country Insights takeaway"))
    return findings


def _apply_r18(prs, finding):
    tf = _shape(prs, finding).text_frame
    idx = _paragraph_index(tf, lambda t: t.lower().startswith("to be filled by csm"))
    if idx is None:
        return
    info = parse_slides(prs)[finding.slide_index]
    header, rows = _country_table(info)
    p_el = tf.paragraphs[idx]._p
    rpr = _first_rpr(p_el)
    runs = [_blue_number_run(text, rpr) if number else _make_run(text, rpr, bold=False)
            for text, number in _country_parts(rows, header)]
    _replace_paragraph_runs(p_el, runs)
    old_ppr = p_el.find(qn("a:pPr"))
    if old_ppr is not None:
        p_el.remove(old_ppr)
    p_el.insert(0, _bullet_ppr())


# ---------------------------------------------------------------- R23 Halo Effect takeaway

# The value the platform leaves unfilled: "X", "ERROR:Division by zero", or a bare "%".
_HALO_PLACEHOLDER_RE = re.compile(r"^\s*(?:x\s*%?|%|.*error:.*)\s*$", re.IGNORECASE)
_HALO_EXPLAINER = "average number of website visits"


def halo_placeholder_index(tf) -> int | None:
    """The takeaway paragraph still holding the platform's placeholder value."""
    for i, p in enumerate(tf.paragraphs):
        if i and _HALO_PLACEHOLDER_RE.match(norm_text(_p_text(p._p)) or "-"):
            return i
    return None


def halo_value(info: SlideInfo) -> tuple[str, str | None]:
    """What the Halo Effect takeaway should say, from the slide's own chart.

    The chart has one bar per engagement category ("Single-channel",
    "Two-channel", ...) with an ACCOUNTS series (accounts in that category)
    and a Site Visits series (their website visits). The takeaway's number
    reads "Average number of website visits per account who engaged
    multi-channel, compared to single-channel", so it is the multiplier

        (multi-channel site visits / multi-channel accounts)
        / (single-channel site visits / single-channel accounts)

    written as e.g. "2.5x", where multi-channel is every category that is
    not single-channel. Returns ("ratio", text), ("single", None) when no
    account engaged on more than one channel (the platform's division by
    zero), or ("unknown", None) when the chart can't give the number.
    """
    for chart_shape in info.charts:
        categories, series = chart_data(chart_shape.chart)
        accounts = next((v for name, v in series if "account" in name.lower()), None)
        visits = next((v for name, v in series if "site visit" in name.lower()), None)
        if not categories or accounts is None or visits is None:
            continue
        totals = {"single": [0.0, 0.0], "multi": [0.0, 0.0]}
        for category, n_accounts, n_visits in zip(categories, accounts, visits):
            bucket = totals["single" if "single" in category.lower() else "multi"]
            bucket[0] += n_accounts or 0
            bucket[1] += n_visits or 0
        (single_accounts, single_visits), (multi_accounts, multi_visits) = totals["single"], totals["multi"]
        if not multi_accounts:
            return "single", None
        if not single_accounts or not single_visits:
            return "unknown", None
        ratio = (multi_visits / multi_accounts) / (single_visits / single_accounts)
        return "ratio", TAKEAWAY_TEMPLATES["halo_ratio"].format(ratio=_round_half_up(ratio, "0.1"))
    return "unknown", None


def _find_r23(deck, report_type, jira_info):
    findings = []
    for info in deck.of_kind("halo"):
        shape = _takeaway_shape(info)
        if shape is None:
            continue
        idx = halo_placeholder_index(shape.text_frame)
        if idx is None:
            continue
        outcome, value = halo_value(info)
        if outcome == "unknown":
            continue
        after = value if outcome == "ratio" else TAKEAWAY_TEMPLATES["halo_single_channel"]
        findings.append(Finding(
            "R23", info.index, f"Fill the Halo Effect takeaway on {_slide_label(info)} from its chart",
            before=norm_text(_p_text(shape.text_frame.paragraphs[idx]._p)), after=after,
            data={"shape_id": shape.shape_id, "outcome": outcome, "value": after},
            summary="Fill the Halo Effect takeaway from its chart"))
    return findings


def _apply_r23(prs, finding):
    tf = _shape(prs, finding).text_frame
    idx = halo_placeholder_index(tf)
    if idx is None:
        return
    p_el = tf.paragraphs[idx]._p
    runs = [r for r in _runs(p_el) if _run_text(r).strip()]
    if not runs:
        return
    single = finding.data["outcome"] == "single"
    # A ratio replaces only the placeholder run(s); the fallback sentence
    # replaces the whole line. Either way it keeps the placeholder run's own
    # formatting (same a:rPr).
    targets = runs if single else ([r for r in runs if _HALO_PLACEHOLDER_RE.match(_run_text(r))] or runs)
    trailing = " " if not single and _run_text(targets[-1]).endswith(" ") and targets[-1] is not runs[-1] else ""
    _set_run_text(targets[0], finding.data["value"] + trailing)
    for extra in targets[1:]:
        p_el.remove(extra)
    if single:
        # The sentence explaining the multiplier no longer applies.
        for p in list(tf.paragraphs[idx + 1:]):
            if norm_text(_p_text(p._p)).lower().startswith(_HALO_EXPLAINER):
                p._p.getparent().remove(p._p)


# ---------------------------------------------------------------- R25 Audience Reach stat

# The generator's unfilled big stat: "0", "0%", "%", "X%", "ERROR:...".
_REACH_STAT_RE = re.compile(r"^\s*(?:0(?:\.0+)?\s*%?|%|x\s*%?|.*error:.*)\s*$", re.IGNORECASE)


def reach_stat_run(tf):
    """The big stat run before "of leads" in the Audience Reach takeaway
    (it may share a paragraph with the sentence or sit on its own line)."""
    for p in tf.paragraphs[1:]:
        for r in _runs(p._p):
            text = _run_text(r)
            if "of leads" in text.lower():
                return None
            if text.strip():
                return r
    return None


def reach_stat_missing(tf) -> bool:
    run = reach_stat_run(tf)
    return run is not None and bool(_REACH_STAT_RE.match(_run_text(run)))


def reach_value(info: SlideInfo) -> str | None:
    """The Audience Reach % from the slide's Account Engagement Summary chart:
    Trending accounts Engaged / All Accounts Engaged, as a whole percent
    (e.g. 478 / 784 -> "61%"). This is the user-confirmed stand-in for "% of
    leads from engaged accounts" (the true lead-weighted share is not in the
    deck). None when the chart is missing or no account engaged."""
    for chart_shape in info.charts:
        categories, series = chart_data(chart_shape.chart)
        cats = [norm_text(c).lower() for c in categories]
        if "engaged" not in cats:
            continue
        col = cats.index("engaged")
        named = {name.strip("'\" ").lower(): values for name, values in series}
        all_accounts, trending = named.get("all accounts"), named.get("trending")
        if not all_accounts or not trending or col >= len(all_accounts) or col >= len(trending):
            continue
        if not all_accounts[col]:
            return None
        return f"{_round_half_up((trending[col] or 0) * 100 / all_accounts[col], '1')}%"
    return None


def _find_r25(deck, report_type, jira_info):
    findings = []
    for info in deck.of_kind("audience_reach"):
        shape = _takeaway_shape(info)
        if shape is None or not reach_stat_missing(shape.text_frame):
            continue
        value = reach_value(info)
        if value is None:
            continue
        findings.append(Finding(
            "R25", info.index, f"Fill the Audience Reach % on {_slide_label(info)} from its chart",
            before=_run_text(reach_stat_run(shape.text_frame)).strip(), after=value,
            data={"shape_id": shape.shape_id, "value": value},
            summary=f"Fill the Audience Reach % ({value}, trending / all engaged accounts)"))
    return findings


def _apply_r25(prs, finding):
    tf = _shape(prs, finding).text_frame
    if not reach_stat_missing(tf):
        return  # never overwrite a real value
    run = reach_stat_run(tf)
    _set_run_text(run, finding.data["value"] + " ")


# ---------------------------------------------------------------- R19 auto-fit

_DEFAULT_SIZE_PT = 18.0
CHAR_WIDTH_EM = 0.58  # Montserrat is a wide face
_LINE_SPACING = 1.2
_EMU_PER_PT = 12700


def _paragraph_size_pt(p_el) -> float:
    sizes = [int(r.find(qn("a:rPr")).get("sz")) for r in _runs(p_el)
             if r.find(qn("a:rPr")) is not None and r.find(qn("a:rPr")).get("sz")]
    if not sizes:
        end = p_el.find(qn("a:endParaRPr"))
        if end is not None and end.get("sz"):
            sizes = [int(end.get("sz"))]
    return max(sizes) / 100 if sizes else _DEFAULT_SIZE_PT


def wrapped_lines(text: str, per_line: int) -> int:
    """Greedy word-wrap line count, the way PowerPoint breaks at spaces and
    only splits a word that is longer than a whole line."""
    if not text.strip():
        return 1
    lines, current = 1, 0
    for word in text.split():
        length = len(word)
        if current and current + 1 + length > per_line:
            lines += 1
            current = 0
        if length > per_line:
            lines += (length - 1) // per_line
            length = length % per_line or per_line
        current = length if not current else current + 1 + length
    return lines


def estimate_text_height(shape) -> int:
    """Rough rendered height (EMU) of a shape's text: characters per line
    from the box width and font size, wrapped line count, 1.2 line spacing."""
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))

    def inset(name, default):
        return int(body.get(name)) if body is not None and body.get(name) else default

    width = int(shape.width) - inset("lIns", 91440) - inset("rIns", 91440)
    height = inset("tIns", 45720) + inset("bIns", 45720)
    for p in shape.text_frame.paragraphs:
        size = _paragraph_size_pt(p._p)
        ppr = p._p.find(qn("a:pPr"))
        margin = int(ppr.get("marL", "0")) if ppr is not None else 0
        usable = max(width - margin, _EMU_PER_PT * size)
        per_line = max(1, int(usable / (size * CHAR_WIDTH_EM * _EMU_PER_PT)))
        lines = wrapped_lines(_p_text(p._p), per_line)
        height += int(lines * size * _LINE_SPACING * _EMU_PER_PT)
    return height


def _takeaway_shapes(deck: DeckInfo):
    for info in deck.of_kind(*_TAKEAWAY_KINDS):
        shape = _takeaway_shape(info)
        if shape is not None:
            yield info, shape


def _has_sp_autofit(shape) -> bool:
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))
    return body is not None and body.find(qn("a:spAutoFit")) is not None


MIN_BOTTOM_INSET = 91440  # 0.1 in
BREATHING_ROOM = 182880  # 0.2 in, about one line, under the last line of text
BOTTOM_MARGIN = 228600  # 0.25 in: a takeaway box never grows past this above the slide's bottom edge


def _bottom_inset(shape) -> int:
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))
    return int(body.get("bIns")) if body is not None and body.get("bIns") else 45720


def _padded_inset(shape) -> int:
    """Bottom inset with the breathing room: at least 0.1 in, plus 0.2 in.
    It lives in the inset (not just the height) so PowerPoint's own
    shape-to-fit keeps it when someone edits the text later."""
    current = _bottom_inset(shape)
    target = MIN_BOTTOM_INSET + BREATHING_ROOM
    return current if current >= target else max(current, MIN_BOTTOM_INSET) + BREATHING_ROOM


def box_height_cap(shape, slide_height: int) -> int:
    return slide_height - BOTTOM_MARGIN - int(shape.top)


def _target_height(shape, slide_height: int, original: int) -> int:
    """Fitted height, never below the original, capped at the bottom margin."""
    return max(original, min(estimate_text_height(shape), box_height_cap(shape, slide_height)))


def _find_r19(deck, report_type, jira_info):
    findings = []
    for info, shape in _takeaway_shapes(deck):
        padded = _bottom_inset(shape) >= MIN_BOTTOM_INSET + BREATHING_ROOM
        if padded and _has_sp_autofit(shape) and int(shape.height) >= _target_height(
                shape, deck.slide_height, 0):
            continue
        findings.append(Finding(
            "R19", info.index, f"Auto-fit the Key Takeaways box on {_slide_label(info)}",
            before=f"height {int(shape.height) / 914400:.2f} in",
            after="resize shape to fit text, 0.2 in clear below the last line",
            data={"shape_id": shape.shape_id, "slide_height": deck.slide_height},
            summary="Fit the Key Takeaways box to its text, with room below the last line"))
    return findings


def _apply_r19(prs, finding):
    shape = _shape(prs, finding)
    original = int(shape.height)
    shape.text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    shape.text_frame._txBody.find(qn("a:bodyPr")).set("bIns", str(_padded_inset(shape)))
    shape.height = _target_height(shape, int(prs.slide_height), original)


# ---------------------------------------------------------------- R20 Thank You

_OWNER_PLACEHOLDERS = ("owner name", "<ml team member name>")


def _thank_you_box(info: SlideInfo):
    """The owner name / title / email box: the one with placeholders or an
    email, else the only other text box with three or more paragraphs - so a
    box already holding someone else's details is replaced too."""
    candidates = [(shape, text.lower()) for shape, text in info.texts
                  if len(shape.text_frame.paragraphs) >= 3 and not text.lower().startswith("thank you")]
    for shape, low in candidates:
        if any(p in low for p in _OWNER_PLACEHOLDERS) or "@" in low:
            return shape
    return candidates[0][0] if len(candidates) == 1 else None


def owner_values(jira_info: dict | None) -> list[str]:
    """The ticket's Thank You Page Deck Owner name, title and email."""
    jira_info = jira_info or {}
    return [norm_text(jira_info.get(k) or "") for k in ("owner_name", "owner_title", "owner_email")]



def _find_r20(deck, report_type, jira_info):
    values = owner_values(jira_info)
    if not all(values):
        return []
    findings = []
    for info in deck.of_kind("thank_you"):
        shape = _thank_you_box(info)
        if shape is None:
            continue
        paras = shape.text_frame.paragraphs[:3]
        current = [norm_text(_p_text(p._p)) for p in paras]
        email_runs = [r for r in paras[2].runs if r.text.strip()]
        linked = bool(email_runs) and all(r.hyperlink.address == f"mailto:{values[2]}" for r in email_runs)
        if current == values and linked:
            continue
        findings.append(Finding(
            "R20", info.index, f"Fill the Thank You page owner details on {_slide_label(info)}",
            before=" / ".join(current), after=" / ".join(values) + " (email linked)",
            data={"shape_id": shape.shape_id, "values": values},
            summary=f"Fill the owner details: {values[0]}, {values[1]}, {values[2]} (email linked)"))
    return findings


def _apply_r20(prs, finding):
    shape = _shape(prs, finding)
    for p, value in zip(shape.text_frame.paragraphs[:3], finding.data["values"]):
        runs = p.runs
        if runs:
            runs[0].text = value
            for extra in runs[1:]:
                extra._r.getparent().remove(extra._r)
        else:
            run = p.add_run()
            run.text = value
    email_p = shape.text_frame.paragraphs[2]
    for run in email_p.runs:
        run.hyperlink.address = f"mailto:{finding.data['values'][2]}"


# ---------------------------------------------------------------- registry

RULES = [
    Rule("R1", "Pacing: add Imps / Leads after unit values", _find_r1, _apply_r1),
    Rule("R2", "Pacing: add a Total row", _find_r2, _apply_r2),
    Rule("R3", "Pacing: column widths, 12.5 in wide, centred", _find_r3, _apply_r3),
    Rule("R24", "Pacing: split a table that runs off the slide", _find_r24, _apply_r24),
    Rule("R22", "Delete slides for channels not in this report", _find_r22, _apply_delete, deletes=True),
    Rule("R4", "Delete slides with no data", _skip_off_channel(_find_r4), _apply_delete, deletes=True),
    Rule("R5", "Delete Country Insights when there is only one country", _skip_off_channel(_find_r5),
         _apply_delete, deletes=True),
    Rule("R6", "Delete Halo Effect slides in single-channel reports", _skip_off_channel(_find_r6), _apply_delete,
         deletes=True),
    Rule("R7", "Delete the Custom Question example slide", _skip_off_channel(_find_r7), _apply_delete,
         deletes=True),
    Rule("R8", "Delete Creative Sets slides in Content Syndication sections", _skip_off_channel(_find_r8),
         _apply_delete, deletes=True),
    Rule("R9", "Top Accounts: drop an all-zero CTV Impressions column (combined reports)", _find_r9, _apply_r9),
    Rule("R10", "Hyperlink the ML Platform URLs", _find_r10, _apply_r10),
    Rule("R11", "Round long decimals to 2 places", _find_r11, _apply_r11),
    Rule("R12", "Remove empty 'site visits generated' lines", _find_r12, _apply_r12),
    Rule("R13", "Campaign Highlights: format company sizes", _find_r13, _apply_r13),
    Rule("R26", "Campaign Highlights: remove duplicate list items", _find_r26, _apply_r26),
    Rule("R14", "Bold chart data labels (Country pie, Program Performance doughnut)", _find_r14, _apply_r14),
    Rule("R15", "Audience Reach takeaway: top intent topics", _find_r15, _apply_r15),
    Rule("R16", "Content Insights takeaway: top job titles", _find_r16, _apply_r16),
    Rule("R17", "Audience Insights takeaway", _find_r17, _apply_r17),
    Rule("R18", "Country Insights takeaway", _find_r18, _apply_r18),
    Rule("R23", "Halo Effect takeaway from its chart", _find_r23, _apply_r23),
    Rule("R25", "Audience Reach: fill a blank % from its chart", _find_r25, _apply_r25),
    Rule("R19", "Auto-fit Key Takeaways boxes", _find_r19, _apply_r19),
    Rule("R20", "Thank You page: owner name, title and linked email", _find_r20, _apply_r20),
    Rule("R21", "Top Accounts / Display tables 12.5 in wide, centred", _find_r21, _apply_r21),
]
RULES_BY_ID = {rule.id: rule for rule in RULES}

# Plain-language groups the review page shows the rules in: (id, title, rule ids).
RULE_GROUPS = [
    ("pacing", "Pacing table", ("R1", "R2", "R3", "R24")),
    ("remove", "Slides to remove", ("R22", "R4", "R5", "R6", "R7", "R8")),
    ("takeaways", "Key takeaways", ("R25", "R15", "R16", "R17", "R18", "R23", "R19")),
    ("links", "Links and Thank You page", ("R10", "R20")),
    ("cleanup", "Number and chart clean-up", ("R9", "R11", "R12", "R13", "R26", "R14", "R21")),
]
GROUP_OF_RULE = {rule_id: group_id for group_id, _title, ids in RULE_GROUPS for rule_id in ids}
# Rules that rewrite visible text: the review list shows their before -> after.
TEXT_CHANGE_RULES = ("R9", "R13", "R25", "R15", "R16", "R17", "R18", "R23")

# Execution order for non-deleting rules: column removal before width fitting,
# takeaway text before the auto-fit that sizes the boxes around it.
EXECUTION_ORDER = ["R1", "R2", "R3", "R9", "R21", "R10", "R11", "R12", "R13", "R26", "R14",
                   "R25", "R15", "R16", "R17", "R18", "R23", "R19", "R20"]
# Rules that add slides run after the deletions, on the final slide order.
AFTER_DELETE_ORDER = ["R24"]
