"""Scan and format a raw PPRA deck (bytes in, bytes out).

scan() reports what every rule would change on the uploaded deck; format()
applies the enabled rules in EXECUTION_ORDER, each one re-reading the deck as
it is at that point, then removes the slides the deleting rules picked (all
decided from the ORIGINAL deck's data), then runs the slide-adding rules
(AFTER_DELETE_ORDER: the pacing split) on the final slide order.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

from pptx import Presentation

from core.app_logging import get_logger
from core.ppra.blanks import DataIssue, find_data_issues
from core.ppra.detect import DeckInfo, SlideInfo, analyze, norm_text
from core.ppra.layout import has_hidden_shapes
from core.ppra.rules import (
    AFTER_DELETE_ORDER, EXECUTION_ORDER, RULES, RULES_BY_ID, Finding, data_rows, box_height_cap,
    estimate_text_height, halo_placeholder_index, owner_values, pacing_tables, reach_stat_missing, row_unit,
)

_P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


@dataclass
class AttentionItem:
    """One thing to do by hand before posting. message is a short action;
    slide_number is in the formatted deck (None = the whole deck)."""
    slide_number: int | None
    slide_title: str
    message: str

    def __str__(self) -> str:
        return f"{f'Slide {self.slide_number}' if self.slide_number else 'Deck'} · {self.message}"


def checklist_lines(items) -> list[str]:
    """One line per action, its slides merged: "Slides 9, 15 · Add the asset
    thumbnails". Deck-wide actions read "Deck · ..."."""
    order: list[str] = []
    slides: dict[str, set] = {}
    for item in items:
        if item.message not in slides:
            order.append(item.message)
            slides[item.message] = set()
        slides[item.message].add(item.slide_number)
    lines = []
    for message in order:
        numbers = sorted(n for n in slides[message] if n)
        if numbers:
            where = f"Slide{'s' if len(numbers) > 1 else ''} {', '.join(str(n) for n in numbers)}"
        else:
            where = "Deck"
        lines.append(f"{where} · {message}")
    return lines


@dataclass
class ScanResult:
    report_type: str | None
    detected_report_type: str | None
    findings: dict[str, list[Finding]] = field(default_factory=dict)  # every rule id present
    # Predicted "Before you post" items: what is still manual once every
    # found rule has run (slide numbers in that formatted deck).
    attention: list[AttentionItem] = field(default_factory=list)
    slide_titles: list[str] = field(default_factory=list)
    # Blank values the generator left in the ORIGINAL deck (slide numbers in
    # that deck; slides the formatter removes are left out).
    data_issues: list[DataIssue] = field(default_factory=list)

    def found_rule_ids(self) -> list[str]:
        return [rid for rid, items in self.findings.items() if items]

    def skipped_rule_ids(self) -> list[str]:
        return [rid for rid, items in self.findings.items() if not items]

    def change_count(self, rule_ids=None) -> int:
        """Findings of the non-deleting rules (all, or only rule_ids)."""
        return sum(len(items) for rid, items in self.findings.items()
                   if not RULES_BY_ID[rid].deletes and (rule_ids is None or rid in rule_ids))

    def slides_to_remove(self, rule_ids=None) -> list[int]:
        """Distinct 1-based slide numbers the deleting rules picked."""
        return sorted({f.slide_index + 1 for rid, items in self.findings.items()
                       if RULES_BY_ID[rid].deletes and (rule_ids is None or rid in rule_ids) for f in items})


def load(pptx_bytes: bytes):
    return Presentation(io.BytesIO(pptx_bytes))


def save(prs) -> bytes:
    # python-pptx names an added slide slide{count + 1}.xml and renumbers slide
    # parts only when prs.slides is first read, so adding slides after a
    # delete can reuse a part name an existing slide still holds; the zip then
    # carries two entries with that name and PowerPoint offers to repair it.
    # Renumber every slide part in deck order right before writing.
    prs.part.rename_slide_parts([entry.rId for entry in prs.slides._sldIdLst])
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def detect_report_type(pptx_bytes: bytes) -> str | None:
    return analyze(load(pptx_bytes)).detected_report_type


def scan(pptx_bytes: bytes, report_type: str | None, jira_info: dict | None = None) -> ScanResult:
    deck = analyze(load(pptx_bytes))
    findings = {rule.id: rule.applies(deck, report_type, jira_info) for rule in RULES}
    # The checklist is about the deck that gets posted, so it comes from a
    # trial format rather than the raw deck: placeholders the rules fill
    # (Thank You owner from the ticket, country / Halo takeaways, ...) must
    # not show up as manual work.
    try:
        _out, _log, attention = format_deck(pptx_bytes, report_type, jira_info)
    except Exception:  # noqa: BLE001 - fall back to the raw deck's view
        get_logger().exception("PPRA trial format failed; checklist taken from the raw deck")
        attention = attention_items(deck, report_type, jira_info)
    removed = {f.slide_index + 1 for rid, items in findings.items() if RULES_BY_ID[rid].deletes for f in items}
    return ScanResult(
        report_type=report_type, detected_report_type=deck.detected_report_type, findings=findings,
        attention=attention, slide_titles=[s.label for s in deck.slides],
        data_issues=find_data_issues(deck, skip_slides=removed),
    )


def delete_slides(prs, indices) -> None:
    """Remove slides cleanly: drop the sldId entry, then the presentation's
    relationship to the slide part (python-pptx only drops a rel once
    nothing references it), plus any section-list reference, so PowerPoint
    opens the file without a repair prompt."""
    sld_id_lst = prs.slides._sldIdLst
    entries = list(sld_id_lst)
    removed = {id(prs.slides[index].part) for index in set(indices)}
    for index in sorted(set(indices), reverse=True):
        entry = entries[index]
        slide_id, r_id = entry.get("id"), entry.rId
        sld_id_lst.remove(entry)
        prs.part.drop_rel(r_id)
        for ref in prs.part._element.iter(f"{{{_P14_NS}}}sldId"):
            if ref.get("id") == slide_id:
                ref.getparent().remove(ref)
    _drop_links_to(prs, removed)


def _drop_links_to(prs, removed: set) -> None:
    """A layout or slide can hold a slide-jump hyperlink to a deleted slide;
    that link keeps the slide part in the package under its old name, which
    a renumbered slide then also takes (two zip entries, repair prompt).
    Drop such links: the hyperlink elements and their relationships."""
    for part in list(prs.part.package.iter_parts()):
        if id(part) in removed:
            continue
        stale = [r_id for r_id, rel in part.rels.items()
                 if not rel.is_external and id(rel.target_part) in removed]
        if not stale:
            continue
        element = getattr(part, "_element", None)
        if element is not None:
            for el in list(element.iter()):
                if el.get(f"{{{_R_NS}}}id") in stale:
                    el.getparent().remove(el)
        for r_id in stale:
            part.rels.pop(r_id)


def _change_line(finding: Finding) -> str:
    detail = f" ({finding.before} -> {finding.after})" if finding.before or finding.after else ""
    return f"{finding.rule_id}: {finding.description}{detail}"


def format_deck(pptx_bytes: bytes, report_type: str | None, jira_info: dict | None = None,
                enabled_rule_ids=None) -> tuple[bytes, list[str], list[AttentionItem]]:
    enabled = set(enabled_rule_ids) if enabled_rule_ids is not None else {r.id for r in RULES}
    prs = load(pptx_bytes)
    original = analyze(prs)
    change_log: list[str] = []
    problems: list[tuple[int, str, str]] = []  # (slide id, label, message): ids survive deletes / inserts

    to_delete: dict[int, Finding] = {}
    for rule in RULES:
        if rule.deletes and rule.id in enabled:
            for finding in rule.applies(original, report_type, jira_info):
                to_delete.setdefault(finding.slide_index, finding)

    def run(rule_id: str, skip=()) -> None:
        rule = RULES_BY_ID[rule_id]
        deck = analyze(prs)
        for finding in rule.applies(deck, report_type, jira_info):
            if finding.slide_index in skip:
                continue
            try:
                rule.apply(prs, finding)
            except Exception:  # noqa: BLE001 - one bad slide must not sink the whole deck
                get_logger().exception("PPRA rule %s failed on slide %s", rule_id, finding.slide_index + 1)
                info = deck.slides[finding.slide_index]
                problems.append((info.slide.slide_id, info.label,
                                 f"Check this slide by hand ({rule_id} could not be applied)"))
                continue
            change_log.append(_change_line(finding))
            if finding.data.get("attention"):  # what the rule could not finish
                info = deck.slides[finding.slide_index]
                problems.append((info.slide.slide_id, info.label, finding.data["attention"]))

    for rule_id in EXECUTION_ORDER:
        if rule_id in enabled:
            run(rule_id, skip=to_delete)

    for index in sorted(to_delete):
        change_log.append(_change_line(to_delete[index]))
    delete_slides(prs, to_delete.keys())

    # Rules that insert slides (the pacing split) run on the final order.
    for rule_id in AFTER_DELETE_ORDER:
        if rule_id in enabled:
            run(rule_id)

    out = save(prs)
    final = analyze(load(out))
    position = {slide.slide_id: n for n, slide in enumerate(prs.slides, start=1)}
    failed = [AttentionItem(position[sid], label, message) for sid, label, message in problems if sid in position]
    return out, change_log, failed + attention_items(final, report_type, jira_info)


format = format_deck  # noqa: A001 - the spec's public name


# ---------------------------------------------------------------- attention list

def _reach_missing(info: SlideInfo) -> bool:
    shape = info.text_shape("key takeaways")
    if shape is None:
        return False
    text = norm_text(shape.text_frame.text)
    blank = "% of leads" in text and not any(ch.isdigit() for ch in text.split("% of leads")[0])
    return blank or reach_stat_missing(shape.text_frame)


def _takeaway_text(info: SlideInfo) -> str:
    shape = info.text_shape("key takeaways")
    return norm_text(shape.text_frame.text) if shape is not None else ""


MSG_PICK_TYPE = "Pick the report type"
MSG_THUMBNAILS = "Add the asset thumbnails"
MSG_OWNER = "Add the owner name, title and email (missing on the Jira ticket)"
MSG_OWNER_BOX = "Fill in the owner name, title and email"
MSG_HALO = "Fill in the Halo Effect takeaway number"
MSG_ERROR = "Replace the leftover ERROR value"
MSG_REACH = "Fill the % of leads from engaged accounts"
MSG_INDUSTRY = "Fill in the Industry Insights takeaway numbers"
MSG_OVERLAP = "Move or shrink the overlapping tables"
MSG_OVERFLOW = "Shorten the Key Takeaways text (it runs off the slide)"


def attention_items(deck: DeckInfo, report_type: str | None, jira_info: dict | None = None) -> list[AttentionItem]:
    """The "Before you post" checklist: only work the tool genuinely can't do
    on this deck (run on the formatted deck, so anything a rule fills is
    already gone). Agenda, Recommended Actions, Creative Sets, logos and
    the country placeholder are never listed."""
    items: list[AttentionItem] = []

    def add(info: SlideInfo | None, message: str) -> None:
        items.append(AttentionItem(info.number if info else None, info.label if info else "", message))

    if report_type is None:
        add(None, MSG_PICK_TYPE)
    for info, shape, rows, cols in pacing_tables(deck):
        for r in data_rows(rows):
            if row_unit(rows[r][cols["campaign"]], rows[r], cols, deck, report_type) is None:
                add(info, f"Add Imps or Leads to the '{norm_text(rows[r][cols['campaign']])}' row")
    owner_missing = not all(owner_values(jira_info))
    for info in deck.slides:
        takeaway = _takeaway_text(info)
        halo_left = False
        if info.kind == "content_insights" and any(len(rows) > 1 for _s, rows in info.tables):
            add(info, MSG_THUMBNAILS)
        elif info.kind == "halo":
            shape = info.text_shape("key takeaways")
            halo_left = shape is not None and halo_placeholder_index(shape.text_frame) is not None
            if halo_left:
                add(info, MSG_HALO)
        elif info.kind == "audience_reach" and _reach_missing(info):
            add(info, MSG_REACH)
        elif info.kind == "industry_insights" and "(% of total)" in takeaway:
            add(info, MSG_INDUSTRY)
        elif info.kind == "thank_you" and owner_missing:
            add(info, MSG_OWNER)
        elif info.kind == "thank_you" and (info.has_text("owner name") or info.has_text("<ml team member name>")):
            add(info, MSG_OWNER_BOX)  # ticket has the details but the owner box wasn't found
        if info.has_text("error:") and not halo_left:
            add(info, MSG_ERROR)
        if has_hidden_shapes(info, deck.slide_width, deck.slide_height):
            add(info, MSG_OVERLAP)
        shape = info.text_shape("key takeaways")
        if shape is not None and (int(shape.top) + int(shape.height) > deck.slide_height
                                  or estimate_text_height(shape) > box_height_cap(shape, deck.slide_height)):
            add(info, MSG_OVERFLOW)
    # Blank values still in the deck ("Manufacturing - %", "{{2}}"), beyond
    # what the items above already cover.
    for issue in find_data_issues(deck, for_checklist=True):
        add(deck.slides[issue.slide_number - 1], issue.checklist_message())
    return items
