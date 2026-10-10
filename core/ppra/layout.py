"""Shapes hidden behind a table or text box that renders taller than its frame.

python-pptx (and the generator) store a table's row heights as minimums: a
row of long asset names wraps to two lines and PowerPoint draws it taller,
down over the heading and table below it. Here every content shape gets an
estimated rendered height (wrapped text, at least the stored height), and a
shape that starts below another one's stored bottom but above its estimated
bottom is hidden. The fix moves the shapes underneath down (cascading down
the column), and when that runs past the slide's usable bottom, fits the
involved tables' rows to their text and steps their body text down (never
below 9 pt). If nothing fits, the slide is left alone and stays on the
checklist.
"""
from __future__ import annotations

from dataclasses import dataclass

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

from core.ppra.detect import SlideInfo, chart_title, norm_text
from core.ppra.rules import (
    Finding, _cell_font_pt, _cell_height, _p_text, _paragraph_size_pt, estimate_text_height, usable_bottom,
)

GAP = 91440  # 0.1 in between a grown shape and the one moved under it
LAYER_TOLERANCE = 45720  # a shape starting this far inside another's frame is layered on it on purpose
TEXT_TOLERANCE = 91440  # text estimates are rough: ignore less than 0.1 in of growth
MIN_MOVE = 9144  # 0.01 in
MIN_TABLE_PT = 9.0
_SKIP_PLACEHOLDERS = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.FOOTER,
                      PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE)


@dataclass
class _Box:
    shape: object
    kind: str  # "table", "text" or "other"
    top: int
    left: int
    width: int
    height: int

    @property
    def right(self) -> int:
        return self.left + self.width


def _skip_placeholder(shape) -> bool:
    if not shape.is_placeholder:
        return False
    try:
        return shape.placeholder_format.type in _SKIP_PLACEHOLDERS
    except ValueError:
        return False


def _content_boxes(info: SlideInfo, slide_width: int, slide_height: int, bottom: int) -> list[_Box]:
    """Top-level tables, text, charts, pictures and groups in the content
    area: not the title, not full-slide backgrounds, not the footer band."""
    boxes = []
    for shape in info.slide.shapes:
        if shape.shape_id == info.title_shape_id or _skip_placeholder(shape):
            continue
        if shape._element.tag == qn("p:cxnSp"):
            continue
        if any(v is None for v in (shape.top, shape.left, shape.width, shape.height)):
            continue
        top, left, width, height = int(shape.top), int(shape.left), int(shape.width), int(shape.height)
        if width >= slide_width * 0.9 and height >= slide_height * 0.8:
            continue  # background picture / panel
        if top >= bottom:
            continue  # footer band: confidential line, logo
        if getattr(shape, "has_table", False) and shape.has_table:
            kind = "table"
        elif shape.has_text_frame and norm_text(shape.text_frame.text):
            kind = "text"
        else:
            kind = "other"
        boxes.append(_Box(shape, kind, top, left, width, height))
    return boxes


def _shrunk(size: float, steps: int) -> float:
    return size if size <= MIN_TABLE_PT else max(MIN_TABLE_PT, size - steps)


def table_row_heights(frame, shrink: int = 0, compact: bool = False) -> list[int]:
    """Estimated rendered row heights (EMU): each row's tallest wrapped cell,
    at least the stored height - or exactly the text height when compact.
    shrink steps the body rows' text down that many points (min 9 pt)."""
    widths = [int(c.width) for c in frame.table.columns]
    heights = []
    for r, tr in enumerate(frame.table._tbl.tr_lst):
        text_height, c = 0, 0
        for tc in tr.findall(qn("a:tc")):
            span = int(tc.get("gridSpan", "1"))
            width = sum(widths[c:c + span])
            c += span
            if tc.get("hMerge") == "1" or tc.get("vMerge") == "1":
                continue
            size = _cell_font_pt(tc)
            if r and shrink:
                size = _shrunk(size, shrink)
            paras = list(tc.iter(qn("a:p")))
            texts = [_p_text(p) for p in paras]
            indents = [int(p.find(qn("a:pPr")).get("marL", "0")) if p.find(qn("a:pPr")) is not None else 0
                       for p in paras]
            text_height = max(text_height, _cell_height(tc, texts, width, size, indents))
        stored = int(tr.get("h", "0"))
        heights.append(text_height if compact and text_height else max(stored, text_height))
    return heights


def _text_extent(shape, stored: int) -> int:
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))
    if body is not None and (body.get("wrap") == "none" or body.find(qn("a:normAutofit")) is not None):
        return stored  # no wrapping, or PowerPoint shrinks the text to fit
    if body is not None and body.get("anchor") == "b":
        return stored  # bottom-anchored text grows upwards
    grow = estimate_text_height(shape) - stored
    # One line of estimating error is normal (narrow letters fit more per
    # line than the estimate allows), so a box must overflow by more.
    one_line = int(max(_paragraph_size_pt(p._p) for p in shape.text_frame.paragraphs) * 1.2 * 12700)
    if grow <= max(TEXT_TOLERANCE, one_line):
        return stored
    if body is not None and body.get("anchor") == "ctr":
        grow //= 2
    return stored + grow


def _heights(boxes: list[_Box], shrink: int, compact: bool, tuned: set) -> list[int]:
    heights = []
    for i, box in enumerate(boxes):
        if box.kind == "table":
            heights.append(sum(table_row_heights(box.shape, shrink if i in tuned else 0, compact and i in tuned)))
        elif box.kind == "text":
            heights.append(_text_extent(box.shape, box.height))
        else:
            heights.append(box.height)
    return heights


def _side_by_side(a: _Box, b: _Box) -> bool:
    return min(a.right, b.right) - max(a.left, b.left) <= MIN_MOVE * 5


def _plan(boxes: list[_Box], heights: list[int]) -> tuple[dict, dict]:
    """New tops (moved boxes only ever go down) and, per moved box, the box
    that pushed it."""
    order = sorted(range(len(boxes)), key=lambda i: (boxes[i].top, boxes[i].left))
    new_top, cause = {}, {}
    for pos, j in enumerate(order):
        b = boxes[j]
        need, pusher = b.top, None
        for i in order[:pos]:
            a = boxes[i]
            if _side_by_side(a, b):
                continue
            stored_bottom = a.top + a.height
            if b.top < stored_bottom - LAYER_TOLERANCE:
                continue  # starts inside a's frame: layered on purpose (a label on a chart, ...)
            if new_top[i] == a.top and heights[i] <= a.height:
                continue  # a is drawn as stored: the original layout stands
            new_bottom = new_top[i] + heights[i]
            if new_bottom <= b.top:
                continue
            original_gap = b.top - stored_bottom
            required = new_bottom + (min(original_gap, GAP) if original_gap > 0 else GAP)
            if required > need:
                need, pusher = required, i
        if need - b.top < MIN_MOVE:
            need, pusher = b.top, None
        new_top[j] = need
        if pusher is not None:
            cause[j] = pusher
    return new_top, cause


@dataclass
class Fix:
    boxes: list
    heights: list
    new_top: dict
    cause: dict
    tuned: set
    shrink: int
    compact: bool
    fits: bool

    @property
    def moved(self) -> list[int]:
        return sorted(self.cause, key=lambda j: self.boxes[j].top)


def _attempt(boxes, bottom, shrink, compact, tuned) -> Fix:
    heights = _heights(boxes, shrink, compact, tuned)
    new_top, cause = _plan(boxes, heights)
    touched = set(cause) | (tuned if compact or shrink else set())
    fits = all(new_top[j] + heights[j] <= bottom for j in touched)
    return Fix(boxes, heights, new_top, cause, tuned, shrink, compact, fits)


def _max_shrink(boxes, tuned) -> int:
    steps = 0
    for i in tuned:
        for tr in boxes[i].shape.table._tbl.tr_lst[1:]:
            for tc in tr.findall(qn("a:tc")):
                steps = max(steps, int(_cell_font_pt(tc) - MIN_TABLE_PT))
    return steps


def plan_slide(info: SlideInfo, slide_width: int, slide_height: int) -> tuple[Fix | None, Fix | None]:
    """(as found, fix): as found is None when nothing is hidden; fix is the
    first layout that fits above the usable bottom, or None."""
    bottom = usable_bottom(info.slide, slide_height)
    boxes = _content_boxes(info, slide_width, slide_height, bottom)
    found = _attempt(boxes, bottom, 0, False, set())
    if not found.cause:
        return None, None
    if found.fits:
        return found, found
    tuned = {i for i in set(found.cause) | set(found.cause.values()) if boxes[i].kind == "table"}
    if not tuned:
        return found, None
    for shrink in range(0, _max_shrink(boxes, tuned) + 1):
        fix = _attempt(boxes, bottom, shrink, True, tuned)
        if fix.fits:
            return found, fix
    return found, None


def has_hidden_shapes(info: SlideInfo, slide_width: int, slide_height: int) -> bool:
    found, _fix = plan_slide(info, slide_width, slide_height)
    return found is not None


def _name(box: _Box) -> str:
    if box.kind == "table":
        first = norm_text(box.shape.table.cell(0, 0).text) if len(box.shape.table.rows) else ""
        return f"the '{first[:40]}' table" if first else "the table"
    if box.kind == "text":
        return f"'{norm_text(box.shape.text_frame.text)[:40]}'"
    if getattr(box.shape, "has_chart", False) and box.shape.has_chart:
        title = chart_title(box.shape.chart)
        return f"the '{title[:40]}' chart" if title else "the chart"
    return f"'{box.shape.name}'"


def describe(fix: Fix) -> str:
    moved = fix.moved
    if not moved:  # fitting the rows (and smaller text) alone cleared it
        names = " and ".join(_name(fix.boxes[i]) for i in sorted(fix.tuned, key=lambda i: fix.boxes[i].top))
        size = f" at {_body_size(fix):g} pt" if fix.shrink else ""
        return f"Fit {names} to {'its' if len(fix.tuned) == 1 else 'their'} text{size} so nothing below is hidden"
    first = moved[0]
    grower = fix.cause[first]
    while grower in fix.cause:  # name the shape that grew, not one it pushed
        grower = fix.cause[grower]
    what = _name(fix.boxes[first])
    if len(moved) > 1:
        what += f" and {len(moved) - 1} more shape{'s' if len(moved) > 2 else ''}"
    text = f"Move {what} below {_name(fix.boxes[grower])} ({'they were' if len(moved) > 1 else 'it was'} hidden)"
    if fix.shrink:
        text += f"; table text down to {_body_size(fix):g} pt"
    elif fix.compact:
        text += "; table rows fitted to their text"
    return text


def _body_size(fix: Fix) -> float:
    """The largest body text size left in the tuned tables."""
    return max((_shrunk(_cell_font_pt(tc), fix.shrink) for i in fix.tuned
                for tr in fix.boxes[i].shape.table._tbl.tr_lst[1:] for tc in tr.findall(qn("a:tc"))),
               default=MIN_TABLE_PT)


def _set_size(rpr, steps: int, default: float) -> None:
    current = int(rpr.get("sz") or default * 100) / 100
    new = _shrunk(current, steps)
    if new != current or not rpr.get("sz"):
        rpr.set("sz", str(int(round(new * 100))))


def _apply_fix(fix: Fix) -> None:
    for i in fix.tuned:
        frame = fix.boxes[i].shape
        if fix.shrink:
            for tr in frame.table._tbl.tr_lst[1:]:
                for tc in tr.findall(qn("a:tc")):
                    size = _cell_font_pt(tc)
                    for r in tc.iter(qn("a:r")):
                        rpr = r.find(qn("a:rPr"))
                        if rpr is None:
                            rpr = OxmlElement("a:rPr")
                            r.insert(0, rpr)
                        _set_size(rpr, fix.shrink, size)
                    for end in tc.iter(qn("a:endParaRPr")):
                        _set_size(end, fix.shrink, size)
        if fix.compact:
            rows = table_row_heights(frame, compact=True)  # after the font change
            for tr, height in zip(frame.table._tbl.tr_lst, rows):
                tr.set("h", str(height))
            frame.height = sum(rows)
    for j, top in fix.new_top.items():
        if top != fix.boxes[j].top:
            fix.boxes[j].shape.top = top


def find_hidden(deck) -> list[Finding]:
    findings = []
    for info in deck.slides:
        found, fix = plan_slide(info, deck.slide_width, deck.slide_height)
        if fix is None:
            continue
        summary = describe(fix)
        findings.append(Finding(
            "R29", info.index, f"{summary} on slide {info.number} ({info.label})",
            before="hidden behind a taller shape", after="moved below it", data={},
            summary=summary))
    return findings


def apply_hidden(prs, finding: Finding) -> None:
    from core.ppra.detect import parse_slides  # the slide as it is now

    info = parse_slides(prs)[finding.slide_index]
    _found, fix = plan_slide(info, int(prs.slide_width), int(prs.slide_height))
    if fix is not None:
        _apply_fix(fix)
