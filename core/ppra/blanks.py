"""Blank values the report generator left in a deck.

The platform fills its templates from data; when a value is missing it
leaves the template showing: a "%" with no number in front ("Manufacturing
- %", "% of companies had 1-9 employees"), a "{{2}}" token, an "ERROR:"
string, or a sentence with an empty substitution ("Top job titles for leads
from  include  and ."). This module only detects them; rules.py (R12, R28)
decides what it may change, and the rest is listed for the user.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from pptx.oxml.ns import qn

from core.ppra.detect import DeckInfo, iter_shapes, norm_text

TOKEN_RE = re.compile(r"\{\{[^{}]*\}\}|\{\{|\}\}")
# A connector word followed by a gap where the value should have been.
_EMPTY_AFTER_RE = re.compile(
    r"\b(?:from|include|includes|including|and|or|of|with|by|at|for|to|in|on|were|was|is|are)\s{2,}(?=\S)",
    re.IGNORECASE)
# (A colon is a list lead-in - "closely aligned with:" - not a gap.)
_EMPTY_END_RE = re.compile(r"\s(?:and|or|include|includes|including|from|of|with|by)\s*[.,;]\s*$", re.IGNORECASE)
_SPACE_DOT_RE = re.compile(r"\S\s+\.\s*$")
_ALNUM_RE = re.compile(r"[A-Za-z0-9]")
# "%" / "K+" with no number in front (spaces skipped), "$" with none after.
_BLANK_PCT_RE = re.compile(r"(?:^|[^\d\s])\s*%")
_BLANK_K_RE = re.compile(r"(?:^|[^\w.])\s*K\+")
_BLANK_DOLLAR_RE = re.compile(r"\$(?!\s*[\d.])")
_PLACEHOLDER_X_RE = re.compile(r"^\s*[xX]\s*%?\s*$")
_ERROR_RE = re.compile(r"error\s*:", re.IGNORECASE)


def raw_text(p_el) -> str:
    """A paragraph's text with its spacing kept (an empty substitution shows
    as a double space), only non-breaking spaces made plain."""
    parts = []
    for child in p_el:
        if child.tag in (qn("a:r"), qn("a:fld")):
            t = child.find(qn("a:t"))
            parts.append(t.text or "" if t is not None else "")
        elif child.tag == qn("a:br"):
            parts.append(" ")
    return "".join(parts).replace("\xa0", " ")


def broken_reason(text: str) -> str | None:
    """Why a sentence is a broken template ('token', 'empty'), or None."""
    if not text.strip():
        return None
    if TOKEN_RE.search(text):
        return "token"
    # One double space can be a typo; a dangling end ("... and .") or two
    # gaps after connector words is a template that got no values.
    if _EMPTY_END_RE.search(text) or _SPACE_DOT_RE.search(text) or len(_EMPTY_AFTER_RE.findall(text)) >= 2:
        return "empty"
    return None


def only_punctuation(text: str) -> bool:
    return bool(text.strip()) and not _ALNUM_RE.search(text)


def blank_kind(text: str) -> str | None:
    """'%' for a blank percentage, 'value' for any other blank value
    (token, ERROR, "$", "K+", "X", empty substitution), else None."""
    if not text.strip():
        return None
    if _BLANK_PCT_RE.search(text):
        return "%"
    if (TOKEN_RE.search(text) or _ERROR_RE.search(text) or _BLANK_K_RE.search(text)
            or _BLANK_DOLLAR_RE.search(text) or _PLACEHOLDER_X_RE.match(text) or broken_reason(text)):
        return "value"
    return None


@dataclass
class DataIssue:
    """Blank values on one slide (numbers are in the deck that was checked)."""
    slide_number: int
    slide_title: str
    count: int
    example: str
    kinds: set = field(default_factory=set)

    @property
    def only_percent(self) -> bool:
        return self.kinds == {"%"}

    def line(self) -> str:
        what = f"{self.count} blank value{'s' if self.count != 1 else ''}"
        return f"Slide {self.slide_number} · {self.slide_title} · {what} (e.g. '{self.example}')"

    def checklist_message(self) -> str:
        noun = "% value" if self.only_percent else "value"
        return f"Fill {self.count} blank {noun}{'s' if self.count != 1 else ''}"


# Slides whose Key Takeaways box a dedicated rule fills or a dedicated
# checklist item already covers (the rest of the slide is still checked).
_HANDLED_KINDS = ("audience_reach", "industry_insights", "halo")


def _paragraph_texts(slide, skip_takeaway: bool = False):
    """Every text paragraph on a slide: shapes (groups included) and table
    body cells (header rows hold labels like "% Delivered", never values)."""
    for shape in iter_shapes(slide.shapes):
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in list(shape.table.rows)[1:]:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        yield raw_text(p._p)
        elif shape.has_text_frame:
            if skip_takeaway and norm_text(shape.text_frame.text).lower().startswith("key takeaways"):
                continue
            for p in shape.text_frame.paragraphs:
                yield raw_text(p._p)


def find_data_issues(deck: DeckInfo, skip_slides=(), for_checklist: bool = False) -> list[DataIssue]:
    """Blank values per slide. skip_slides: 1-based numbers left out (slides
    the formatter removes). for_checklist leaves out what another checklist
    item or rule already covers (Audience Reach %, Industry and Halo
    takeaways, ERROR values)."""
    issues = []
    for info in deck.slides:
        if info.number in skip_slides:
            continue
        count, example, kinds = 0, "", set()
        for text in _paragraph_texts(info.slide, for_checklist and info.kind in _HANDLED_KINDS):
            if for_checklist and _ERROR_RE.search(text):
                continue
            kind = blank_kind(text)
            if kind is None:
                continue
            count += 1
            kinds.add(kind)
            example = example or norm_text(text)[:60]
        if count:
            issues.append(DataIssue(info.number, info.label, count, example, kinds))
    return issues
